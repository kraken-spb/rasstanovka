import argparse
import json
from pathlib import Path

import openpyxl


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser(description="Подготовить справочники из файла расстановки")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    reference = workbook[workbook.sheetnames[0]]
    placement = workbook[workbook.sheetnames[1]]

    objects = {}
    organizations = set()
    professions = set()
    for row in reference.iter_rows(min_row=2, values_only=True):
        object_name, subobject_name, contractor, employer, profession = map(clean, row[:5])
        if object_name and object_name.lower() != "список":
            objects.setdefault(object_name, set())
            if subobject_name and subobject_name.lower() != "список":
                objects[object_name].add(subobject_name)
        for organization in (contractor, employer):
            if organization and organization.lower() != "список":
                organizations.add(organization)
        if profession and profession.lower() != "список":
            professions.add(profession)

    workers = []
    initial_assignments = []
    for row in placement.iter_rows(min_row=3, values_only=True):
        values = [clean(value) for value in row]
        if not values[5] or not values[6]:
            continue
        worker = {
            "full_name": values[5],
            "personnel_no": values[6],
            "contractor": values[3],
            "employer": values[4],
            "profession": values[9] or values[8] or values[7],
        }
        workers.append(worker)
        if values[1] and values[1].lower() != "список":
            objects.setdefault(values[1], set())
            if values[2] and values[2].lower() != "список":
                objects[values[1]].add(values[2])
        initial_assignments.append(
            {
                "object": values[1],
                "subobject": values[2],
                "personnel_no": values[6],
                "employer": values[4],
                "foreman": values[11],
                "shift": values[12],
            }
        )

    payload = {
        "objects": [
            {"name": name, "subobjects": sorted(subobjects)}
            for name, subobjects in sorted(objects.items())
        ],
        "organizations": sorted(organizations),
        "professions": sorted(professions),
        "workers": workers,
        "initial_assignments": initial_assignments,
    }
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"Готово: {len(payload['objects'])} объектов, "
        f"{sum(len(item['subobjects']) for item in payload['objects'])} подобъектов, "
        f"{len(workers)} сотрудников"
    )


if __name__ == "__main__":
    main()
