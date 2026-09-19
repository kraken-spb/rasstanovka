"""Scoped hostel requests: explicit travel details, immutable export snapshots."""
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from textwrap import TextWrapper
from math import ceil
from PIL import ImageFont

from flask import abort, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter

from workforce_bulk import ids_from
from workforce_core import (EDITORS, PROFILE_SELECT, actor_scope, audit, date_value,
                            digest, plain, remember, replay, text_value, uuid_value)

REQUEST_ROLES = EDITORS | {'hr_viewer'}
COMPANY = 'АО «Ленгазспецстрой»'
TITLE = 'ЗАЯВКА НА ЗАСЕЛЕНИЕ В ХОСТЕЛ (ПВП)'
BRAND_BLUE = '0073BC'  # Official website: lgss-spb.ru/static/main.css, header__logo.
HEADERS = ('№', 'Табельный №', 'ФИО', 'Должность', 'Должность ГДЛР', 'Дата рождения',
           'Гражданство', 'Телефон', 'Объект строительства', 'Категория работника (ИТР/ рабочий)',
           'Дата и время заезда', 'Дата выезда (ориентировочно)', 'Ответственный работник УРП', 'Основание')
EDITABLE = {'object', 'worker_category', 'arrival_date', 'arrival_time',
            'departure_date', 'responsible', 'basis'}


def request_actor(db):
    actor, scope, args = actor_scope(db)
    if actor['role'] not in REQUEST_ROLES:
        abort(403, description='Заявки на проживание доступны службам учёта персонала и администраторам.')
    return actor, scope, args


def source_rows(db, ids):
    actor, scope, args = request_actor(db)
    profiles = plain(db.native(PROFILE_SELECT + ' WHERE w.id=ANY(%s) AND (' + scope + ')',
                               [ids, *args]).fetchall())
    if len(profiles) != len(ids):
        abort(404, description='Один из выбранных сотрудников недоступен в ваших СМУ. Обновите список.')
    if any(not row['active'] for row in profiles):
        abort(409, description='В выборе есть сотрудник, выведенный из состава. Обновите список.')
    # Classification follows current database category bindings and its hierarchy,
    # never a guessed profession or an outdated Excel qualification.
    nodes = {r['id']: r for r in db.native('SELECT id,parent_id,name_key,category_id FROM gdlr_hierarchy_nodes')}
    by_category = {r['category_id']: r for r in nodes.values() if r['category_id'] is not None}
    def category_for(category_id):
        node, seen = by_category.get(category_id), set()
        while node and node['id'] not in seen:
            seen.add(node['id'])
            if node['name_key'] in ('итр', 'рабочие'):
                return 'ИТР' if node['name_key'] == 'итр' else 'Рабочий'
            node = nodes.get(node['parent_id'])
        return ''
    planned = {r['worker_id']: plain(r['planned_arrival']) for r in db.native('''
        SELECT DISTINCT ON(worker_id) worker_id,planned_arrival FROM workforce_pvp_stays
        WHERE worker_id=ANY(%s) AND departed_on IS NULL AND planned_arrival IS NOT NULL
        ORDER BY worker_id,updated_at DESC,id''', (ids,))}
    author = db.native('SELECT full_name FROM users WHERE id=%s', (actor['id'],)).fetchone()[0]
    profiles = {p['id']: p for p in profiles}
    rows = []
    for worker_id in ids:
        p = profiles[worker_id]
        row = {key: p.get(key) or '' for key in ('id','personnel_no','full_name','profession','category',
                                               'birth_date','citizenship','phone','accommodation')}
        row.update(object=p.get('project') or '', worker_category=category_for(p['category_id']),
                   arrival_date=planned.get(worker_id) or '', arrival_time='', departure_date='',
                   responsible=author or '', basis='')
        row['token'] = digest([p, row])
        rows.append(row)
    return actor, rows, author


def clean_row(value):
    if not isinstance(value, dict) or set(value) != EDITABLE | {'id', 'token'}:
        abort(400, description='Проверьте поля строки заявки.')
    result = {key: text_value(value[key], label, maximum) for key, label, maximum in (
        ('object', 'Объект строительства', 500), ('responsible', 'Ответственный УРП', 240),
        ('basis', 'Основание', 2000))}
    if value['worker_category'] not in ('', 'ИТР', 'Рабочий'):
        abort(400, description='Уточните категорию работника: ИТР или Рабочий.')
    result['worker_category'] = value['worker_category']
    result['arrival_date'] = date_value(value['arrival_date'], 'Дата заезда') or ''
    result['departure_date'] = date_value(value['departure_date'], 'Ориентировочная дата выезда') or ''
    try:
        if value['arrival_time'] != '':
            clock = time.fromisoformat(value['arrival_time'])
            if clock.tzinfo is not None or clock.isoformat(timespec='minutes') != value['arrival_time']:
                raise ValueError()
    except (ValueError, TypeError):
        abort(400, description='Укажите время заезда в формате ЧЧ:ММ.')
    result['arrival_time'] = value['arrival_time']
    if result['departure_date'] and result['arrival_date'] and result['departure_date'] < result['arrival_date']:
        abort(400, description='Ориентировочный выезд не может быть раньше заезда.')
    return result


def request_workbook(rows, day):
    """A4 landscape form; continue down pages rather than shrinking the whole list."""
    book = Workbook(); sheet = book.active; sheet.title = 'Заявка на заселение'
    book.properties.creator = COMPANY; book.properties.title = TITLE
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells('A1:D2')
    logo = Image(str(Path(__file__).parent / 'static' / 'lgss-print-logo.png'))
    logo.width = 300; logo.height = 300 * 48 / 273
    sheet.add_image(logo, 'A1')
    sheet.merge_cells('E1:N1'); sheet['E1'] = COMPANY
    sheet['E1'].font = Font(name='Arial', size=18, bold=True, color=BRAND_BLUE)
    sheet.merge_cells('E2:N2'); sheet['E2'] = TITLE
    sheet['E2'].font = Font(name='Arial', size=13, bold=True, color='001437')
    for address in ('E1','E2'):
        sheet[address].alignment = Alignment(horizontal='right', vertical='center')
    sheet.row_dimensions[1].height = 26; sheet.row_dimensions[2].height = 24
    sheet.merge_cells('A3:H3'); sheet['A3'] = 'Дата заявки: ' + date.fromisoformat(day).strftime('%d.%m.%Y')
    sheet.merge_cells('I3:N3'); sheet['I3'] = 'Сотрудников: ' + str(len(rows))
    for address in ('A3','I3'):
        sheet[address].font = Font(name='Arial',size=11,color='001437')
        sheet[address].alignment = Alignment(horizontal='right' if address=='I3' else 'left',vertical='center')
    sheet.row_dimensions[3].height = 24; sheet.row_dimensions[4].height = 7
    sheet.row_dimensions[5].height = 70
    widths = (4, 12, 23, 20, 18, 13, 14, 16, 15, 13, 22, 17, 21, 22)
    border = Border(*(Side(style='thin', color='BCCCD8') for _ in range(4)))
    font = Font(name='Arial', size=11, color='001437')
    measure_font = ImageFont.truetype(str(Path(__file__).parent / 'fonts' / 'DejaVuSans.ttf'),15)
    alignment = Alignment(vertical='center', wrap_text=True)
    for index, (header, width) in enumerate(zip(HEADERS, widths), 1):
        cell = sheet.cell(5, index, header);cell.font = Font(name='Arial',size=10,bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor=BRAND_BLUE)
        cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True);cell.border = border
        sheet.column_dimensions[get_column_letter(index)].width = width
    cursor = 6
    for index, row in enumerate(rows, 1):
        if row['arrival_date'] and row['arrival_time']:
            arrival = datetime.fromisoformat(row['arrival_date'] + 'T' + row['arrival_time'])
        elif row['arrival_date']:
            arrival = date.fromisoformat(row['arrival_date'])
        elif row['arrival_time']:
            arrival = time.fromisoformat(row['arrival_time'])
        else:
            arrival = None
        values = (index,row['personnel_no'],row['full_name'],row['profession'],row['category'],
                  date.fromisoformat(row['birth_date']) if row['birth_date'] else None,
                  row['citizenship'],row['phone'],row['object'],row['worker_category'],arrival,
                  date.fromisoformat(row['departure_date']) if row['departure_date'] else None,
                  row['responsible'],row['basis'])
        # Explicit continuation rows keep long reasons readable within Excel's row-height limit.
        lines = []
        for value,width in zip(values,widths):
            display = value.strftime('%d.%m.%Y\n%H:%M') if isinstance(value,datetime) else value.strftime('%d.%m.%Y') if isinstance(value,date) else value.strftime('%H:%M') if isinstance(value,time) else str(value or '')
            # Use the bundled, slightly wider font to leave safe space for Excel/Arial.
            for chars in range(max(2,int(width)-1),0,-1):
                wrapper=TextWrapper(width=chars,drop_whitespace=False,replace_whitespace=False)
                wrapped=[part for line in display.split('\n') for part in (wrapper.wrap(line) or [''])]
                if all(measure_font.getlength(part)<=width*7-10 for part in wrapped):break
            lines.append(wrapped)
        parts = max(1,ceil(max(map(len,lines))/20))
        for part in range(parts):
            sheet.row_dimensions[cursor].height = max(36,10+15*max(len(items[part*20:(part+1)*20]) for items in lines))
            for column,value in enumerate(values,1):
                if parts>1 and isinstance(value,str) and len(lines[column-1])>20:
                    value = '\n'.join(lines[column-1][part*20:(part+1)*20])
                elif part and column not in (1,2,3):
                    value = None
                if part and column==3:value = 'Продолжение: ' + str(value)
                cell=sheet.cell(cursor,column,value);cell.font=font;cell.alignment=alignment;cell.border=border
                cell.fill=PatternFill('solid',fgColor='F0F6FA' if index%2==0 else 'FFFFFF')
                if isinstance(value,str):cell.data_type='s'
                elif isinstance(value,datetime):cell.number_format='dd.mm.yyyy hh:mm'
                elif isinstance(value,date):cell.number_format='dd.mm.yyyy'
                elif isinstance(value,time):cell.number_format='hh:mm'
                if column in (1,2,6,7,10,11,12):cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cursor+=1
    sheet.freeze_panes='D6';sheet.auto_filter.ref=f'A5:N{max(5,cursor-1)}'
    sheet.print_title_rows='1:5';sheet.print_area=f'A1:N{max(5,cursor-1)}'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.page_setup.horizontalDpi=300;sheet.page_setup.verticalDpi=300
    sheet.page_margins=PageMargins(left=.2,right=.2,top=.25,bottom=.3,header=.1,footer=.12)
    sheet.print_options.horizontalCentered=True
    sheet.oddFooter.left.text=COMPANY;sheet.oddFooter.right.text='Страница &P из &N'
    for footer in (sheet.oddFooter.left,sheet.oddFooter.right):
        footer.size=8;footer.font='Arial';footer.color='0073BC'
    stream=BytesIO();book.save(stream);return stream.getvalue()


def register_accommodation_routes(app, database, roles_required):
    @app.post('/api/workforce/accommodation/prepare')
    @roles_required(*REQUEST_ROLES)
    def workforce_accommodation_prepare():
        data=request.get_json(silent=True)
        if not isinstance(data,dict) or set(data)!={'ids','date'}:
            abort(400,description='Выберите сотрудников и дату заявки.')
        ids=ids_from(data['ids']);day=date_value(data['date'],'Дата заявки',True)
        db=database()
        with db:
            db.execute('BEGIN');_,rows,author=source_rows(db,ids)
        return jsonify(date=day,rows=rows,responsible_default=author)

    @app.post('/api/workforce/accommodation/export')
    @roles_required(*REQUEST_ROLES)
    def workforce_accommodation_export():
        data=request.get_json(silent=True)
        if not isinstance(data,dict) or set(data)!={'date','rows','request_key'} or not isinstance(data['rows'],list):
            abort(400,description='Проверьте состав заявки.')
        if any(not isinstance(r,dict) for r in data['rows']):
            abort(400,description='Проверьте строки заявки.')
        ids=ids_from([r.get('id') for r in data['rows']]);day=date_value(data['date'],'Дата заявки',True)
        uuid_value(data['request_key'],'Ключ заявки')
        edits={r['id']:clean_row(r) for r in data['rows']}
        db=database()
        with db:
            db.execute('BEGIN IMMEDIATE');actor,current,_=source_rows(db,ids)
            # Fresh role and scope are checked even when replaying a completed export.
            snapshot=replay(db,actor,'accommodation_request',data)
            if snapshot is None:
                posted={r['id']:r for r in data['rows']}
                if any(posted[r['id']]['token']!=r['token'] for r in current):
                    abort(409,description='Данные сотрудников изменились. Откройте заявку заново и проверьте значения.')
                rows=[{**r,**edits[r['id']]} for r in current]
                snapshot={'date':day,'rows':rows}
                audit(db,actor,'export','accommodation_request',data['request_key'],None,snapshot,
                      reason='Формирование заявки на заселение в хостел')
                remember(db,actor,'accommodation_request',data,snapshot)
        response=send_file(BytesIO(request_workbook(snapshot['rows'],snapshot['date'])),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,download_name='Заявка_на_заселение_'+day+'.xlsx',max_age=0)
        response.headers['Cache-Control']='private, no-store'
        response.headers['X-Export-Row-Count']=str(len(snapshot['rows']))
        return response
