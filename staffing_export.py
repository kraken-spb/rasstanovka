from filter_values import argument as filter_argument, values as filter_values, matches as filter_matches, label as filter_label
"""XLSX export of the actual daily placement records."""
import io
import re
from user_smu_access import worker_clause
from datetime import date

from flask import abort, g, request, send_file
from report_queries import SHIFT_LABELS, assignment_rows, filter_assignment_rows
from report_matrix import summary_sheet
from native_pivot import save_workbook_with_slicer
from staffing_import import active_members_sql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    '№\nп/п', 'Группа подобъектов', 'Подобъект', 'Компания подрядчик',
    'Организация-работодатель', 'ФИО работника', 'Таб. № с префиксом',
    'Должность по штатному расписанию', 'Профессия ГСП', 'Категория ГДЛР',
    'ФИО линейного ИТР', 'ФИО бригадира', 'Смена', 'СМУ', 'Выполняемые операции',
]


def _period():
    raw_day = request.args.get('date', '')
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', raw_day or ''):
        abort(400, description='Укажите дату расстановки в формате ГГГГ-ММ-ДД.')
    try:
        day = date.fromisoformat(raw_day).isoformat()
    except ValueError:
        abort(400, description='Укажите существующую дату расстановки.')
    shift = filter_argument('shift', allowed={'all','1 смена','2 смена'})
    if 'all' in filter_values(shift) or len(filter_values(shift))==2: shift='all'
    if shift not in {'all', '1 смена', '2 смена'}:
        abort(400, description='Выберите смену или все смены.')
    return day, shift

def _rows(db, day, shifts):
    # Compatibility for existing callers; the query belongs to the read model.
    return assignment_rows(db, day, shifts, user=g.user)


def _string(cell, value):
    cell.value = '' if value is None else str(value)
    cell.data_type = 's'


def _unassigned_rows(db, day, requested_shift):
    """Current placement roster, once per worker, with no assignment on this date."""
    access, scope_params = worker_clause(db)
    effective_shift = "CASE WHEN ss.worker_id IS NULL THEN '1 смена' ELSE ss.shift END"
    return db.execute(f'''
        WITH roster AS (
            SELECT worker_id FROM ({active_members_sql()})
            UNION SELECT worker_id FROM outstaff_members
            UNION SELECT worker_id FROM manual_employees
        )
        SELECT NULL id, NULL object_id, NULL subobject_id,
               w.employer display_company,
               '' object_name, '' subobject_name,
               w.full_name,w.personnel_no,COALESCE(ct.name,w.contractor) contractor,
               w.employer assignment_employer,w.profession,w.gsp_profession,w.department,
               COALESCE(gc.name,w.category) category,
               d.linear_itr_override,d.brigadier_override,
               c.linear_itr crew_linear_itr,c.brigadier crew_brigadier,
               {effective_shift} normalized_shift,pw.description performed_work
        FROM roster r JOIN workers w ON w.id=r.worker_id
        LEFT JOIN crew_members m ON m.worker_id=w.id LEFT JOIN crews c ON c.id=m.crew_id
        LEFT JOIN employee_contractors ew ON ew.worker_id=w.id LEFT JOIN contractors ct ON ct.id=ew.contractor_id
        LEFT JOIN employee_gdlr eg ON eg.worker_id=w.id LEFT JOIN gdlr_categories gc ON gc.id=eg.category_id
        LEFT JOIN staffing_row_details d ON d.worker_id=w.id
        LEFT JOIN staffing_shifts ss ON ss.worker_id=w.id AND ss.work_date=?
        LEFT JOIN staffing_performed_work pw ON pw.worker_id=w.id AND pw.work_date=? AND pw.shift={effective_shift}
        WHERE w.active=1 AND NOT EXISTS (
            SELECT 1 FROM assignments a WHERE a.worker_id=w.id AND a.work_date=?)
          AND (?='all' OR {effective_shift}=?) AND ({access})
        ORDER BY w.full_name COLLATE NOCASE,w.personnel_no COLLATE NOCASE,w.id
    ''', [day, day, day, requested_shift, requested_shift, *scope_params]).fetchall()


def _sheet(workbook, title, rows, include_unassigned=False):
    sheet = workbook.create_sheet(title)
    fill = PatternFill('solid', fgColor='0070C0')
    font = Font(name='Times New Roman', size=12, bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    headers = HEADERS + (['Статус расстановки'] if include_unassigned else [])
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column)
        _string(cell, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = header_alignment
    sheet.row_dimensions[1].height = 60
    for number, row in enumerate(rows, 1):
        values = (
            number, row['object_name'], row['subobject_name'], row['contractor'],
            row['assignment_employer'], row['full_name'], row['personnel_no'], row['profession'],
            row['gsp_profession'], row['category'],
            row['linear_itr_override'] if row['linear_itr_override'] is not None else row['crew_linear_itr'],
            row['brigadier_override'] if row['brigadier_override'] is not None else row['crew_brigadier'],
            'Без смены' if row['normalized_shift'] is None else SHIFT_LABELS[row['normalized_shift']], row['department'], row['performed_work'],
        )
        if include_unassigned:
            values += ('Не расставлен' if row['id'] is None else 'Расставлен',)
        for column, value in enumerate(values, 1):
            cell = sheet.cell(number + 1, column)
            if column == 1:
                cell.value = value
            else:
                _string(cell, value)
                if column == 7:
                    cell.number_format = '@'
            cell.font = Font(name='Times New Roman', size=12)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        capacities = (6, 20, 24, 20, 24, 24, 18, 28, 20, 20, 24, 24, 10, 28, 44, 20)
        lines = max(1, *(sum(max(1, (len(line) + capacity - 1) // capacity)
                            for line in str(value or '').split('\n'))
                         for value, capacity in zip(values, capacities)))
        sheet.row_dimensions[number + 1].height = min(409, max(32, 16 * lines))
    widths = (7, 24, 28, 24, 28, 30, 20, 32, 24, 24, 28, 28, 12, 32, 48)
    if include_unassigned:
        widths += (24,)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'


def register_staffing_export_route(app, get_db, roles_required):
    @app.get('/api/staffing/export')
    @roles_required('admin', 'foreman')
    def staffing_export():
        day, requested_shift = _period()
        kind = request.args.get('kind', 'report')
        if kind not in ('report', 'staffing', 'summary'):
            abort(400, description='Выберите вид отчёта.')
        include_unassigned = request.args.get('include_unassigned', '0')
        if include_unassigned not in ('0', '1'):
            abort(400, description='Некорректная настройка выгрузки нерасставленных сотрудников.')
        include_unassigned = include_unassigned == '1'
        category = filter_argument('category')
        department = filter_argument('department',500,ignore_empty=True)
        contractor = filter_argument('contractor',500,ignore_empty=True)
        query = request.args.get('query', '')
        if len(query) > 500:
            abort(400, description='Слишком длинный фильтр отчёта.')
        db = get_db()
        shifts = tuple(SHIFT_LABELS) if requested_shift == 'all' else (requested_shift,)
        db.execute('BEGIN')
        try:
            rows = _rows(db, day, shifts)
            unassigned = _unassigned_rows(db, day, requested_shift) if include_unassigned else []
            rows = filter_assignment_rows([*rows, *unassigned], category=category,
                                          department=department, contractor=contractor, query=query)
            unassigned = [row for row in rows if row['id'] is None]
        finally:
            db.commit()
        count = len(rows)
        if not count:
            if contractor:
                return {'error': 'По выбранным фильтрам компании-подрядчика нет сотрудников для выгрузки.'}, 404
            if category is not None or query:
                return {'error': 'По выбранным фильтрам нет сотрудников для выгрузки.'}, 404
            if department:
                return {'error': 'По выбранным дате, смене и СМУ нет сотрудников для выгрузки.'}, 404
            if include_unassigned:
                return {'error': 'За выбранные дату и смену нет сотрудников для выгрузки.'}, 404
            return {'error': 'За выбранные дату и смену нет расставленных сотрудников.'}, 404
        workbook = Workbook()
        workbook.remove(workbook.active)
        # Both tabs use the exact same authorized, filtered snapshot.
        _sheet(workbook, 'Список сотрудников', rows, include_unassigned)
        summary_sheet(workbook, rows, day, filter_label(category,'','Без категории') if category is not None else None, query)
        output = io.BytesIO()
        save_workbook_with_slicer(workbook, output)
        output.seek(0)
        filename = f'Расстановка на {day}'
        if requested_shift != 'all':
            filename += ' день' if requested_shift == '1 смена' else ' ночь'
        response = send_file(output, as_attachment=True, download_name=filename + '.xlsx',
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response.headers['X-Export-Count'] = str(count)
        response.headers['X-Export-Unassigned-Count'] = str(len(unassigned))
        return response
