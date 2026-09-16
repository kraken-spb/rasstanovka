"""The agreed two-sheet URP report, from relational records at a selected date."""
import io
from datetime import date, timedelta

from flask import jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from workforce_core import READERS, actor_scope, date_value


REPORT_CATEGORIES = (
    'Антикоррозийщик', 'Арматурщик', 'Бетонщик', 'Дефектоскописты', 'Изолировщик',
    'Монтажник СиЖБК', 'Монтажник ТТ', 'Прочие монтажники', 'Прочие основные рабочие',
    'Прочие сварщики и газорезчики', 'Сварщик АиПАМ', 'Сварщик МК', 'Сварщик ТТ',
    'Электромонтажник', 'Линейные ИТР', 'Мастер строительных и монтажных работ', 'Прораб', 'Старший прораб',
)
FIRST_HEADERS = ['Проект', 'Компания подрядчик', 'Организация-работодатель', 'Гражданство',
                 'Дата заезда', 'Прогноз окончания вахты', 'ФИО', 'Табельный номер',
                 'Должность/профессия', 'Профессия согласно ГДЛР 3 уровень', 'Примечания']
SECOND_HEADERS = ['Проект', 'Компания подрядчик', 'Организация-работодатель', 'Статус сотрудника',
                  'Гражданство', 'ФИО', 'Таб. Номер', 'Должность/профессия',
                  'Профессия согласно ГДЛР 3 уровень', 'Тип заезда/выезда', 'Плановая дата',
                  'Заезд/ выезд', 'Статус заезда/выезда', 'Примечания']


def report_records(db, day):
    actor, scope, params = actor_scope(db)
    rows = db.native('''SELECT w.id,w.full_name,w.personnel_no::text personnel_no,w.personnel_is_internal,
        w.profession,w.department,COALESCE(ct.name,w.contractor) contractor,
        o.name employer,ci.label citizenship,em.label employment,p.employment_code,
        p.arrival_date,p.forecast_departure_date,p.leave_end_date,p.notes,p.pure_outstaff,
        COALESCE(gc.name,w.category) category,pr.label project,st.stage_code,
        trip.id movement_id,trip.direction,trip.planned_date,ba.label basis,re.label result,
        old.planned_date previous_planned_date,ro.planned_end_date rotation_end,ro.next_arrival_date,
        EXISTS(SELECT 1 FROM workforce_source_records sr WHERE sr.worker_id=w.id AND sr.source_role<>'outstaff') mixed_sources,
        (SELECT string_agg(DISTINCT cf.description,E'\\n') FROM workforce_conflicts cf WHERE cf.worker_id=w.id AND cf.state='open') conflicts
        FROM workers w JOIN workforce_profiles p ON p.worker_id=w.id
        LEFT JOIN workforce_organizations o ON o.id=p.employer_id
        LEFT JOIN workforce_catalog ci ON ci.code=p.citizenship_code
        LEFT JOIN workforce_catalog em ON em.code=p.employment_code
        LEFT JOIN employee_contractors ec ON ec.worker_id=w.id LEFT JOIN contractors ct ON ct.id=ec.contractor_id
        LEFT JOIN employee_gdlr eg ON eg.worker_id=w.id LEFT JOIN gdlr_categories gc ON gc.id=eg.category_id
        LEFT JOIN employee_smu es ON es.worker_id=w.id
        LEFT JOIN workforce_smu_projects sp ON sp.smu_id=es.smu_id LEFT JOIN workforce_catalog pr ON pr.code=sp.project_code
        LEFT JOIN LATERAL(SELECT e.stage_code FROM workforce_stage_events e WHERE e.worker_id=w.id AND e.effective_date<=%s
            AND e.confirmed AND NOT e.retracted AND NOT EXISTS(SELECT 1 FROM workforce_stage_events r WHERE r.replaces_id=e.id AND (r.confirmed OR r.retracted) AND r.effective_date<=%s)
            ORDER BY e.effective_date DESC,e.sequence DESC LIMIT 1) st ON TRUE
        LEFT JOIN LATERAL(SELECT m.* FROM workforce_movements m WHERE m.worker_id=w.id AND m.destination_kind<>'pvp'
            AND NOT EXISTS(SELECT 1 FROM workforce_movements next WHERE next.rescheduled_from=m.id)
            ORDER BY (m.result_code IS NULL) DESC,(m.planned_date>=%s) DESC NULLS LAST,
                CASE WHEN m.planned_date>=%s THEN m.planned_date END ASC,
                CASE WHEN m.planned_date<%s THEN m.planned_date END DESC,m.updated_at DESC,m.id LIMIT 1) trip ON TRUE
        LEFT JOIN workforce_movements old ON old.id=trip.rescheduled_from
        LEFT JOIN workforce_catalog ba ON ba.code=trip.basis_code LEFT JOIN workforce_catalog re ON re.code=trip.result_code
        LEFT JOIN LATERAL(SELECT r.planned_end_date,r.next_arrival_date FROM workforce_rotations r WHERE r.worker_id=w.id
            AND NOT r.cancelled AND r.start_date<=%s AND COALESCE(r.actual_end_date,r.next_arrival_date)>=%s
            ORDER BY r.start_date DESC LIMIT 1) ro ON TRUE
        WHERE w.active=1 AND lower(COALESCE(gc.name,w.category))=ANY(%s) AND (''' + scope + ''')
        ORDER BY w.name_search,w.id''', [day, day, day, day, day, day, day,
                                      [name.lower() for name in REPORT_CATEGORIES], *params]).fetchall()
    if actor['role'] in {'foreman', 'viewer'}:
        return [{**dict(row), 'notes':'', 'conflicts':''} for row in rows]
    return rows


def rows_for_report(records):
    tabs = [[], []]
    for record in records:
        row = dict(record)
        tab = row['personnel_no'] if not row['personnel_is_internal'] and row['personnel_no'].isdigit() else ''
        notes = [row['notes'] or '', row['conflicts'] or '']
        if not row['project']:
            notes.append('Проект не определён по СМУ: ' + (row['department'] or 'СМУ не указан') + '.')
        if row['previous_planned_date']:
            notes.append('Поездка перенесена с ' + row['previous_planned_date'].strftime('%d.%m.%Y') + '.')
        common = [row['project'] or '', row['contractor'] or '', row['employer'] or '']
        citizenship = (row['citizenship'] or '').upper()
        pure_outstaff = row.get('pure_outstaff', False) or row['employment_code'] in ('employment.external', 'employment.internal') and not row['mixed_sources']
        first = row['stage_code'] == 'stage.onsite' or pure_outstaff
        if first:
            values = common + [citizenship, row['arrival_date'], row['rotation_end'] or row['forecast_departure_date'],
                               row['full_name'], tab, row['profession'], row['category']]
        else:
            planned = row['planned_date']
            direction = 'Заезд' if planned and row['direction'] == 'arrival' else 'Выезд' if planned and row['direction'] == 'departure' else ''
            basis = row['basis'] or ''
            if not planned and row['stage_code'] == 'stage.leave' and row['leave_end_date']:
                planned, direction = row['leave_end_date'] + timedelta(days=2), 'Заезд'
            if not planned and row['next_arrival_date']:
                planned, direction, basis = row['next_arrival_date'], 'Заезд', 'по графику'
            if not row['stage_code']:
                notes.append('Фактическое состояние сотрудника на отчётную дату не подтверждено.')
            values = common + ['Штат' if tab else row['employment'] or '', citizenship, row['full_name'], tab,
                               row['profession'], row['category'], basis, planned, direction, row['result'] or '']
        values.append('\n'.join(dict.fromkeys(text.strip() for text in notes if text.strip())))
        tabs[0 if first else 1].append(values)
    return tabs


def workbook_bytes(tabs):
    book = Workbook()
    book.remove(book.active)
    for number, (name, headers, rows) in enumerate(zip(('Явка и аутстаффинг', 'Неявка, заезд и ПВП'),
                                                       (FIRST_HEADERS, SECOND_HEADERS), tabs)):
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='176D50')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[1].height = 45
        for row_number, values in enumerate(rows, 2):
            sheet.append(values)
            for column in range(1, len(headers) + 1):
                cell = sheet.cell(row_number, column)
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Workbook data cannot become executable Excel formulas.
                if isinstance(cell.value, date):
                    cell.number_format = 'DD.MM.YYYY'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11)
            sheet.cell(row_number, 8 if number == 0 else 7).number_format = '@'
        widths = (22, 22, 26, 18, 18, 22, 34, 20, 34, 32, 64) if number == 0 else (22, 22, 26, 24, 18, 34, 20, 34, 32, 20, 18, 16, 20, 64)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1,len(rows)+1)}'
        sheet.print_title_rows = '1:1'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        if number == 1:
            for column, values in [('D', 'Штат,Трудоустройство ИРС,Трудоустройство,Аутстафф внешний,Аутстафф внутренний'),
                                    ('J', 'на билете,заявка,по графику,без заявки'), ('L', 'Заезд,Выезд'),
                                    ('M', 'состоялся,перенос,отмена')]:
                validation = DataValidation(type='list', formula1='"' + values + '"', allow_blank=True)
                validation.errorTitle = 'Выберите значение из списка'
                validation.error = 'Используйте утверждённый перечень.'
                validation.showErrorMessage = True
                sheet.add_data_validation(validation)
                validation.add(f'{column}2:{column}{max(2,len(rows)+1)}')
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    return content


def register_workforce_export(app, database, roles_required):
    @app.get('/api/workforce/export')
    @roles_required(*READERS)
    def workforce_export():
        day = date_value(request.args.get('date'), 'Отчётная дата', True)
        db = database()
        with db:
            db.execute('BEGIN')
            tabs = rows_for_report(report_records(db, day))
        response = send_file(workbook_bytes(tabs), as_attachment=True,
                             download_name=f'УРП — учёт персонала — {day}.xlsx',
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response.headers['Cache-Control'] = 'no-store'
        response.headers['X-Export-Count'] = str(sum(map(len, tabs)))
        response.headers['X-Export-First-Count'] = str(len(tabs[0]))
        response.headers['X-Export-Second-Count'] = str(len(tabs[1]))
        return response
