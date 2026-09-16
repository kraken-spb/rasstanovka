"""Explicit Excel source mapping. Unknown fields and conflicts remain reviewable."""
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ALIASES = {
    'name': ('фио', 'фиоработника', 'физическоелицо', 'сотрудник'),
    'tab': ('таб', 'табномер', 'табельныйномер'),
    'employer': ('организация', 'работодатель', 'организацияработодатель', 'контрагент'),
    'profession': ('должность', 'должностьпрофессия'),
    'category': ('категориягдлр', 'категориягдлрлгсс', 'гдлр', 'профессиясогласногдлр3уровень'),
    'dob': ('датарождения',), 'citizenship': ('гражданство',), 'phone': ('телефон', 'номертелефона'),
    'smu': ('сму',), 'department': ('подразделение', 'подразделениеорганизации'),
    'arrival': ('датазаезда', 'датазаезданаучасток', 'датаприбытиянаучасток'),
    'departure': ('датавыезда', 'датавыездасучастка', 'выездсппс15'),
    'planned_arrival': ('пландатазаезда',), 'ticket_arrival': ('датаприбытияпобилетам',),
    'leave_start': ('датаначалаотпускамо',), 'leave_end': ('датаокончания',),
    'ticket_home': ('статусбилетадомой',), 'ticket_site': ('статусбилетанаучасток',),
    'home_trip': ('датавремягородвыезда', 'датавремявыезда'),
    'site_flight': ('датавылетанаучасток', 'датавремявылета'),
    'transit_arrival': ('датавремяприбытиявиркутск', 'датавремяприбытия'),
    'city': ('город',), 'origin': ('городотправки',), 'action': ('приемперевод',),
    'responsible': ('ответственныйсотрудник', 'спецт'), 'notes': ('примечание', 'дополнительно', 'примечания'),
    'ok': ('ок',), 'migration': ('му',), 'medical': ('медкомиссия',), 'safety': ('от',),
    'accommodation': ('статуспребывания', 'хостел'), 'attestation': ('аттестация',),
    'hangar': ('ангар',), 'employment_type': ('вид',), 'site': ('объектвыполненияработ',),
    'vehicle': ('регномер',), 'umiat': ('умиат',), 'messenger': ('мессенджер',), 'application': ('заявление',),
}
REVERSE = {alias: field for field, values in ALIASES.items() for alias in values}
DATE_FIELDS = {'dob', 'arrival', 'departure', 'planned_arrival', 'ticket_arrival', 'leave_start', 'leave_end', 'migration'}
SHEETS = {'rotation': {'явка', 'неявка', 'заезд', 'пвп', 'аустаффинг', 'аутстаффинг'},
          'recruitment': {'пвп', 'патентыпвп', 'патентывпвп'}}


class SourceProblem(ValueError):
    pass


def normalize(value):
    return re.sub(r'[^а-яa-z0-9]', '', str(value or '').casefold().replace('ё', 'е'))


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def source_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = text(value)
    for fmt in ('%Y-%m-%d', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S', '%d.%m.%Y', '%d.%m.%Y %H:%M'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    prefix = re.match(r'^(\d{2}\.\d{2}\.\d{4})(?=\s|$)', value)
    if prefix:
        try:
            return datetime.strptime(prefix[1], '%d.%m.%Y').date()
        except ValueError:
            pass
    return None


def parse_workbook(content, filename, service):
    if service not in SHEETS:
        raise SourceProblem('Укажите службу: перевахта или комплектация.')
    if not filename.lower().endswith('.xlsx') or len(content) > 20 * 1024 * 1024:
        raise SourceProblem('Загрузите файл XLSX размером до 20 МБ.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise SourceProblem('Слишком большой распакованный файл Excel.')
            if any('vbaproject' in item.filename.lower() for item in entries):
                raise SourceProblem('Файл с макросами не поддерживается.')
            for item in entries:
                if item.filename.lower().endswith('.xml'):
                    with archive.open(item) as stream:
                        carry = b''
                        while chunk := stream.read(65536):
                            scanned = carry + chunk
                            if b'<!DOCTYPE' in scanned or b'<!ENTITY' in scanned:
                                raise SourceProblem('Внешние сущности и DTD в Excel не поддерживаются.')
                            carry = scanned[-16:]
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except SourceProblem:
        raise
    except (BadZipFile, KeyError, OSError, ValueError) as error:
        raise SourceProblem('Не удалось прочитать структуру Excel.') from error
    records, excluded, mappings = [], [], []
    try:
        for sheet in book:
            normalized_sheet = normalize(sheet.title)
            if normalized_sheet not in SHEETS[service]:
                excluded.append({'sheet': sheet.title, 'reason': 'Лист исключён правилами импорта.'})
                continue
            if (sheet.max_column or 0) > 128:
                raise SourceProblem(f'Лист «{sheet.title}»: более 128 столбцов. Уберите лишние форматированные столбцы.')
            rows = sheet.iter_rows(values_only=True)
            headers = None
            for header_row, values in enumerate(rows, 1):
                mapping = {i: REVERSE.get(normalize(value)) for i, value in enumerate(values)}
                if 'name' in mapping.values() and len(set(mapping.values()) - {None}) >= 2:
                    headers = values
                    break
                if header_row >= 25:
                    break
            if headers is None:
                raise SourceProblem(f'Лист «{sheet.title}»: не найдены заголовки ФИО и кадровых полей.')
            for i, value in enumerate(headers):
                if normalize(value).startswith('графикпо1с'):
                    mapping[i] = 'schedule'
            name_column = next(i for i, field in mapping.items() if field == 'name')
            mappings.append({'sheet': sheet.title, 'header_row': header_row,
                             'columns': [{'column': get_column_letter(i+1), 'header': text(value), 'field': mapping[i]}
                                         for i, value in enumerate(headers) if text(value)]})
            section = ''
            for number, values in enumerate(rows, header_row + 1):
                if number > 25000 or len(records) > 20000:
                    raise SourceProblem('Слишком много строк. Разделите файл по подразделениям.')
                name = text(values[name_column]) if name_column < len(values) else ''
                if service == 'recruitment' and not name and len(values) > 2 and text(values[2]):
                    section = text(values[2])
                if not name or normalize(name) in ALIASES['name']:
                    continue
                if normalized_sheet in {'аустаффинг', 'аутстаффинг'} and (len(values) < 7 or not text(values[6])):
                    excluded.append({'sheet': sheet.title, 'row': number, 'reason': 'Пустой исходный столбец G с ГДЛР.'})
                    continue
                fields, raw, notes = {}, {}, []
                for i, value in enumerate(values[:24] if service == 'recruitment' else values):
                    if value is None:
                        continue
                    raw[get_column_letter(i+1)] = {'header': text(headers[i]) if i < len(headers) else '', 'value': text(value)}
                    field = mapping.get(i)
                    if not field:
                        continue
                    parsed_date = source_date(value) if field in DATE_FIELDS else None
                    mapped = parsed_date.isoformat() if parsed_date else text(value)
                    if fields.get(field) and fields[field] != mapped:
                        notes.append(f'Конфликт столбцов «{field}»: {fields[field]} / {mapped}.')
                    elif mapped:
                        fields[field] = mapped
                fields['name'] = name
                fields.setdefault('tab', '')
                if fields['tab'] and not re.fullmatch('[0-9]+', fields['tab']):
                    notes.append('Отметка вместо табельного номера: ' + fields['tab'])
                    fields['tab'] = ''
                if fields['tab'] and not fields.get('employer'):
                    fields['employer'] = 'ЛГСС'
                if not fields.get('smu'):
                    match = re.search(r'(?:СМУ\s*[-№]?\s*|Строительно-монтажный участок\s*№?\s*)(\d+(?:\.\d+)*)', fields.get('department', ''), re.I)
                    if match:
                        fields['smu'] = 'СМУ-' + match[1]
                records.append({'filename': filename, 'sheet': sheet.title, 'row': number, 'service': service,
                                'fields': fields, 'raw': raw, 'mapping_notes': notes, 'section': section})
    finally:
        book.close()
    if not records:
        raise SourceProblem('На разрешённых листах нет подходящих сотрудников.')
    return {'records': records, 'excluded': excluded, 'mappings': mappings}
