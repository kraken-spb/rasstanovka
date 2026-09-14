"""Reviewed outstaff imports, with source lineage and conservative catalog matching."""
import hashlib
from user_smu_access import require_all
import io
import json
import re
import secrets
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from flask import abort, g, jsonify, request
from itsdangerous import BadSignature, URLSafeTimedSerializer

from staffing_import import ImportProblem, backup_database


def normalized(value):
    return ' '.join(str(value or '').casefold().replace('ё', 'е').split())


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


def migrate_outstaff(db):
    db.executescript('''
        CREATE TABLE IF NOT EXISTS outstaff_imports (
            id INTEGER PRIMARY KEY, request_key TEXT NOT NULL UNIQUE, filename TEXT NOT NULL,
            sheet TEXT NOT NULL, file_hash TEXT NOT NULL, selection_json TEXT NOT NULL,
            imported_by INTEGER NOT NULL REFERENCES users(id), imported_at TEXT NOT NULL,
            selected_count INTEGER NOT NULL
        );
        CREATE TABLE IF NOT EXISTS outstaff_members (
            worker_id INTEGER PRIMARY KEY REFERENCES workers(id), identity_key TEXT NOT NULL UNIQUE,
            import_id INTEGER NOT NULL REFERENCES outstaff_imports(id), source_row INTEGER NOT NULL,
            source_department TEXT NOT NULL, source_category TEXT NOT NULL, work_kind TEXT NOT NULL,
            vendor TEXT NOT NULL, staff_type TEXT NOT NULL, arrival TEXT NOT NULL, departure TEXT NOT NULL,
            edit_token TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS outstaff_import_rows (
            import_id INTEGER NOT NULL REFERENCES outstaff_imports(id), source_row INTEGER NOT NULL,
            worker_id INTEGER NOT NULL REFERENCES workers(id), source_json TEXT NOT NULL,
            mapping_json TEXT NOT NULL, PRIMARY KEY(import_id,source_row)
        );
    ''')


def parse_outstaff(content, filename):
    if len(content) > 20_000_000:
        raise ImportProblem('Размер файла превышает 20 МБ.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 150_000_000:
                raise ImportProblem('Распакованный файл слишком велик.')
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (ValueError, OSError, zipfile.BadZipFile, KeyError) as exc:
        raise ImportProblem('Не удалось прочитать XLSX.') from exc
    try:
        sheets = [s for s in book if normalized(s.title) in ('аутстафф', 'аутстаффинг', 'аустаффинг')]
        if len(sheets) != 1:
            raise ImportProblem('Нужен один лист «Аутстафф», «Аутстаффинг» или «Аустаффинг».')
        sheet = sheets[0]
        if (sheet.max_row or 0) > 50_000 or (sheet.max_column or 0) > 100:
            raise ImportProblem('Лист превышает допустимые размеры: 50 000 строк и 100 столбцов.')
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, ())
        headers = [str(v).strip() if v is not None else '' for v in raw_headers]
        aliases = {'full_name': ('фио', 'ф.и.о.'), 'profession': ('должность',),
                   'work_kind': ('вид работ',), 'department': ('сму',), 'vendor': ('контрагент',),
                   'category': ('категория гдлр лгсс',), 'staff_type': ('вид',),
                   'arrival': ('дата заезда',), 'departure': ('дата выезда',)}
        columns = {}
        for field, names in aliases.items():
            found = [i for i, h in enumerate(headers) if normalized(h) in names]
            if len(found) != 1:
                raise ImportProblem(f'Нужен один столбец «{names[0]}».')
            columns[field] = found[0]
        rows = []
        for number, values in enumerate(iterator, 2):
            cells = [v.isoformat()[:10] if isinstance(v, (date, datetime)) else
                     str(int(v)) if isinstance(v, float) and v.is_integer() else
                     str(v).strip() if v is not None else '' for v in values]
            if not any(cells):
                continue
            if len(rows) >= 20_000 or any(len(v) > 1000 for v in cells):
                raise ImportProblem('Слишком много строк или слишком длинное значение ячейки.')
            row = {field: cells[index] if index < len(cells) else '' for field, index in columns.items()}
            if not row['full_name']:
                raise ImportProblem(f'Строка {number}: не заполнено ФИО.')
            rows.append({**row, 'source_row': number, 'cells': cells})
        if not rows:
            raise ImportProblem('На листе нет сотрудников.')
        return {'filename': Path(filename).name, 'sheet': sheet.title,
                'file_hash': hashlib.sha256(content).hexdigest(), 'headers': headers, 'rows': rows}
    finally:
        book.close()


def filtered_rows(parsed, filters):
    if not isinstance(filters, list) or len(filters) > 30:
        raise ImportProblem('Укажите не более 30 условий отбора.')
    rows = parsed['rows']
    for condition in filters:
        if not isinstance(condition, dict):
            raise ImportProblem('Некорректное условие отбора.')
        column, operation, values = condition.get('column'), condition.get('op'), condition.get('values', [])
        if (type(column) is not int or not 0 <= column < len(parsed['headers'])
                or operation not in ('in', 'not_in', 'contains', 'empty', 'not_empty', 'before', 'after')
                or not isinstance(values, list) or len(values) > 20000
                or any(not isinstance(v, str) or len(v) > 1000 for v in values)):
            raise ImportProblem('Некорректное условие отбора.')
        if operation not in ('empty', 'not_empty') and not values:
            raise ImportProblem('Выберите значение условия отбора.')
        if operation in ('before', 'after'):
            try:
                date.fromisoformat(values[0])
            except ValueError:
                raise ImportProblem('Для сравнения дат укажите дату в формате ГГГГ-ММ-ДД.')
        choices = set(values)
        def matches(row):
            value = row['cells'][column] if column < len(row['cells']) else ''
            if operation == 'in': return value in choices
            if operation == 'not_in': return value not in choices
            if operation == 'contains': return normalized(values[0]) in normalized(value)
            if operation == 'empty': return not value
            if operation == 'not_empty': return bool(value)
            try:
                actual = date.fromisoformat(value)
            except ValueError:
                return False
            boundary = date.fromisoformat(values[0])
            return actual <= boundary if operation == 'before' else actual >= boundary
        rows = [row for row in rows if matches(row)]
    return rows


def department_number(value):
    match = re.fullmatch(r'(?:сму|строительно[-– ]монтажный участок)\s*№?\s*(\d+(?:\.\d+)*)', normalized(value))
    return match[1] if match else None


def reference_data(db):
    categories = [dict(r) for r in db.execute('SELECT * FROM gdlr_categories WHERE active=1 ORDER BY name_key')]
    from smu_api import active_smu_names
    departments = active_smu_names(db)
    return categories, departments


def prepare_import(db, parsed, filters, decisions):
    if not isinstance(decisions, dict) or any(not isinstance(decisions.get(k, {}), dict) for k in ('categories', 'departments', 'identities')):
        raise ImportProblem('Некорректные ручные соответствия.')
    rows = filtered_rows(parsed, filters)
    categories, departments = reference_data(db)
    aliases = db.execute('''SELECT a.name source,c.name target FROM smu_aliases a
        JOIN smu_catalog c ON c.id=a.smu_id WHERE c.active=1''').fetchall()
    active_ids = {r['id'] for r in categories}
    exact = defaultdict(set)
    for category in categories:
        exact[normalized(category['name'])].add(category['id'])
    workers = [dict(r) for r in db.execute('''SELECT w.*,ec.category_id,ec.edit_token category_token,
        m.crew_id,o.identity_key,o.edit_token outstaff_token FROM workers w
        LEFT JOIN employee_gdlr ec ON ec.worker_id=w.id LEFT JOIN crew_members m ON m.worker_id=w.id
        LEFT JOIN outstaff_members o ON o.worker_id=w.id''')]
    identities = {r['identity_key']: r for r in workers if r['identity_key']}
    by_name, by_category, by_profession = defaultdict(list), defaultdict(set), defaultdict(set)
    for worker in workers:
        by_name[normalized(worker['full_name'])].append(worker)
        targets = {worker['category_id']} if worker['category_id'] in active_ids else exact[normalized(worker['category'])]
        if worker['active']:
            by_category[normalized(worker['category'])].update(targets)
            by_profession[normalized(worker['profession'])].update(targets)
    for old in db.execute('''SELECT o.*,ec.category_id FROM outstaff_members o
            JOIN employee_gdlr ec ON ec.worker_id=o.worker_id JOIN workers w ON w.id=o.worker_id WHERE w.active=1'''):
        if old['category_id'] in active_ids:
            by_category[normalized(old['source_category'] or old['work_kind'])].add(old['category_id'])
    counts = Counter((normalized(r['full_name']), normalized(r['vendor'])) for r in rows)
    ready, used = [], set()
    for source in rows:
        item = {k: v for k, v in source.items() if k != 'cells'}
        errors = []
        key = digest([normalized(source['full_name']), normalized(source['vendor'])])
        old = identities.get(key)
        candidates = by_name[normalized(source['full_name'])]
        identity = decisions.get('identities', {}).get(str(source['source_row']))
        if not old and candidates:
            if identity == 'new':
                pass
            elif type(identity) is int and any(w['id'] == identity and not w['identity_key'] for w in candidates):
                old = next(w for w in candidates if w['id'] == identity)
            else:
                errors.append('Выберите: существующий сотрудник или отдельный однофамилец.')
        if old and not old['active']:
            errors.append('Сотрудник отключён; нужна ручная проверка.')
        if old and old['id'] in used:
            errors.append('Один сотрудник выбран для нескольких строк.')
        if old: used.add(old['id'])
        if not source['vendor']:
            errors.append('Не заполнен контрагент.')
        if counts[(normalized(source['full_name']), normalized(source['vendor']))] > 1:
            errors.append('Повтор ФИО и контрагента в выбранных строках. Уточните фильтры или исходный файл.')
        department = decisions.get('departments', {}).get(source['department'])
        department_reason = 'Выбрано вручную'
        if department is None:
            matches = sorted({a['target'] for a in aliases if normalized(a['source']) == normalized(source['department'])
                       or (department_number(source['department']) is not None
                           and department_number(a['source']) == department_number(source['department']))})
            department = matches[0] if len(matches) == 1 else ''
            department_reason = 'Совпадение номера СМУ' if department else 'Выберите участок'
            if old and old['department'] in departments:
                department, department_reason = old['department'], 'Текущая привязка сотрудника'
        if department not in departments:
            errors.append('Не сопоставлен СМУ.')
        category_key = source['category'] or source['work_kind']
        category_id = decisions.get('categories', {}).get(category_key)
        category_reason = 'Выбрано вручную'
        if old and old['category_id'] is not None:
            category_id, category_reason = old['category_id'], 'Сохранена текущая ручная привязка'
        elif category_id is None:
            evidence = exact[normalized(category_key)] | by_category[normalized(category_key)] if category_key else set()
            category_reason = 'Точное название / существующие привязки' if source['category'] else 'По виду работ: точное название / существующие привязки'
            if not evidence and source['profession']:
                evidence = by_profession[normalized(source['profession'])]
                category_reason = 'Однозначная категория сотрудников с той же должностью'
            category_id = next(iter(evidence)) if len(evidence) == 1 else None
            if not category_id: category_reason = 'Нет однозначного соответствия'
        if type(category_id) is not int or category_id not in active_ids:
            errors.append('Не сопоставлена категория ГДЛР.')
        item.update(identity_key=key, worker_id=old['id'] if old else None,
                    existing=old, identity_candidates=[{'id': w['id'], 'name': w['full_name'], 'number': w['personnel_no'],
                                                      'vendor': w['employer']} for w in candidates if not w['identity_key']],
                    needs_identity=not identities.get(key) and bool(candidates),
                    mapped_department=department, department_reason=department_reason,
                    category_key=category_key, category_id=category_id, category_reason=category_reason, errors=errors)
        ready.append(item)
    return {'rows': ready, 'categories': categories, 'departments': departments,
            'total': len(parsed['rows']), 'selected': len(rows), 'unresolved': sum(bool(r['errors']) for r in ready)}


def register_outstaff_routes(app, get_db, roles_required, utc_now):
    signer = URLSafeTimedSerializer(app.secret_key, salt='outstaff-review-v1')

    def read_upload():
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.xlsx'):
            abort(400, description='Выберите файл XLSX перевахтовки.')
        return parse_outstaff(upload.read(), upload.filename)

    @app.post('/api/outstaff/inspect')
    @roles_required('admin')
    def inspect():
        try:
            parsed = read_upload()
            columns = [{'index': i, 'name': name, 'values': sorted({r['cells'][i] if i < len(r['cells']) else '' for r in parsed['rows']})}
                       for i, name in enumerate(parsed['headers']) if name]
            return jsonify({'columns': columns, 'sheet': parsed['sheet'], 'total': len(parsed['rows'])})
        except ImportProblem as exc:
            abort(400, description=str(exc))

    @app.post('/api/outstaff/preview')
    @app.post('/api/outstaff/apply')
    @roles_required('admin')
    def import_rows():
        try:
            parsed = read_upload()
            filters = json.loads(request.form.get('filters', '[]'))
            decisions = json.loads(request.form.get('decisions', '{}'))
            db = get_db()
            key = digest([parsed['file_hash'], filters, decisions])
            applying = request.path.endswith('/apply')
            if applying:
                require_all(db)
                try:
                    review = signer.loads(request.form.get('token', ''), max_age=1800)
                except BadSignature:
                    abort(409, description='Предпросмотр устарел. Выполните проверку заново.')
                if review.get('user') != g.user['id'] or review.get('key') != key:
                    abort(409, description='Файл или настройки изменились. Выполните проверку заново.')
                done = db.execute('SELECT id,selected_count FROM outstaff_imports WHERE request_key=?', (key,)).fetchone()
                if done:
                    return jsonify({'id': done['id'], 'selected': done['selected_count'], 'already_imported': True})
                backup_database(db)
            with db:
                if applying:
                    db.execute('BEGIN IMMEDIATE')
                    require_all(db)
                # Recheck after taking the write lock, including current bindings and catalog tokens.
                done = db.execute('SELECT id,selected_count FROM outstaff_imports WHERE request_key=?', (key,)).fetchone()
                if applying and done:
                    return jsonify({'id': done['id'], 'selected': done['selected_count'], 'already_imported': True})
                prepared = prepare_import(db, parsed, filters, decisions)
                snapshot = digest(prepared)
                if not applying:
                    prepared['token'] = signer.dumps({'key': key, 'snapshot': snapshot, 'user': g.user['id']})
                    for row in prepared['rows']: row.pop('existing')
                    return jsonify(prepared)
                if snapshot != review['snapshot']:
                    abort(409, description='Сотрудники или справочники изменились. Проверьте соответствия заново.')
                if not prepared['selected'] or prepared['unresolved']:
                    abort(400, description='Укажите соответствия для всех выбранных строк перед импортом.')
                batch = db.execute('''INSERT INTO outstaff_imports(request_key,filename,sheet,file_hash,selection_json,
                    imported_by,imported_at,selected_count) VALUES (?,?,?,?,?,?,?,?)''',
                    (key, parsed['filename'], parsed['sheet'], parsed['file_hash'], json.dumps(filters, ensure_ascii=False),
                     g.user['id'], utc_now(), prepared['selected'])).lastrowid
                sources = {r['source_row']: r for r in parsed['rows']}
                for row in prepared['rows']:
                    worker_id = row['worker_id']
                    if worker_id is None:
                        worker_id = db.execute('''INSERT INTO workers(full_name,personnel_no,contractor,employer,profession,category,department)
                            VALUES (?,?,?,?,?,?,?)''', (row['full_name'], 'OUT-' + secrets.token_hex(8).upper(), 'ЛГСС',
                            row['vendor'], row['profession'], row['category'], row['mapped_department'])).lastrowid
                        # Keep new employees assignable immediately; their crew can be changed in the directory.
                        crew_key = 'outstaff:' + digest([row['vendor'], row['mapped_department']])
                        crew = db.execute('SELECT id FROM crews WHERE import_key=?', (crew_key,)).fetchone()
                        crew_id = crew['id'] if crew else db.execute('''INSERT INTO crews(name,owner_user_id,created_at,import_key)
                            VALUES (?,?,?,?)''', ('Аутстафф · ' + row['vendor'] + ' · ' + row['mapped_department'],
                            g.user['id'], utc_now(), crew_key)).lastrowid
                        db.execute('INSERT INTO crew_members(crew_id,worker_id) VALUES (?,?)', (crew_id, worker_id))
                    # Existing workers retain their manual fields, brigades, assignments and category bindings.
                    db.execute('UPDATE workers SET department=? WHERE id=?', (row['mapped_department'], worker_id))
                    db.execute('''INSERT OR IGNORE INTO employee_gdlr(worker_id,category_id,edit_token,updated_by,updated_at)
                        VALUES (?,?,?,?,?)''', (worker_id, row['category_id'], secrets.token_hex(16), g.user['id'], utc_now()))
                    db.execute('''INSERT INTO outstaff_members(worker_id,identity_key,import_id,source_row,source_department,
                        source_category,work_kind,vendor,staff_type,arrival,departure,edit_token) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                        ON CONFLICT(worker_id) DO UPDATE SET import_id=excluded.import_id,source_row=excluded.source_row,
                        source_department=excluded.source_department,source_category=excluded.source_category,
                        work_kind=excluded.work_kind,vendor=excluded.vendor,staff_type=excluded.staff_type,
                        arrival=excluded.arrival,departure=excluded.departure,edit_token=excluded.edit_token''',
                        (worker_id, row['identity_key'], batch, row['source_row'], row['department'], row['category'],
                         row['work_kind'], row['vendor'], row['staff_type'], row['arrival'], row['departure'], secrets.token_hex(16)))
                    db.execute('''INSERT INTO outstaff_import_rows(import_id,source_row,worker_id,source_json,mapping_json)
                        VALUES (?,?,?,?,?)''', (batch, row['source_row'], worker_id,
                        json.dumps(sources[row['source_row']], ensure_ascii=False),
                        json.dumps({k: row[k] for k in ('category_id', 'category_reason', 'mapped_department', 'department_reason')}, ensure_ascii=False)))
                from contractor_api import sync_contractors
                sync_contractors(db, g.user['id'])
            return jsonify({'id': batch, 'selected': prepared['selected'], 'already_imported': False})
        except (ImportProblem, json.JSONDecodeError) as exc:
            abort(400, description=str(exc))

    @app.put('/api/outstaff/<int:worker_id>/department')
    @roles_required('admin')
    def change_department(worker_id):
        require_all(get_db())
        payload = request.get_json(silent=True)
        if not isinstance(payload, dict): abort(400, description='Выберите участок.')
        db = get_db()
        with db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT * FROM outstaff_members WHERE worker_id=?', (worker_id,)).fetchone()
            if not row: abort(404, description='Сотрудник аутстаффа не найден.')
            if row['edit_token'] != payload.get('expected_token'):
                abort(409, description='Данные изменились. Обновите список.')
            if payload.get('department') not in reference_data(db)[1]:
                abort(400, description='Выберите существующий строительно-монтажный участок.')
            db.execute('UPDATE workers SET department=? WHERE id=?', (payload['department'], worker_id))
            db.execute('UPDATE outstaff_members SET edit_token=? WHERE worker_id=?', (secrets.token_hex(16), worker_id))
        return jsonify({'updated': worker_id})
