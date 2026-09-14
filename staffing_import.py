"""Explicit attendance import. Source workbook is never modified."""
import hashlib
import io
import re
import secrets
import sqlite3
import zipfile
from collections import Counter
from contextlib import closing
from pathlib import Path

import openpyxl


class ImportProblem(ValueError):
    pass


def clean(value):
    return "" if value is None else str(value).strip()


def natural_key(value):
    return tuple((0, int(x)) if x.isdigit() else (1, x.casefold())
                 for x in re.split(r"(\d+)", value))


def parse_attendance(content, filename):
    if len(content) > 20_000_000:
        raise ImportProblem("Размер файла превышает 20 МБ.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 150_000_000:
                raise ImportProblem("Распакованный файл слишком велик.")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (ValueError, OSError, zipfile.BadZipFile, KeyError) as error:
        raise ImportProblem("Не удалось прочитать XLSX.") from error
    try:
        if "Явка" not in workbook.sheetnames:
            raise ImportProblem("В файле нет листа «Явка».")
        sheet = workbook["Явка"]
        if sheet.max_row > 50_000 or sheet.max_column > 100:
            raise ImportProblem("Лист «Явка» превышает допустимые размеры.")
        headers = [clean(v).casefold() for v in next(sheet.iter_rows(values_only=True))]
        headers = ['таб. номер' if value == 'табельный номер' else value for value in headers]
        required = {"personnel_no": "таб. номер", "full_name": "фио", "profession": "должность",
                    "category": "категория гдлр", "department": "подразделение", "qualification": "квалификация"}
        columns = {}
        for key, label in required.items():
            if headers.count(label) != 1:
                raise ImportProblem(f"На листе «Явка» нужен один столбец «{label}».")
            columns[key] = headers.index(label)
        # The supplied attendance file has an unnamed brigade column Y.
        named = [i for i, name in enumerate(headers) if name in {"бригада", "номер бригады", "№ бригады"}]
        if len(named) > 1:
            raise ImportProblem("В файле несколько столбцов номера бригады.")
        crew_column = named[0] if named else 24
        if not named and len(headers) > 24 and headers[24] == 'примечание':
            crew_column = None
        elif not named and len(headers) > 24 and headers[24]:
            raise ImportProblem("Столбец Y занят другим полем. Укажите заголовок «Номер бригады».")
        counts = Counter()
        rows, people, issues, seen = [], [], [], set()
        for source_row, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            values = {key: clean(cells[i]) for key, i in columns.items()}
            if not values["full_name"]:
                continue
            counts["source"] += 1
            cell = cells[columns["personnel_no"]]
            if isinstance(cell, (float, int)) and not isinstance(cell, bool) and float(cell).is_integer():
                values["personnel_no"] = str(int(cell))
            source_crew = clean(cells[crew_column]) if crew_column is not None and crew_column < len(cells) else ""
            people.append({**values, "source_row": source_row, "source_crew": source_crew})
            if not re.match(r"^строительно[-– ]монтажный участок(?:\s|$)", values["department"], re.I):
                counts["department_excluded"] += 1
                continue
            if values["qualification"].casefold() != "рабочие":
                counts["qualification_excluded"] += 1
                continue
            category = values["category"].casefold()
            if any(part in category for part in ("водител", "машинист", "вспомогател")):
                counts["category_excluded"] += 1
                continue
            if not category or category.startswith("#"):
                issues.append(f"Строка {source_row}: не указана категория ГДЛР.")
            cell = cells[columns["personnel_no"]]
            if isinstance(cell, (float, int)) and not isinstance(cell, bool) and float(cell).is_integer():
                values["personnel_no"] = str(int(cell))
            number = values["personnel_no"]
            if not number or number.startswith("#"):
                issues.append(f"Строка {source_row}: не указан табельный номер.")
            if number.casefold() in seen:
                issues.append(f"Строка {source_row}: повтор табельного номера {number}.")
            seen.add(number.casefold())
            crew_name = source_crew
            if crew_name in {"", "#N/A", "#Н/Д"}:
                crew_name = ""
                counts["without_crew"] += 1
            elif not re.fullmatch(r"Бригада\s*№\s*\S.*", crew_name, re.I):
                issues.append(f"Строка {source_row}: непонятный номер бригады «{crew_name}».")
            rows.append({**values, "source_row": source_row, "crew_name": crew_name,
                         "contractor": "ЛГСС", "employer": "ЛГСС" if re.fullmatch(r"[0-9]+", number) else ""})
        if not rows:
            issues.append("После фильтрации нет сотрудников для расстановки.")
        groups = Counter(row["crew_name"] for row in rows)
        return {"sha256": hashlib.sha256(content).hexdigest(), "filename": Path(filename).name,
                "sheet": "Явка", "rows": rows, "people": people, "issues": issues,
                "summary": {**dict(counts), "selected": len(rows), "crews": len([x for x in groups if x]),
                            "numeric_personnel": sum(bool(re.fullmatch(r"[0-9]+", r["personnel_no"])) for r in rows)},
                "groups": [{"name": name or "Бригада не указана", "count": groups[name]}
                           for name in sorted(groups, key=natural_key)]}
    finally:
        workbook.close()


def migrate_staffing(db):
    additions = {
        "workers": {"pps": "TEXT NOT NULL DEFAULT ''", "category": "TEXT NOT NULL DEFAULT ''", "department": "TEXT NOT NULL DEFAULT ''",
                    "gsp_profession": "TEXT NOT NULL DEFAULT ''"},
        "crews": {"import_key": "TEXT", "linear_itr": "TEXT NOT NULL DEFAULT ''",
                  "brigadier": "TEXT NOT NULL DEFAULT ''", "details_token": "TEXT NOT NULL DEFAULT ''"},
    }
    for table, fields in additions.items():
        existing = {row["name"] for row in db.execute(f"PRAGMA table_info({table})")}
        for name, definition in fields.items():
            if name not in existing:
                db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")
    db.executescript("""
        CREATE UNIQUE INDEX IF NOT EXISTS idx_crew_import_key ON crews(import_key) WHERE import_key IS NOT NULL;
        CREATE TABLE IF NOT EXISTS staffing_imports (
            id INTEGER PRIMARY KEY, sha256 TEXT NOT NULL UNIQUE, filename TEXT NOT NULL,
            imported_by INTEGER NOT NULL REFERENCES users(id), imported_at TEXT NOT NULL,
            selected_count INTEGER NOT NULL, summary_json TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS staffing_import_members (
            import_id INTEGER NOT NULL REFERENCES staffing_imports(id),
            worker_id INTEGER NOT NULL REFERENCES workers(id), source_row INTEGER NOT NULL,
            source_crew TEXT NOT NULL, PRIMARY KEY(import_id, worker_id)
        );
        CREATE TABLE IF NOT EXISTS staffing_row_details (
            worker_id INTEGER PRIMARY KEY REFERENCES workers(id),
            brigadier_override TEXT, linear_itr_override TEXT, edit_token TEXT NOT NULL,
            updated_by INTEGER NOT NULL REFERENCES users(id), updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS staffing_people (
            id INTEGER PRIMARY KEY, import_id INTEGER NOT NULL REFERENCES staffing_imports(id),
            source_row INTEGER NOT NULL, full_name TEXT NOT NULL, personnel_no TEXT NOT NULL,
            profession TEXT NOT NULL, department TEXT NOT NULL, qualification TEXT NOT NULL,
            source_crew TEXT NOT NULL, UNIQUE(import_id, source_row)
        );
        -- Manual entries retain a legacy import anchor and a negative source_row.
        -- Their provenance and lifetime are independent of the anchored workbook.
        CREATE TABLE IF NOT EXISTS manual_staffing_people (
            person_id INTEGER PRIMARY KEY REFERENCES staffing_people(id),
            created_by INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS staffing_shifts (
            work_date TEXT NOT NULL, worker_id INTEGER NOT NULL REFERENCES workers(id),
            shift TEXT NOT NULL CHECK(shift IN ('1 смена','2 смена')),
            edit_token TEXT NOT NULL, updated_by INTEGER NOT NULL REFERENCES users(id), updated_at TEXT NOT NULL,
            PRIMARY KEY(work_date,worker_id)
        );
    """)
    if 'source_label' not in {r['name'] for r in db.execute('PRAGMA table_info(staffing_imports)')}:
        db.execute("ALTER TABLE staffing_imports ADD COLUMN source_label TEXT NOT NULL DEFAULT ''")
    existing = {row["name"] for row in db.execute("PRAGMA table_info(staffing_row_details)")}
    if "linear_itr_override" not in existing:
        db.execute("ALTER TABLE staffing_row_details ADD COLUMN linear_itr_override TEXT")
    for table in ("crews", "staffing_row_details"):
        columns = {row["name"] for row in db.execute(f"PRAGMA table_info({table})")}
        for field in ("linear_itr", "brigadier"):
            if field + "_person_id" not in columns:
                db.execute(f"ALTER TABLE {table} ADD COLUMN {field}_person_id INTEGER REFERENCES staffing_people(id)")


def import_conflicts(db, parsed):
    issues = list(parsed["issues"])
    removed = {r[0] for r in db.execute('SELECT worker_id FROM employee_removals')}
    existing = {r["personnel_no"].casefold(): dict(r) for r in db.execute("""
        SELECT w.*, c.import_key, c.name crew_name FROM workers w
        LEFT JOIN crew_members m ON m.worker_id=w.id LEFT JOIN crews c ON c.id=m.crew_id
    """)}
    for row in parsed["rows"]:
        old = existing.get(row["personnel_no"].casefold())
        if not old:
            continue
        if old["full_name"].strip().casefold() != row["full_name"].casefold():
            issues.append(f"Строка {row['source_row']}: табельный номер {row['personnel_no']} уже принадлежит другому ФИО.")
        if old['id'] in removed:
            continue
        if not old["active"]:
            issues.append(f"Строка {row['source_row']}: сотрудник отключён. Нужна ручная проверка.")
        label = parsed.get('source_label', '')
        if old['pps'] and old['pps'] != label:
            issues.append(f"Строка {row['source_row']}: сотрудник уже отмечен как {old['pps']}.")
        target = import_crew_key(row, label)
        if old["crew_name"] and old["import_key"] != target:
            issues.append(f"Строка {row['source_row']}: сотрудник уже состоит в другой бригаде «{old['crew_name']}».")
    return issues


def backup_database(db):
    from backup_api import create_backup
    return create_backup(db, reason='automatic')


def removal_import_summary(db, parsed):
    removed = {r[0].casefold() for r in db.execute(
        'SELECT w.personnel_no FROM employee_removals er JOIN workers w ON w.id=er.worker_id')}
    skipped = sum(row['personnel_no'].casefold() in removed for row in parsed['rows'])
    return {**parsed['summary'], 'selected': len(parsed['rows']) - skipped, 'skipped_removed': skipped}


def store_people(db, batch_id, parsed):
    db.executemany("""INSERT OR IGNORE INTO staffing_people
        (import_id,source_row,full_name,personnel_no,profession,department,qualification,source_crew)
        VALUES (?,?,?,?,?,?,?,?)""", [(batch_id, p['source_row'], p['full_name'], p['personnel_no'],
            p['profession'], p['department'], p['qualification'], p['source_crew']) for p in parsed['people']])


def active_import_ids_sql():
    return 'SELECT MAX(id) FROM staffing_imports GROUP BY source_label'


def active_members_sql():
    return f'''SELECT sm.worker_id,sm.source_row,sm.source_crew FROM staffing_import_members sm
        WHERE sm.import_id IN ({active_import_ids_sql()}) AND sm.import_id=(
            SELECT MAX(other.import_id) FROM staffing_import_members other WHERE other.worker_id=sm.worker_id
            AND other.import_id IN ({active_import_ids_sql()}))'''


def import_crew_key(row, label):
    key = row['crew_name'] or 'attendance:missing-crew'
    return label + ':' + key if label else key


def apply_attendance(db, parsed, user_id, utc_now, authorize=None):
    import json
    label = parsed.get('source_label', '')
    if label not in ('', 'ППС15', 'ППС19'):
        raise ImportProblem('Выберите ППС15 или ППС19.')
    existing = db.execute("SELECT id,source_label FROM staffing_imports WHERE sha256=?", (parsed["sha256"],)).fetchone()
    if existing:
        if existing['source_label'] != label:
            raise ImportProblem('Этот файл уже загружен с другой отметкой ППС.')
        count = db.execute("""SELECT COUNT(*) FROM staffing_people p WHERE import_id=?
            AND NOT EXISTS (SELECT 1 FROM manual_staffing_people m WHERE m.person_id=p.id)""",
            (existing['id'],)).fetchone()[0]
        result = {"id": existing["id"], "already_imported": True, "selected": parsed["summary"]["selected"],
                  "people_count": len(parsed['people']), "people_added": len(parsed['people']) - count}
        if count < len(parsed['people']):
            result['backup'] = backup_database(db).name
            with db:
                db.execute('BEGIN IMMEDIATE')
                if authorize:
                    authorize(db)
                store_people(db, existing['id'], parsed)
        return result
    issues = import_conflicts(db, parsed)
    if issues:
        raise ImportProblem("\n".join(issues[:20]))
    backup = backup_database(db)
    with db:
        db.execute("BEGIN IMMEDIATE")
        if authorize:
            authorize(db)
        # A second process may have completed this import while the backup ran.
        existing = db.execute("SELECT id,source_label FROM staffing_imports WHERE sha256=?", (parsed["sha256"],)).fetchone()
        if existing:
            if existing['source_label'] != label:
                raise ImportProblem('Этот файл уже загружен с другой отметкой ППС.')
            store_people(db, existing['id'], parsed)
            return {"id": existing["id"], "already_imported": True, "selected": parsed["summary"]["selected"]}
        issues = import_conflicts(db, parsed)
        if issues:
            raise ImportProblem("\n".join(issues[:20]))
        summary = removal_import_summary(db, parsed)
        batch = db.execute("""INSERT INTO staffing_imports(sha256,filename,imported_by,imported_at,selected_count,summary_json,source_label)
                            VALUES (?,?,?,?,?,?,?)""", (parsed["sha256"], parsed["filename"], user_id, utc_now(),
                            summary['selected'], json.dumps(summary, ensure_ascii=False), label)).lastrowid
        store_people(db, batch, parsed)
        crews = {r["import_key"]: r["id"] for r in db.execute("SELECT id,import_key FROM crews WHERE import_key IS NOT NULL")}
        removed = {r[0].casefold() for r in db.execute('SELECT w.personnel_no FROM employee_removals er JOIN workers w ON w.id=er.worker_id')}
        for row in parsed["rows"]:
            if row['personnel_no'].casefold() in removed:
                continue
            key = import_crew_key(row, label)
            if key not in crews:
                crews[key] = db.execute("INSERT INTO crews(name,owner_user_id,created_at,import_key) VALUES (?,?,?,?)",
                                       ((label + ' · ' if label else '') + (row["crew_name"] or "Бригада не указана"), user_id, utc_now(), key)).lastrowid
            db.execute("""INSERT INTO workers(full_name,personnel_no,contractor,employer,profession,category,department)
                VALUES (:full_name,:personnel_no,:contractor,:employer,:profession,:category,:department)
                ON CONFLICT(personnel_no) DO UPDATE SET contractor=excluded.contractor,employer=excluded.employer,
                profession=excluded.profession,category=excluded.category,department=excluded.department""", row)
            worker = db.execute("SELECT id FROM workers WHERE personnel_no=? COLLATE NOCASE", (row["personnel_no"],)).fetchone()["id"]
            db.execute('UPDATE workers SET pps=? WHERE id=?', (label, worker))
            db.execute("INSERT OR IGNORE INTO crew_members(crew_id,worker_id) VALUES (?,?)", (crews[key], worker))
            db.execute("INSERT INTO staffing_import_members(import_id,worker_id,source_row,source_crew) VALUES (?,?,?,?)",
                       (batch, worker, row["source_row"], row["crew_name"]))
        from contractor_api import sync_contractors
        sync_contractors(db, user_id)
    return {"id": batch, "selected": summary['selected'], 'skipped_removed': summary['skipped_removed'], "backup": backup.name, "already_imported": False}
