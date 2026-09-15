"""Native OOXML PivotTable backed by the exported Excel table."""
from collections import Counter
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.fields import Index, Missing, Number, Text
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import DataField, FieldItem, Location, PivotField, PivotTableStyle, RowColField, RowColItem, TableDefinition
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.escape import escape
from openpyxl.worksheet.table import Table, TableStyleInfo


ROW_FIELDS = (1, 2)
COL_FIELDS = (3, 4, 12)
NAME_FIELD = 5
SOURCE_NAME = 'StaffingSource'
GDLR_FIELD = 9
SHARED_FIELDS = (*ROW_FIELDS, *COL_FIELDS, GDLR_FIELD)
SUMMARY_OFFSET = 8  # Space above the report for the visible category slicer.


def _value(value):
    return None if value in (None, '') else value


def _label(value):
    return '(пусто)' if value is None else str(value)


def _axis_paths(records, fields, top_subtotals=False):
    leaves = {tuple(record[field] for field in fields) for record in records}
    paths = []

    def visit(prefix):
        children = sorted({leaf[len(prefix)] for leaf in leaves if leaf[:len(prefix)] == prefix},
                          key=lambda value: _label(value).casefold())
        for value in children:
            child = (*prefix, value)
            leaf = len(child) == len(fields)
            if leaf or top_subtotals:
                paths.append(child)
            if not leaf:
                visit(child)
                if not top_subtotals:
                    paths.append(child)
    visit(())
    return [*paths, ()]


def _snapshot(workbook, records, day, category, query, row_paths, col_paths):
    """Save visible values as well as the pivot cache for readers that cannot refresh."""
    sheet = workbook.create_sheet('Сводная таблица')
    for row, pair in enumerate((('Дата расстановки', '.'.join(reversed(day.split('-')))),
                               ('Выборка по расстановке', category if category else ('Без категории' if category == '' else 'Все категории ГДЛР')),
                               ('Поиск по позициям', query or 'Все позиции'),
                               ('Источник: лист «Список сотрудников»', f'Строк в списке: {len(records)}. После изменения списка: Данные → Обновить всё')), 1 + SUMMARY_OFFSET):
        for col, value in enumerate(pair, 1):
            cell = sheet.cell(row, col, value)
            cell.data_type = 's'
    counts = Counter()
    for record in records:
        if record[NAME_FIELD] is None:
            continue
        row_key = tuple(record[field] for field in ROW_FIELDS)
        col_key = tuple(record[field] for field in COL_FIELDS)
        for r in range(len(ROW_FIELDS) + 1):
            for c in range(len(COL_FIELDS) + 1):
                counts[(row_key[:r], col_key[:c])] += 1
    first = 5 + SUMMARY_OFFSET
    depth = len(COL_FIELDS)
    for col, path in enumerate(col_paths, 2):
        for level in range(depth):
            if not path:
                value = 'Общий итог' if level == 0 else ''
            elif level < len(path):
                value = _label(path[level])
            elif level == len(path):
                value = 'ИТОГО'
            else:
                value = ''
            cell = sheet.cell(first + 1 + level, col, value)
            cell.data_type = 's'
    sheet.cell(first, 1, 'Количество сотрудников')
    sheet.cell(first + depth, 1, 'Позиция')
    for row, path in enumerate(row_paths, first + depth + 1):
        label = 'Общий итог' if not path else ('Не расставлены' if path == (None,) else _label(path[-1]))
        cell = sheet.cell(row, 1, label)
        cell.data_type = 's'
        cell.alignment = Alignment(wrap_text=False, indent=1 if len(path) == len(ROW_FIELDS) else 0)
        for col, column_path in enumerate(col_paths, 2):
            sheet.cell(row, col, counts[(path, column_path)])
    last_col = len(col_paths) + 1
    line = Side(style='thin', color='B8C9CC')
    for row in sheet.iter_rows(min_row=first, max_col=last_col):
        for cell in row:
            subtotal = cell.row > first + depth and len(row_paths[cell.row - first - depth - 1]) < len(ROW_FIELDS)
            header = cell.row <= first + depth
            cell.font = Font(name='Arial', size=10, bold=header or subtotal)
            cell.border = Border(left=line, right=line, top=line, bottom=line)
            if header or subtotal:
                cell.fill = PatternFill('solid', fgColor='DAE3F3' if header else 'E8EEF8')
            if cell.column > 1:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.column_dimensions['A'].width = 80
    for col in range(2, last_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 10
    sheet.sheet_format.defaultRowHeight = 15
    sheet.freeze_panes = f'B{9 + SUMMARY_OFFSET}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f'1:{8 + SUMMARY_OFFSET}'
    sheet.print_area = f'A1:{get_column_letter(last_col)}{sheet.max_row}'
    return sheet


def summary_sheet(workbook, rows, day, category=None, query=''):
    source = workbook['Список сотрудников']
    headers = [str(cell.value) for cell in source[1]]
    records = [tuple(_value(value) for value in row) for row in source.iter_rows(min_row=2, values_only=True)]
    table = Table(displayName=SOURCE_NAME, ref=source.auto_filter.ref)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    source.add_table(table)
    # Excel does not allow a worksheet AutoFilter overlapping a table AutoFilter.
    source.auto_filter.ref = None
    indexes, cache_fields = [], []
    def cache_value(value):
        return Missing() if value is None else Number(v=value) if isinstance(value, (int, float)) else Text(v=escape(str(value)))

    for field, header in enumerate(headers):
        values = list(dict.fromkeys(record[field] for record in records))
        indexes.append({value: index for index, value in enumerate(values)})
        items = [cache_value(value) for value in values] if field in SHARED_FIELDS else []
        strings = any(isinstance(value, str) for value in values)
        numbers = any(isinstance(value, (int, float)) for value in values)
        cache_fields.append(CacheField(name=escape(header), numFmtId=49 if field == 6 else 0, sharedItems=SharedItems(
            _fields=items, containsString=None if strings else False, containsNumber=True if numbers else None,
            containsInteger=True if numbers else None, containsBlank=True if None in values else None,
            containsNonDate=False if not strings and not numbers else None,
            containsSemiMixedTypes=False if numbers and not strings else None,
            minValue=min(value for value in values if value is not None) if numbers else None,
            maxValue=max(value for value in values if value is not None) if numbers else None)))
    cache = CacheDefinition(cacheSource=CacheSource(type='worksheet', worksheetSource=WorksheetSource(name=SOURCE_NAME)),
                            cacheFields=cache_fields, recordCount=len(records), saveData=True,
                            enableRefresh=True, refreshOnLoad=True, createdVersion=6, refreshedVersion=6,
                            minRefreshableVersion=3, missingItemsLimit=0)
    cache.records = RecordList(r=[Record(_fields=[Index(v=indexes[field][value]) if field in SHARED_FIELDS
                                                else cache_value(value) for field, value in enumerate(record)]) for record in records])
    row_paths = _axis_paths(records, ROW_FIELDS, top_subtotals=True)
    col_paths = _axis_paths(records, COL_FIELDS)
    sheet = _snapshot(workbook, records, day, category, query, row_paths, col_paths)
    fields = []
    for field in range(len(headers)):
        axis = 'axisRow' if field in ROW_FIELDS else 'axisCol' if field in COL_FIELDS else None
        subtotal = field in (*ROW_FIELDS[:-1], *COL_FIELDS[:-1])
        items = [FieldItem(x=index, n='Не расставлены' if field == ROW_FIELDS[0] and value is None else None)
                 for value, index in indexes[field].items()] if axis or field == GDLR_FIELD else []
        if subtotal:
            items.append(FieldItem(t='default'))
        fields.append(PivotField(axis=axis, dataField=field == NAME_FIELD, items=items,
                                 defaultSubtotal=subtotal, showAll=False, sortType='ascending'))

    def axis_items(paths, axes):
        result, previous = [], ()
        for path in paths:
            repeated = 0
            while repeated < min(len(previous), len(path) - 1) and previous[repeated] == path[repeated]:
                repeated += 1
            kind = 'grand' if not path else 'data' if axes == ROW_FIELDS or len(path) == len(axes) else 'default'
            result.append(RowColItem(t=kind, r=repeated, x=[Index(v=indexes[axes[level]][path[level]])
                                                         for level in range(repeated, len(path))] or [Index(v=0)]))
            previous = path
        return result

    pivot = TableDefinition(name='StaffingPivot', cacheId=1, dataCaption='Количество сотрудников',
                            grandTotalCaption='Общий итог', rowHeaderCaption='Позиция',
                            location=Location(ref=f'A{5 + SUMMARY_OFFSET}:{get_column_letter(sheet.max_column)}{sheet.max_row}',
                                              firstHeaderRow=1, firstDataRow=4, firstDataCol=1),
                            pivotFields=fields, rowFields=[RowColField(x=index) for index in ROW_FIELDS],
                            colFields=[RowColField(x=index) for index in COL_FIELDS],
                            rowItems=axis_items(row_paths, ROW_FIELDS), colItems=axis_items(col_paths, COL_FIELDS),
                            dataFields=[DataField(name='Количество сотрудников', fld=NAME_FIELD, subtotal='count', baseField=0, baseItem=0)],
                            createdVersion=6, updatedVersion=6, minRefreshableVersion=3,
                            compact=True, compactData=True, outline=True, outlineData=True,
                            disableFieldList=False, enableDrill=True, useAutoFormatting=False, preserveFormatting=True,
                            pivotTableStyleInfo=PivotTableStyle(name='PivotStyleMedium9', showRowHeaders=True,
                                                               showColHeaders=True, showRowStripes=False,
                                                               showColStripes=False, showLastColumn=True))
    pivot.cache = cache
    sheet.add_pivot(pivot)
    workbook.active = sheet
    return len(records)


def save_workbook_with_slicer(workbook, output):
    """Add Excel's x14 slicer parts after openpyxl saves the two-sheet report.

    openpyxl cannot serialize slicers. Keep this final packaging step after all
    workbook edits; loading and saving the result with openpyxl drops the slicer.
    The cache and drawing follow the OOXML produced by Excel for a non-OLAP
    PivotTable slicer, including the separate x14 pivotCacheId relationship.
    """
    if workbook.sheetnames != ['Список сотрудников', 'Сводная таблица']:
        raise ValueError('The staffing slicer requires the two-sheet report')
    raw = BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as source:
        parts = {info.filename: source.read(info) for info in source.infolist()}

    main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    x14 = 'http://schemas.microsoft.com/office/spreadsheetml/2009/9/main'
    rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    package = 'http://schemas.openxmlformats.org/package/2006/relationships'
    content = 'http://schemas.openxmlformats.org/package/2006/content-types'

    def node(parent, ns, tag, **attrs):
        return ET.SubElement(parent, f'{{{ns}}}{tag}', attrs)

    def write(path, root):
        parts[path] = ET.tostring(root, encoding='utf-8', xml_declaration=True)

    def relationship(path, kind, target):
        root = ET.fromstring(parts[path])
        used = {child.get('Id') for child in root}
        number = 1
        while f'rId{number}' in used:
            number += 1
        rid = f'rId{number}'
        node(root, package, 'Relationship', Id=rid, Type=kind, Target=target)
        write(path, root)
        return rid

    def extension(root, uri):
        extensions = root.find(f'{{{main}}}extLst')
        if extensions is None:
            extensions = node(root, main, 'extLst')
        return node(extensions, main, 'ext', uri=uri)

    cache_path = 'xl/pivotCache/pivotCacheDefinition1.xml'
    cache = ET.fromstring(parts[cache_path])
    field = cache.find(f'{{{main}}}cacheFields')[GDLR_FIELD]
    shared_items = field.find(f'{{{main}}}sharedItems')
    category_name = field.get('name')
    cache_id = '1'
    node(extension(cache, '{725AE2AE-9491-48be-B2B4-4EB974FC3084}'),
         x14, 'pivotCacheDefinition', pivotCacheId=cache_id)
    write(cache_path, cache)

    slicer_cache = ET.Element(f'{{{x14}}}slicerCacheDefinition',
                              name='Slicer_GDLR', sourceName=category_name)
    pivot_tables = node(slicer_cache, x14, 'pivotTables')
    node(pivot_tables, x14, 'pivotTable', tabId='2', name='StaffingPivot')
    data = node(slicer_cache, x14, 'data')
    tabular = node(data, x14, 'tabular', pivotCacheId=cache_id)
    items = node(tabular, x14, 'items', count=str(len(shared_items)))
    for index in sorted(range(len(shared_items)), key=lambda i: (
            shared_items[i].tag == f'{{{main}}}m', shared_items[i].get('v', '').casefold())):
        node(items, x14, 'i', x=str(index), s='1')
    write('xl/slicerCaches/slicerCache1.xml', slicer_cache)

    slicers = ET.Element(f'{{{x14}}}slicers')
    node(slicers, x14, 'slicer', name='GDLR', cache='Slicer_GDLR', caption='Категория ГДЛР',
         columnCount='3', rowHeight='228600')
    write('xl/slicers/slicer1.xml', slicers)

    # Fixed top panel (rows 1-7); report metadata begins at row 9. It stays
    # visible above the frozen report and never covers source cells or totals.
    parts['xl/drawings/drawing1.xml'] = b'''<?xml version="1.0" encoding="UTF-8"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <xdr:twoCellAnchor editAs="oneCell">
  <xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
  <xdr:to><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>7</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
  <mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
   <mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">
    <xdr:graphicFrame macro="">
     <xdr:nvGraphicFramePr><xdr:cNvPr id="2" name="GDLR"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>
     <xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>
     <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">
      <sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" name="GDLR"/>
     </a:graphicData></a:graphic>
    </xdr:graphicFrame>
   </mc:Choice>
  </mc:AlternateContent>
  <xdr:clientData/>
 </xdr:twoCellAnchor>
</xdr:wsDr>'''

    book = ET.fromstring(parts['xl/workbook.xml'])
    names = book.find(f'{{{main}}}definedNames')
    # Print areas/titles already provide definedNames in this report.
    node(names, main, 'definedName', name='Slicer_GDLR').text = '#N/A'
    cache_rid = relationship('xl/_rels/workbook.xml.rels',
                             'http://schemas.microsoft.com/office/2007/relationships/slicerCache',
                             'slicerCaches/slicerCache1.xml')
    caches = node(extension(book, '{BBE1A952-AA13-448e-AADC-164F8A28A991}'), x14, 'slicerCaches')
    node(caches, x14, 'slicerCache', **{f'{{{rel}}}id': cache_rid})
    write('xl/workbook.xml', book)

    sheet_path = 'xl/worksheets/sheet2.xml'
    sheet = ET.fromstring(parts[sheet_path])
    sheet_rels = 'xl/worksheets/_rels/sheet2.xml.rels'
    drawing_rid = relationship(sheet_rels, rel + '/drawing', '../drawings/drawing1.xml')
    # The summary has no tables, legacy drawings or controls after pageSetup.
    drawing = ET.Element(f'{{{main}}}drawing', {f'{{{rel}}}id': drawing_rid})
    sheet.append(drawing)
    slicer_rid = relationship(sheet_rels, 'http://schemas.microsoft.com/office/2007/relationships/slicer',
                              '../slicers/slicer1.xml')
    slicer_list = node(extension(sheet, '{A8765BA9-456A-4dab-B4F3-ACF838C121DE}'), x14, 'slicerList')
    node(slicer_list, x14, 'slicer', **{f'{{{rel}}}id': slicer_rid})
    write(sheet_path, sheet)

    types = ET.fromstring(parts['[Content_Types].xml'])
    for name, mime in (('drawings/drawing1.xml', 'application/vnd.openxmlformats-officedocument.drawing+xml'),
                       ('slicers/slicer1.xml', 'application/vnd.ms-excel.slicer+xml'),
                       ('slicerCaches/slicerCache1.xml', 'application/vnd.ms-excel.slicerCache+xml')):
        node(types, content, 'Override', PartName='/xl/' + name, ContentType=mime)
    write('[Content_Types].xml', types)
    with ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for name, payload in parts.items():
            target.writestr(name, payload)
