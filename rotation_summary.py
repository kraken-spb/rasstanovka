"""Read-only rotation summary over dated stages, current catalogs and SMG plans."""
from collections import defaultdict
from datetime import date, timedelta
from io import BytesIO
import json

from flask import abort, jsonify, request, send_file
from werkzeug.datastructures import MultiDict

from workforce_core import READERS, actor_scope, date_value, plain


STAGES = {'stage.onsite': 'onsite', 'stage.leave': 'leave', 'stage.inbound': 'inbound',
          'stage.pvp': 'pvp', 'stage.outbound': 'outbound'}
FACTS = [('total', 'Всего в составе'), ('onsite', 'Явка'), ('inbound', 'Заезд'),
         ('pvp', 'ПВП'), ('leave', 'Неявка'), ('outbound', 'Выезд'), ('unknown', 'Статус не задан'),
         ('outstaff', 'В т. ч. аутстафф'), ('arrived', 'Заезд факт'), ('departed', 'Выезд факт')]
FUTURE = [('arrivals', 'Заезд план'), ('departures', 'Выезд план'), ('forecast', 'Прогноз явки'),
          ('plan', 'План СМГ'), ('delta', 'Отклонение')]
MONTHS = ('Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
          'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь')
DRILL_METRICS = {k for k, _ in FACTS} | {k + str(i) for k in ('arrivals', 'departures', 'forecast') for i in (1, 2)}
NOTES = [
    'Состав — действующие сотрудники. Категория ГДЛР, СМУ и занятость берутся из текущих привязок базы; исторические изменения этих реквизитов не восстанавливаются.',
    'Явка и другие этапы — последнее подтверждённое состояние на отчётную дату с учётом исправлений. Билет не подтверждает явку. Аутстафф входит в общую численность.',
    'Заезд/выезд факт — уникальные сотрудники с состоявшейся поездкой с начала месяца по отчётную дату. Заезд учитывается только на участок.',
    'План движения — уникальные сотрудники с действующими поездками после отчётной даты. Отменённые и заменённые поездки исключены. Даты из карточки и графика без оформленной поездки не добавляются автоматически.',
    'Прогноз — последовательное применение поездок к явке на дату, без понижающего коэффициента. Встречные поездки в один день и неизвестный исходный этап делают прогноз неопределённым (—).',
    'План СМГ — последняя загруженная редакция на месяц окончания периода. Суммируется только при наличии всех значений. При выборе отдельных СМУ или ограниченном доступе общий план ППС не распределяется.',
    'ППС определяется только связью СМУ со справочником ППС. «Без ППС» означает отсутствие этой связи. Аренда, увольнения и высвобождение отдельно не типизированы: —, а не ноль.',
]


def month_after(day):
    return date(day.year + (day.month == 12), day.month % 12 + 1, 1)


def report_columns(day, end1, end2):
    """Calendar headers like the source summary; custom ranges stay explicit."""
    def display(value):
        return date.fromisoformat(value).strftime('%d.%m.%Y')

    def month_title(value):
        stamp = date.fromisoformat(value)
        return f'{MONTHS[stamp.month - 1]} {stamp.year}'

    columns = [dict(key=k, label=v, group='На ' + display(day), period=display(day))
               for k, v in FACTS if k not in ('arrived', 'departed')]
    lower = day
    for index, end in enumerate((end1, end2), 1):
        start = str(date.fromisoformat(lower) + timedelta(days=1))
        month_end = str(month_after(date.fromisoformat(end)) - timedelta(days=1))
        calendar = end == month_end and (lower[:7] == end[:7] if index == 1 else start == end[:8] + '01')
        title = month_title(end) if calendar else f'{display(start)} — {display(end)}'
        period = f'{display(start)} — {display(end)}' if start <= end else 'После отчётной даты в этом периоде дней нет'
        if index == 1:
            columns.extend(dict(key=k, label=v, group=title if calendar else month_title(day) + ' · факт',
                                period=f'{display(day[:8] + "01")} — {display(day)}')
                           for k, v in FACTS if k in ('arrived', 'departed'))
        for key, label in FUTURE:
            detail = 'На ' + display(end) if key in ('forecast', 'delta') else month_title(end) if key == 'plan' else period
            columns.append(dict(key=key + str(index), label=label, group=title, period=detail))
        lower = end
    third = str(month_after(date.fromisoformat(end2)))
    columns.append(dict(key='plan3', label='План СМГ', group=month_title(third), period=month_title(third)))
    columns.extend(dict(key=k, label=v, group='Отдельный учёт', period='На ' + display(day))
                   for k, v in [('rental', 'Аренда'), ('dismissed', 'Увольнения'), ('released', 'Высвобождение')])
    return columns


def arguments(query):
    day = date.fromisoformat(date_value(query.get('date'), required=True))
    end1 = date.fromisoformat(date_value(query.get('end1') or str(month_after(day) - timedelta(days=1)), required=True))
    end2 = date.fromisoformat(date_value(query.get('end2') or str(month_after(month_after(end1)) - timedelta(days=1)), required=True))
    if not day <= end1 < end2 or (end2 - day).days > 366:
        abort(400, description='Окончание первого периода должно быть не раньше отчётной даты, второго — позже первого. Горизонт: до 366 дней.')
    filters = {}
    for key in ('pps', 'smu'):
        values = list(dict.fromkeys(query.getlist(key)))
        if len(values) > 100 or any(v != 'none' and (len(v) > 19 or not v.isascii() or not v.isdecimal() or not 0 < int(v) < 2**63) for v in values):
            abort(400, description='Выберите до 100 значений ППС / СМУ из списка.')
        filters[key] = values
    return str(day), str(end1), str(end2), filters


def facts_for(person, movements, day, ends):
    values = {k: 0 for k, _ in FACTS}
    values['total'] = 1
    values[STAGES.get(person['stage_code'], 'unknown')] = 1
    values['outstaff'] = int(person['employment_code'] in ('employment.internal', 'employment.external'))
    live = [m for m in movements if not m.get('replaced') and m.get('result_code') != 'result.cancelled']
    for direction, key in [('arrival', 'arrived'), ('departure', 'departed')]:
        values[key] = int(any(m['direction'] == direction and m.get('result_code') == 'result.happened'
                             and m.get('actual_date') and day[:8] + '01' <= m['actual_date'] <= day
                             and (direction == 'departure' or m.get('destination_kind') == 'site') for m in live))
    pending = [m for m in live if m.get('result_code') not in ('result.happened', 'result.postponed')
               and not m.get('actual_date') and m.get('planned_date') and day < m['planned_date'] <= ends[-1]
               and (m['direction'] == 'departure' or m.get('destination_kind') == 'site')]
    events = defaultdict(set)
    for m in pending:
        events[m['planned_date']].add(m['direction'])
    state = int(person['stage_code'] == 'stage.onsite') if person['stage_code'] in STAGES else None
    lower = day
    for index, end in enumerate(ends, 1):
        period = [m for m in pending if lower < m['planned_date'] <= end]
        values[f'arrivals{index}'] = int(any(m['direction'] == 'arrival' for m in period))
        values[f'departures{index}'] = int(any(m['direction'] == 'departure' for m in period))
        for when in sorted(k for k in events if lower < k <= end):
            directions = events[when]
            state = None if len(directions) > 1 else int('arrival' in directions)
        values[f'forecast{index}'] = state
        lower = end
    return values


def calculate(people, movements, nodes, day, ends, plans, drill=None):
    """Identity sets prevent duplicate counts through hierarchy ancestors."""
    by_id = {n['id']: n for n in nodes}
    category_nodes = {n['category_id']: n['id'] for n in nodes if n['category_id'] is not None}
    root = next((n['id'] for n in nodes if not n['parent_id'] and n['is_group']), '__total__')
    if root == '__total__':
        by_id[root] = dict(id=root, parent_id=None, label='Итого', is_group=True, category_id=None)
    buckets = defaultdict(set)
    worker_values = {}
    for p in people:
        worker_values[p['id']] = facts_for(p, movements.get(p['id'], []), day, ends)
        node = category_nodes.get(p['category_id'])
        if not node:
            node = 'extra:' + str(p['category_id'])
            by_id.setdefault(node, dict(id=node, parent_id=root, label=p['category'] or 'Без категории ГДЛР', is_group=False, category_id=p['category_id']))
        visited = set()
        while node and node not in visited:
            visited.add(node)
            buckets[node].add(p['id'])
            node = by_id[node]['parent_id']
        buckets[root].add(p['id'])
    children = defaultdict(list)
    for n in by_id.values():
        if n['id'] != root:
            children[n['parent_id'] or root].append(n['id'])
    order = []
    def walk(node, depth, visited):
        if node in visited:
            return
        visited.add(node); order.append((node, depth))
        for child in children[node]:
            walk(child, depth + 1, visited)
    walk(root, 0, set())
    rows = []
    keys = [k for k, _ in FACTS] + ['arrivals1', 'departures1', 'forecast1', 'arrivals2', 'departures2', 'forecast2']
    for node, depth in order:
        n = by_id[node]
        values = {}
        for key in keys:
            numbers = [worker_values[pid][key] for pid in buckets[node]]
            values[key] = None if any(v is None for v in numbers) else sum(numbers)
        for index, plan in enumerate(plans, 1):
            values[f'plan{index}'] = plan.get(node)
            if index < 3:
                forecast, required = values[f'forecast{index}'], values[f'plan{index}']
                values[f'delta{index}'] = forecast - required if forecast is not None and required is not None else None
        values.update(rental=None, dismissed=None, released=None)
        rows.append(dict(id=node, label=n.get('category_name') or n['label'], group=n['is_group'], depth=depth, values=values))
        if drill is not None and str(node) == drill['node']:
            if values[drill['metric']] is None:
                abort(400, description='Для этой ячейки состав сотрудников не определён. Обновите свод.')
            drill['ids'] = sorted(pid for pid in buckets[node] if worker_values[pid][drill['metric']] == 1)
    return rows


def build_report(db, day, end1, end2, filters, drill=None):
    _, scope, args = actor_scope(db)
    # Single permission-scoped population; filters intersect rather than broadening rights.
    people = plain(db.native(f'''SELECT w.id,eg.category_id,gc.name category,p.employment_code,
        sc.id smu_id,sc.name smu,pc.id pps_id,pc.name pps,st.stage_code
        FROM workers w JOIN workforce_profiles p ON p.worker_id=w.id
        LEFT JOIN employee_gdlr eg ON eg.worker_id=w.id LEFT JOIN gdlr_categories gc ON gc.id=eg.category_id
        LEFT JOIN employee_smu es ON es.worker_id=w.id LEFT JOIN smu_catalog sc ON sc.id=es.smu_id
        LEFT JOIN pps_catalog pc ON pc.id=sc.pps_id
        LEFT JOIN LATERAL (SELECT e.stage_code FROM workforce_stage_events e
            WHERE e.worker_id=w.id AND e.confirmed AND NOT e.retracted AND e.effective_date<=%s
              AND NOT EXISTS(SELECT 1 FROM workforce_stage_events r WHERE r.replaces_id=e.id
                AND (r.confirmed OR r.retracted) AND r.effective_date<=%s)
            ORDER BY e.effective_date DESC,e.sequence DESC LIMIT 1) st ON TRUE
        WHERE w.active=1 AND ({scope})''', [day, day, *args]).fetchall())
    catalog = plain(db.native('SELECT id,name FROM pps_catalog ORDER BY id').fetchall()) if scope == 'TRUE' else []
    options = {}
    for key in ('pps', 'smu'):
        source = people if key == 'pps' or not filters['pps'] else [p for p in people if str(p['pps_id'] or 'none') in filters['pps']]
        labels = {str(p[key + '_id'] or 'none'): p[key] or ('Без ППС' if key == 'pps' else 'Без СМУ') for p in source}
        if key == 'pps':
            labels.update({str(p['id']): p['name'] for p in catalog})
        options[key] = [{'value': value, 'label': label} for value, label in sorted(labels.items(), key=lambda item: item[1])]
    selected = [p for p in people if all(not filters[k] or str(p[k + '_id'] or 'none') in filters[k] for k in ('pps', 'smu'))]
    ids = [p['id'] for p in selected]
    moves = defaultdict(list)
    if ids:
        for row in plain(db.native('''SELECT m.worker_id,m.direction,m.destination_kind,m.result_code,m.planned_date,m.actual_date,
            EXISTS(SELECT 1 FROM workforce_movements n WHERE n.rescheduled_from=m.id) replaced
            FROM workforce_movements m WHERE m.worker_id=ANY(%s)
              AND (m.planned_date BETWEEN %s AND %s OR m.actual_date BETWEEN %s AND %s)''',
            [ids, day, end2, day[:8] + '01', day]).fetchall()):
            moves[row['worker_id']].append(row)
    nodes = plain(db.native('''SELECT n.*,gc.name category_name FROM gdlr_hierarchy_nodes n
        LEFT JOIN gdlr_categories gc ON gc.id=n.category_id ORDER BY n.sort_order,n.id''').fetchall())
    periods = [end1[:8] + '01', end2[:8] + '01', str(month_after(date.fromisoformat(end2)))]
    plans = [{}, {}, {}]
    scopes = []
    # A partial SMU/actor slice must never receive a whole-PPS plan.
    plan_available = scope == 'TRUE' and not filters['smu']
    if plan_available:
        scopes = [p['name'] for p in catalog if not filters['pps'] or str(p['id']) in filters['pps']]
        # With unknown PPS membership there is no valid plan/fact comparison.
        plan_available = bool(scopes) and not any(p['pps_id'] is None for p in selected) and 'none' not in filters['pps']
    if plan_available and drill is None:
        records = db.native('''WITH versions AS (SELECT DISTINCT ON(scope,period) id,scope,period
            FROM smg_plan_versions WHERE scope=ANY(%s) AND period=ANY(%s::date[])
            ORDER BY scope,period,version DESC)
            SELECT v.scope,v.period,p.node_id,p.value FROM versions v
            JOIN smg_plan_values p ON p.version_id=v.id''', [scopes, periods]).fetchall()
        for index, period in enumerate(periods):
            by_node = defaultdict(dict)
            for r in records:
                if str(r['period']) == period:
                    by_node[r['node_id']][r['scope']] = float(r['value'])
            plans[index] = {node: sum(values.values()) for node, values in by_node.items() if len(values) == len(scopes)}
    columns = report_columns(day, end1, end2)
    warnings = []
    unbound = sum(p['pps_id'] is None for p in selected)
    if unbound:
        warnings.append(f'Без привязки СМУ к ППС: {unbound} сотрудников. Они учтены в общем своде и фильтре «Без ППС».')
    if not plan_available:
        warnings.append('План ППС не распределяется на отдельные СМУ.' if filters['smu'] else
                        'Общий план ППС недоступен при ограниченном доступе.' if scope != 'TRUE' else
                        'Сравнение с планом требует полной привязки СМУ к ППС.')
    elif not any(plans):
        warnings.append('На выбранные месяцы нет загруженного плана СМГ.')
    rows = calculate(selected, moves, nodes, day, (end1, end2), plans, drill)
    if any(row['values']['forecast1'] is None or row['values']['forecast2'] is None for row in rows):
        warnings.append('Есть неопределённый прогноз: проверьте неизвестные этапы и встречные поездки в один день.')
    return dict(date=day, end1=end1, end2=end2, options=options, filters=filters, rows=rows,
                columns=columns, notes=NOTES, warnings=warnings, total=len(selected), plan_periods=periods)


def drilldown_ids(db, raw, day):
    """Recompute one permission-scoped cell; client IDs never define membership."""
    try:
        if len(raw) > 4000:
            raise ValueError()
        data = json.loads(raw)
        required = {'date', 'end1', 'end2', 'pps', 'smu', 'node', 'metric'}
        if not isinstance(data, dict) or not required <= data.keys() or data.keys() - required - {'label'}:
            raise ValueError()
        if any(not isinstance(data[k], str) or len(data[k]) > 500 for k in required - {'pps', 'smu'}):
            raise ValueError()
        if 'label' in data and (not isinstance(data['label'], str) or len(data['label']) > 1000):
            raise ValueError()
        if data['date'] != day or data['metric'] not in DRILL_METRICS or not data['node']:
            raise ValueError()
        query = MultiDict((k, data[k]) for k in ('date', 'end1', 'end2'))
        for key in ('pps', 'smu'):
            if not isinstance(data[key], list) or len(data[key]) > 100 or any(not isinstance(v, str) for v in data[key]):
                raise ValueError()
            query.setlist(key, data[key])
    except (ValueError, TypeError):
        abort(400, description='Некорректный фильтр свода перевахтовки. Откройте ячейку отчёта заново.')
    args = arguments(query)
    drill = {'node': data['node'], 'metric': data['metric']}
    build_report(db, *args, drill=drill)
    if 'ids' not in drill:
        abort(400, description='Категория свода больше недоступна. Обновите отчёт.')
    return drill['ids']


def workbook(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook(); sheet = book.active; sheet.title = 'Свод перевахтовки'
    sheet.append(['Свод перевахтовки на ' + report['date']])
    sheet.append(['Категория ГДЛР'] + [c['group'] for c in report['columns']])
    sheet.append([None] + [c['label'] for c in report['columns']])
    sheet.merge_cells('A2:A3')
    start = 2
    for index in range(1, len(report['columns']) + 1):
        if index == len(report['columns']) or report['columns'][index]['group'] != report['columns'][index - 1]['group']:
            if index + 1 > start:
                sheet.merge_cells(start_row=2, start_column=start, end_row=2, end_column=index + 1)
            start = index + 2
    for row in report['rows']:
        sheet.append([row['label']] + [row['values'][c['key']] if row['values'][c['key']] is not None else '—' for c in report['columns']])
        for cell in sheet[sheet.max_row]:
            if cell.data_type == 'f':
                cell.data_type = 's'
            cell.font = Font(name='Calibri', size=10, bold=row['group'])
            if row['group']:
                cell.fill = PatternFill('solid', fgColor='E7EEF5')
    sheet.freeze_panes = 'B4'
    sheet.auto_filter.ref = f'A3:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    sheet.print_title_rows = '1:3'
    sheet.column_dimensions['A'].width = 42
    for index in range(2, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 16
    sheet.row_dimensions[2].height = 26
    sheet.row_dimensions[3].height = 42
    for row in (sheet[2], sheet[3]):
        for cell in row:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='315878')
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    note = book.create_sheet('Параметры и правила')
    note.append(['Параметр', 'Значение'])
    for key in ('date', 'end1', 'end2'):
        note.append([{'date': 'Отчётная дата', 'end1': 'Конец периода 1', 'end2': 'Конец периода 2'}[key], report[key]])
    for key in ('pps', 'smu'):
        options = {o['value']: o['label'] for o in report['options'][key]}
        note.append([key.upper(), ', '.join(options.get(v, 'Недоступное значение') for v in report['filters'][key]) or 'Все доступные'])
    for text in report['warnings'] + report['notes']:
        note.append(['Примечание', text])
    note.column_dimensions['A'].width = 24; note.column_dimensions['B'].width = 110
    for row in note:
        for cell in row:
            if cell.data_type == 'f': cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    stream = BytesIO(); book.save(stream); stream.seek(0)
    return stream


def register_rotation_summary(app, database, roles_required):
    @app.get('/api/workforce/rotation-summary')
    @app.get('/api/workforce/rotation-summary.xlsx')
    @roles_required(*READERS)
    def rotation_summary():
        args = arguments(request.args)
        db = database()
        with db:
            db.execute('BEGIN')
            report = build_report(db, *args)
        if request.path.endswith('.xlsx'):
            response = send_file(workbook(report), as_attachment=True,
                                 download_name='Свод перевахтовки ' + report['date'] + '.xlsx',
                                 mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        else:
            response = jsonify(report)
        response.headers['Cache-Control'] = 'private, no-store'
        return response
