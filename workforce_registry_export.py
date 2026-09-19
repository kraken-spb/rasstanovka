"""Excel copies of the scoped HR registry, using the same rows as the screen."""
from datetime import date
from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


MAX_ROWS = 25000
SECTIONS = {'rotation': 'Перевахта', 'recruitment': 'Комплектация'}
HEADERS = (
    '№ п/п', 'ФИО', 'Табельный номер', 'Телефон', 'E-mail', 'Гражданство', 'Город отправления', 'Проект', 'СМУ', 'Организация-работодатель',
    'Должность', 'Категория ГДЛР', 'Статус сотрудника', 'Состояние',
    'Дата начала статуса', 'Дата заезда', 'Прогноз окончания вахты',
    'Заезд / выезд', 'Плановая дата поездки', 'Тип заезда/выезда',
    'График вахтования', 'Дата окончания МО', 'Следующий заезд', 'Проживание', 'Дата начала МО', 'Подразделение',
)
WIDTHS = (8, 38, 19, 24, 32, 24, 28, 24, 35, 28, 40, 32, 24, 26, 19, 19, 22, 18, 22, 24, 26, 20, 20, 22, 20, 38)


COLUMN_ORDER = (0, 9, 12, 13, 20, 2, 1, 10, 11, 5, 3, 4, 8, 25, 6, 24, 21, 15, 16, 7, 14, 17, 18, 19, 22, 23)
HEADERS = tuple(HEADERS[i] for i in COLUMN_ORDER)
WIDTHS = tuple(WIDTHS[i] for i in COLUMN_ORDER)

def excel_date(value):
    return date.fromisoformat(value) if isinstance(value, str) and value else value or None


def registry_workbook(rows, day, section):
    """Write values only: source text must never be interpreted as an Excel formula."""
    book = Workbook(write_only=True)
    sheet = book.create_sheet(SECTIONS[section])
    sheet.freeze_panes = 'D5'
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for index, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    font = Font(name='Calibri', size=11, color='243E30')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='176D50')
    alternate = PatternFill('solid', fgColor='F0F6F3')
    alignment = Alignment(vertical='top', wrap_text=True)

    def style(*, header=False, stripe=False, date_value=False):
        template = WriteOnlyCell(sheet)
        template.font = header_font if header else font
        template.alignment = alignment
        if header or stripe:
            template.fill = header_fill if header else alternate
        if date_value:
            template.number_format = 'dd.mm.yyyy'
        return template._style

    styles = {(False, False): style(), (False, True): style(stripe=True),
              (True, False): style(header=True), (True, True): style(header=True, date_value=True),
              (False, 'date'): style(date_value=True), (True, 'date'): style(stripe=True, date_value=True)}

    def cells(values, *, header=False, stripe=False):
        result = []
        plain_style = styles[(header, stripe)]
        date_style = styles[(header, stripe) if header else (True, 'date') if stripe else (False, 'date')]
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell._style = date_style if isinstance(value, date) else plain_style
            result.append(cell)
        return result

    sheet.row_dimensions[1].height = 24
    sheet.append(cells([None, SECTIONS[section], 'Отчётная дата', excel_date(day)]))
    sheet.append(cells(['Всего', len(rows), None, None, 'Весь список по фильтрам']))
    sheet.append([])
    sheet.row_dimensions[4].height = 44
    sheet.append(cells(HEADERS, header=True))
    for index, row in enumerate(rows, 1):
        movement, rotation = row.get('movement') or {}, row.get('rotation') or {}
        direction = {'arrival': 'Заезд', 'departure': 'Выезд'}.get(movement.get('direction'), '')
        values = [index, row['full_name'], row.get('personnel_no') or '', row.get('phone'), row.get('email'), row.get('citizenship'), row.get('origin_city'), row.get('project'),
                  row.get('department'), row.get('employer'), row.get('profession'), row.get('category'),
                  row.get('employment'), row.get('stage') or 'Без подтверждённого состояния',
                  excel_date(row.get('effective_date')), excel_date(row.get('arrival_date')),
                  excel_date(row.get('forecast_departure_date')), direction,
                  excel_date(movement.get('planned_date')), movement.get('basis'), rotation.get('schedule'),
                  excel_date(rotation.get('leave_end_date')), excel_date(rotation.get('next_arrival_date')), row.get('accommodation'), excel_date(row.get('leave_start_date')), row.get('division')]
        sheet.append(cells([values[i] for i in COLUMN_ORDER], stripe=index % 2 == 0))
    sheet.auto_filter.ref = f'A4:{get_column_letter(len(HEADERS))}{len(rows) + 4}'
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def registry_download(rows, day, section):
    response = send_file(BytesIO(registry_workbook(rows, day, section)),
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=f'{SECTIONS[section]}_{day}.xlsx', max_age=0)
    response.headers['Cache-Control'] = 'private, no-store'
    response.headers['X-Export-Row-Count'] = str(len(rows))
    return response
