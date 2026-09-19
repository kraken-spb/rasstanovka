import os
import json
from urllib.parse import urlencode
import unittest
from collections import defaultdict
from io import BytesIO
from uuid import uuid4

from flask import Flask
from werkzeug.datastructures import MultiDict
from werkzeug.exceptions import BadRequest

from rotation_summary import arguments, calculate, facts_for, register_rotation_summary, report_columns, workbook
from rotation_summary import DRILL_METRICS, drilldown_ids


DAY = '2026-09-18'
ENDS = ('2026-09-30', '2026-10-31')


def person(id=1, **changes):
    return dict(dict(id=id, category_id=1, category='Монтажник', stage_code='stage.onsite', employment_code='employment.staff'), **changes)


def move(direction='arrival', **changes):
    return dict(dict(direction=direction, destination_kind='site', planned_date='2026-09-20', actual_date=None, result_code=None), **changes)


class RotationSummaryCalculationTest(unittest.TestCase):
    def test_drill_members_match_hierarchy_and_each_metric(self):
        people = [person(), person(2, stage_code='stage.leave'), person(3, category_id=2)]
        trips = {1:[move('departure')], 2:[move(),move()], 3:[]}
        nodes = [dict(id='all',parent_id=None,label='Все',category_id=None,is_group=True),
                 dict(id='fit',parent_id='all',label='Монтажник',category_id=1,is_group=False)]
        for metric in DRILL_METRICS:
            for node in ('all','fit','extra:2'):
                drill = dict(node=node,metric=metric)
                rows = calculate(people,trips,nodes,DAY,ENDS,[{},{},{}],drill)
                row = next(r for r in rows if r['id']==node)
                self.assertEqual(len(drill['ids']),row['values'][metric])
                candidates = people if node=='all' else [p for p in people if p['category_id']==(1 if node=='fit' else 2)]
                self.assertEqual(drill['ids'],[p['id'] for p in candidates if facts_for(p,trips[p['id']],DAY,ENDS)[metric]==1])

    def test_invalid_drill_never_falls_back_to_all_workers(self):
        valid = dict(date=DAY,end1=ENDS[0],end2=ENDS[1],pps=[],smu=[],node='all',metric='onsite')
        for value in ['', 'null', '[]', '{', json.dumps(dict(valid,metric='plan1')),json.dumps(dict(valid,date='2026-09-17')),
                      json.dumps(dict(valid,pps=[1])),json.dumps(dict(valid,worker_ids=[1]))]:
            with self.subTest(value=value), Flask(__name__).test_request_context(), self.assertRaises(BadRequest):
                drilldown_ids(None,value,DAY)

    def test_unknown_forecast_cannot_return_a_partial_cohort(self):
        nodes=[dict(id='all',parent_id=None,label='Все',category_id=None,is_group=True)]
        with Flask(__name__).test_request_context(), self.assertRaises(BadRequest):
            calculate([person(),person(2,stage_code=None)],{},nodes,DAY,ENDS,[{},{},{}],dict(node='all',metric='forecast1'))

    def test_calendar_periods_year_end_and_validation(self):
        with Flask(__name__).test_request_context():
            self.assertEqual(arguments(MultiDict({'date':'2026-12-18'}))[:3], ('2026-12-18','2026-12-31','2027-01-31'))
            for query in [dict(date='2026-09-18',end1='2026-09-17'),dict(date='2026-09-18',pps='1 OR 1=1'),dict(date='2026-09-18',end2='2028-01-01')]:
                with self.assertRaises(BadRequest): arguments(MultiDict(query))

    def test_ticket_is_not_presence_and_duplicate_trips_count_one_person(self):
        values=facts_for(person(stage_code='stage.inbound'),[move(),move()],DAY,ENDS)
        self.assertEqual(values['onsite'],0)
        self.assertEqual(values['inbound'],1)
        self.assertEqual(values['arrived'],0)
        self.assertEqual(values['arrivals1'],1)
        self.assertEqual(values['forecast1'],1)

    def test_month_groups_keep_fact_and_forecast_periods_distinct(self):
        columns = report_columns(DAY, *ENDS)
        by_key = {c['key']: c for c in columns}
        self.assertEqual(len(by_key), len(columns))
        self.assertEqual(by_key['onsite']['group'], 'На 18.09.2026')
        self.assertEqual(by_key['arrived']['group'], 'Сентябрь 2026')
        self.assertEqual(by_key['arrived']['period'], '01.09.2026 — 18.09.2026')
        self.assertEqual(by_key['arrivals1']['period'], '19.09.2026 — 30.09.2026')
        self.assertEqual(by_key['forecast1']['group'], 'Сентябрь 2026')
        self.assertEqual(by_key['forecast2']['group'], 'Октябрь 2026')
        self.assertEqual(by_key['plan3']['group'], 'Ноябрь 2026')
        december = {c['key']: c for c in report_columns('2026-12-31', '2026-12-31', '2027-01-31')}
        self.assertEqual(december['forecast1']['group'], 'Декабрь 2026')
        self.assertEqual(december['forecast2']['group'], 'Январь 2027')
        self.assertEqual(december['plan3']['group'], 'Февраль 2027')
        custom = {c['key']: c for c in report_columns(DAY, '2026-10-05', '2026-11-10')}
        self.assertEqual(custom['arrivals1']['group'], '19.09.2026 — 05.10.2026')
        self.assertEqual(custom['arrived']['group'], 'Сентябрь 2026 · факт')
        self.assertEqual(custom['arrivals2']['group'], '06.10.2026 — 10.11.2026')

    def test_cancelled_rescheduled_and_pvp_do_not_raise_forecast(self):
        trips=[move(result_code='result.cancelled'),move(replaced=True),move(result_code='result.postponed'),move(destination_kind='pvp')]
        values=facts_for(person(stage_code='stage.leave'),trips,DAY,ENDS)
        self.assertEqual(values['forecast1'],0)
        self.assertEqual(values['arrivals1'],0)

    def test_forecast_follows_sequence_without_negative_or_double_count(self):
        trips=[move('departure'),move('departure'),move(planned_date='2026-10-03'),move(planned_date='2026-10-04')]
        values=facts_for(person(),trips,DAY,ENDS)
        self.assertEqual(values['departures1'],1)
        self.assertEqual(values['forecast1'],0)
        self.assertEqual(values['forecast2'],1)

    def test_ambiguous_same_day_and_unknown_initial_state_are_not_zero(self):
        values=facts_for(person(),[move(),move('departure')],DAY,ENDS)
        self.assertIsNone(values['forecast1']);self.assertIsNone(values['forecast2'])
        values=facts_for(person(stage_code=None),[],DAY,ENDS)
        self.assertEqual(values['unknown'],1);self.assertIsNone(values['forecast1'])

    def test_actual_movement_is_dated_and_distinct_from_stage(self):
        values=facts_for(person(stage_code='stage.leave'),[move(result_code='result.happened',actual_date='2026-09-17'),move(result_code='result.happened',actual_date='2026-09-19')],DAY,ENDS)
        self.assertEqual(values['arrived'],1);self.assertEqual(values['onsite'],0)
        self.assertEqual(values['arrivals1'],0)

    def test_hierarchy_current_categories_missing_and_explicit_zero_plan(self):
        nodes=[dict(id='total',parent_id=None,label='Итого',category_id=None,is_group=True),
               dict(id='workers',parent_id='total',label='Рабочие',category_id=None,is_group=True),
               dict(id='fitter',parent_id='workers',label='Монтажник',category_id=1,is_group=False)]
        rows=calculate([person(),person(2,category_id=2,category='Новая категория'),person(3,category_id=None,category=None)],defaultdict(list),nodes,DAY,ENDS,[{'total':0},{},{}])
        self.assertEqual(rows[0]['values']['total'],3)
        self.assertEqual(rows[0]['values']['delta1'],3)
        self.assertIsNone(rows[0]['values']['delta2'])
        self.assertEqual(rows[1]['values']['total'],1)
        self.assertEqual({r['label'] for r in rows[-2:]},{'Новая категория','Без категории ГДЛР'})
        self.assertIsNone(rows[0]['values']['dismissed'])

    def test_excel_matches_report_and_does_not_execute_catalog_formula(self):
        from openpyxl import load_workbook
        report=dict(date=DAY,end1=ENDS[0],end2=ENDS[1],columns=[dict(key='total',label='Всего',group='Факт')],
                    rows=[dict(label='=1+1',group=False,values={'total':7})],filters={'pps':[],'smu':[]},options={'pps':[],'smu':[]},warnings=[],notes=['Правило'])
        book=load_workbook(BytesIO(workbook(report).read()))
        self.assertEqual(book.worksheets[0]['A4'].data_type,'s')
        self.assertEqual(book.worksheets[0]['B4'].value,7)
        self.assertEqual(book.worksheets[0].freeze_panes,'B4')

    def test_excel_month_headers_preserve_every_report_value(self):
        from openpyxl import load_workbook
        columns = report_columns(DAY, *ENDS)
        values = {column['key']: index for index, column in enumerate(columns)}
        values['plan3'] = None
        report = dict(date=DAY, end1=ENDS[0], end2=ENDS[1], columns=columns,
                      rows=[dict(label='Итого', group=True, values=values)],
                      filters={'pps':[], 'smu':[]}, options={'pps':[], 'smu':[]}, warnings=[], notes=[])
        sheet = load_workbook(workbook(report)).worksheets[0]
        self.assertEqual(sheet['J2'].value, 'Сентябрь 2026')
        self.assertIn('J2:P2', {str(area) for area in sheet.merged_cells.ranges})
        self.assertEqual(sheet['Q2'].value, 'Октябрь 2026')
        self.assertEqual(sheet['V2'].value, 'Ноябрь 2026')
        for index, column in enumerate(columns, 2):
            self.assertEqual(sheet.cell(3, index).value, column['label'])
            self.assertEqual(sheet.cell(4, index).value, values[column['key']] if values[column['key']] is not None else '—')


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select isolated PostgreSQL database.')
class RotationSummaryPostgresTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import test_workforce_api as fixture
        fixture.WorkforceApiTest.setUpClass()

    def setUp(self):
        import test_workforce_api as fixture
        self.f=fixture.WorkforceApiTest();self.f.setUp();self.addCleanup(self.f.tearDown)
        f=self.f
        register_rotation_summary(f.app,lambda:f.db,lambda *roles:lambda fn:fn)
        self.pps=f.db.native('INSERT INTO pps_catalog(name,name_key) VALUES (%s,%s) RETURNING id',('Тест '+f.suffix,f.suffix)).fetchone()['id']
        self.smu=f.db.native("INSERT INTO smu_catalog(name,pps_id,edit_token) VALUES (%s,%s,'test') RETURNING id",('СМУ '+f.suffix,self.pps)).fetchone()['id']
        f.db.native("INSERT INTO employee_smu(worker_id,smu_id,source_department) VALUES (%s,%s,'TEST-SMU') ON CONFLICT(worker_id) DO UPDATE SET smu_id=excluded.smu_id",(f.worker,self.smu))
        self.event=f.db.native("INSERT INTO workforce_stage_events(worker_id,stage_code,effective_date,confirmed,reason,created_by,request_key) VALUES (%s,'stage.onsite','2026-09-15',TRUE,'Проверка',%s,%s) RETURNING id",(f.worker,f.users['admin']['id'],uuid4())).fetchone()['id']
        self.path='rotation-summary?date='+DAY+'&pps='+str(self.pps)

    def test_pps_smu_intersection_dated_corrections_and_export(self):
        f=self.f
        response=f.request('get',self.path,role='rotation');self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.json['total'],1)
        self.assertEqual(response.json['rows'][0]['values']['onsite'],1)
        f.db.native("INSERT INTO workforce_stage_events(worker_id,stage_code,effective_date,confirmed,reason,created_by,request_key,replaces_id) VALUES (%s,'stage.leave','2026-09-16',TRUE,'Исправление',%s,%s,%s)",(f.worker,f.users['admin']['id'],uuid4(),self.event))
        response=f.request('get',self.path,role='rotation')
        self.assertEqual(response.json['rows'][0]['values']['onsite'],0)
        self.assertEqual(response.json['rows'][0]['values']['leave'],1)
        self.assertEqual(f.request('get',self.path+'&smu=999999999',role='rotation').json['total'],0)
        exported=f.request('get',self.path.replace('rotation-summary?', 'rotation-summary.xlsx?'),role='rotation')
        self.assertEqual(exported.status_code,200)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(exported.data))
        self.assertEqual(book.worksheets[0]['B4'].value,1)
        self.assertIn('no-store',exported.headers['Cache-Control'])

    def test_scope_and_plan_never_leak_through_filters(self):
        f=self.f
        f.db.native("INSERT INTO employee_smu(worker_id,smu_id,source_department) VALUES (%s,%s,'OTHER-SMU') ON CONFLICT(worker_id) DO UPDATE SET smu_id=excluded.smu_id",(f.ids[1],self.smu))
        # Even when an out-of-scope worker is in the same PPS, actor_scope applies first.
        result=f.request('get',self.path,role='rotation').json
        self.assertEqual(result['total'],1)
        self.assertIsNone(result['rows'][0]['values']['plan1'])
        f.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s",(f.users['rotation']['id'],))
        result=f.request('get',self.path,role='rotation').json
        self.assertEqual(result['total'],0);self.assertEqual(result['options']['pps'],[])
        f.db.native('UPDATE users SET active=0 WHERE id=%s',(f.users['rotation']['id'],))
        self.assertEqual(f.request('get',self.path,role='rotation').status_code,403)

    def test_drill_exact_cohort_scope_filters_and_zero(self):
        f=self.f
        report=f.request('get',self.path,role='rotation').json
        selector=dict(date=DAY,end1=ENDS[0],end2=ENDS[1],pps=[str(self.pps)],smu=[],node=str(report['rows'][0]['id']),metric='onsite')
        def path(**changes):
            return 'people?'+urlencode(dict(date=DAY,section='rotation',rotation_summary=json.dumps(dict(selector,**changes))))
        # Report members also include profiles without rotation source membership.
        for metric in DRILL_METRICS:
            response=f.request('get',path(metric=metric),role='rotation')
            self.assertEqual(response.status_code,200,response.data)
            count=report['rows'][0]['values'][metric]
            self.assertEqual(response.json['totals']['total'],count)
            self.assertEqual([r['id'] for r in response.json['rows']],[f.worker] if count else [])
        self.assertEqual(f.request('get',path()+'&stage=stage.leave',role='rotation').json['totals']['total'],0)
        self.assertEqual(f.request('get',path()+'&offset=1&limit=1',role='rotation').json['rows'],[])
        identity=f.request('get',path()+'&identity_options=full_name',role='rotation')
        self.assertEqual(identity.status_code,200,identity.data)
        exported=f.request('get',path().replace('people?','people/export?'),role='rotation')
        self.assertEqual(exported.status_code,200,exported.data[:500])
        self.assertEqual(f.request('get',path(node='deleted'),role='rotation').status_code,400)
        f.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s",(f.users['rotation']['id'],))
        self.assertEqual(f.request('get',path(),role='rotation').json['totals']['total'],0)

    def test_latest_monthly_plan_explicit_zero_and_smu_slice(self):
        f=self.f
        node=f.db.native('SELECT id FROM gdlr_hierarchy_nodes WHERE parent_id IS NULL AND is_group ORDER BY sort_order LIMIT 1').fetchone()['id']
        scope=f.db.native('SELECT name FROM pps_catalog WHERE id=%s',(self.pps,)).fetchone()['name']
        for version,value in ((1,10),(2,0)):
            plan=f.db.native("INSERT INTO smg_plan_versions(scope,period,version,source,created_by,request_key) VALUES (%s,'2026-09-01',%s,'{}',%s,%s) RETURNING id",(scope,version,f.users['admin']['id'],uuid4())).fetchone()['id']
            f.db.native('INSERT INTO smg_plan_values(version_id,node_id,value) VALUES (%s,%s,%s)',(plan,node,value))
        result=f.request('get',self.path).json
        self.assertEqual(result['rows'][0]['values']['plan1'],0)
        self.assertEqual(result['rows'][0]['values']['delta1'],1)
        self.assertIsNone(result['rows'][0]['values']['plan2'])
        self.assertIsNone(f.request('get',self.path+'&smu='+str(self.smu)).json['rows'][0]['values']['plan1'])
