import os
import unittest
from urllib.parse import urlencode
from uuid import uuid4

import test_workforce_api as fixtures


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class WorkforceDateFiltersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        self.f.app.config['TESTING'] = False
        for worker in self.f.ids:
            self.f.add_source_record(worker, 'urp:К15', role='recruitment', sheet='ПВП')
        self.f.db.native('''UPDATE workforce_profiles SET arrival_date='2026-09-01',
            forecast_departure_date='2026-10-01' WHERE worker_id=%s''', (self.f.worker,))

    def get(self, values=(), role='admin', extra=''):
        query = [('date', '2026-09-16'), ('section', 'recruitment'), ('date_options', '1'),
                 ('department', 'TEST-SMU'), ('department', 'OTHER-SMU'), *values]
        response = self.f.request('get', 'people?' + urlencode(query) + extra, role=role)
        self.assertEqual(response.status_code, 200, response.data)
        return response.json

    def test_date_selection_pagination_totals_scope_and_cached_options(self):
        all_rows = self.get(extra='&limit=1&offset=1')
        self.assertEqual(all_rows['totals']['total'], 2)
        self.assertEqual(all_rows['date_options']['arrival_date'], ['2026-09-01', '__none__'])
        selected = self.get([('arrival_date', '2026-09-01')])
        self.assertEqual([r['id'] for r in selected['rows']], [self.f.worker])
        self.assertEqual(selected['totals']['total'], 1)
        self.assertEqual(selected['date_options'], all_rows['date_options'])
        both = [('arrival_date', '2026-09-01'), ('arrival_date', '__none__')]
        self.assertEqual(self.get(both)['totals']['total'], 2)
        self.assertEqual(self.get([('arrival_date', '__none__')])['rows'][0]['id'], self.f.ids[1])
        self.assertEqual(self.get(both + [('forecast_departure_date', '2026-10-01')])['totals']['total'], 1)
        self.assertEqual(self.get([('arrival_date', '2026-09-01'), ('forecast_departure_date', '__none__')])['totals']['total'], 0)
        scoped = self.get(role='recruitment')
        self.assertEqual(scoped['totals']['total'], 1)
        self.assertEqual(scoped['date_options']['arrival_date'], ['2026-09-01'])
        self.assertEqual(self.get(role='recruitment'), scoped)
        self.assertEqual(self.get([('arrival_date', '__none__')], role='recruitment')['rows'], [])

    def test_forecast_filter_uses_displayed_rotation_date_and_export(self):
        f = self.f
        schedule = f.db.native('''INSERT INTO workforce_rotation_schedules
            (name, onsite_days, leave_days, travel_days, updated_by) VALUES (%s,30,30,2,%s) RETURNING id''',
            (f.suffix, f.users['admin']['id'])).fetchone()[0]
        f.db.native('''INSERT INTO workforce_rotations
            (worker_id,schedule_id,schedule_snapshot,start_date,planned_end_date,leave_end_date,
             next_arrival_date,request_key,created_by,updated_by)
            VALUES (%s,%s,'{}','2026-09-01','2026-09-30','2026-10-30','2026-11-01',%s,%s,%s)''',
            (f.worker, schedule, uuid4(), f.users['admin']['id'], f.users['admin']['id']))
        selected = self.get([('forecast_departure_date', '2026-09-30')])
        self.assertEqual([r['id'] for r in selected['rows']], [f.worker])
        self.assertEqual(selected['rows'][0]['forecast_departure_date'], '2026-09-30')
        self.assertEqual(selected['date_options']['forecast_departure_date'], ['2026-09-30', '__none__'])
        self.assertEqual(self.get([('forecast_departure_date', '2026-10-01')])['totals']['total'], 0)
        self.assertEqual(self.get([('forecast_departure_date', '2026-09-30'), ('forecast_departure_date', '__none__')])['totals']['total'], 2)
        response = f.request('get', 'people/export?date=2026-09-16&section=recruitment&department=TEST-SMU&forecast_departure_date=2026-09-30')
        self.assertEqual(response.status_code, 200, response.data[:200])
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.active.max_row, 5)
        self.assertIn('Тестовый Сотрудник', [cell.value for cell in book.active[5]])

    def test_leave_start_matches_profile_filter_sort_and_export_in_both_sections(self):
        import json
        from io import BytesIO
        from openpyxl import load_workbook
        f = self.f
        f.db.native("UPDATE workforce_profiles SET leave_start_date='2026-09-20' WHERE worker_id=%s", (f.worker,))
        for worker in f.ids:
            f.add_source_record(worker, 'urp:П15', role='rotation', sheet='Неявка')
        for section in ('rotation', 'recruitment'):
            query = urlencode({'section':section,'date':'2026-09-16','department':'TEST-SMU',
                               'leave_start_date':'2026-09-20','date_options':'1',
                               'sort':json.dumps([{'field':'leave_start_date','direction':'asc'}])})
            response = f.request('get', 'people?' + query)
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual([r['id'] for r in response.json['rows']], [f.worker])
            self.assertEqual(response.json['rows'][0]['leave_start_date'], '2026-09-20')
            self.assertEqual(response.json['date_options']['leave_start_date'], ['2026-09-20'])
            exported = f.request('get', 'people/export?' + query)
            self.assertEqual(exported.status_code, 200, exported.data[:100])
            book = load_workbook(BytesIO(exported.data))
            labels = [c.value for c in book.active[4]]
            position = labels.index('Дата начала МО')
            self.assertEqual(labels[position+1], 'Дата окончания МО')
            self.assertEqual(book.active.cell(5,position+1).value.date().isoformat(), '2026-09-20')
            book.close()
        missing = self.get([('leave_start_date','__none__')])
        self.assertEqual([r['id'] for r in missing['rows']], [f.ids[1]])
        self.assertIsNone(missing['rows'][0]['leave_start_date'])

    def test_invalid_dates_rejected(self):
        for query in ['arrival_date=2026-02-30', 'forecast_departure_date=bad',
                      'arrival_date=2026-09-01&arrival_date=bad',
                      urlencode([('forecast_departure_date', str(i)) for i in range(101)])]:
            with self.subTest(query=query):
                self.assertEqual(self.f.request('get', 'people?' + query).status_code, 400)
