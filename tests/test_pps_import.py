import io
import unittest

import openpyxl
import test_staffing
from staffing_import import ImportProblem, apply_attendance, parse_attendance


class PpsImportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        test_staffing.StaffingWorkflowTest.setUpClass()

    @classmethod
    def tearDownClass(cls):
        test_staffing.StaffingWorkflowTest.tearDownClass()

    def setUp(self):
        self.c = test_staffing.StaffingWorkflowTest()
        self.c.setUp()
        self.addCleanup(self.c.doCleanups)

    def apply(self, label, rows, content=None):
        parsed = parse_attendance(content if content is not None else test_staffing.workbook_bytes(rows), label + '.xlsx')
        parsed['source_label'] = label
        with self.c.app.app_context():
            return apply_attendance(self.c.module.get_db(), parsed, self.c.admin_id, self.c.module.utc_now)

    def test_latest_import_per_pps_and_people_directory(self):
        self.apply('ППС15', [{'number': '701'}, {'number': '702', 'name': 'Второй Работник'}])
        self.apply('ППС19', [{'number': '801', 'name': 'Новый Работник', 'crew': ''}])
        rows = self.c.group_table()
        self.assertEqual({r['personnel_no']: r['pps'] for r in rows}, {'701': 'ППС15', '702': 'ППС15', '801': 'ППС19'})
        self.assertEqual(next(r['crew_name'] for r in rows if r['pps'] == 'ППС19'), 'ППС19 · Бригада не указана')
        self.assertEqual(len(self.c.admin.get('/api/staffing/people').json['people']), 3)
        self.assertEqual(len(self.c.foreman.get('/api/staffing?date=2026-09-12&shift=all').json['rows']), 0)
        self.apply('ППС19', [{'number': '802', 'name': 'Другой Новый Работник', 'crew': ''}])
        rows = self.c.group_table()
        self.assertEqual({r['personnel_no'] for r in rows}, {'701', '702', '802'})
        summary = self.c.admin.get('/api/staffing?date=2026-09-12&shift=all&view=summary').json
        self.assertEqual({r['pps'] for r in summary['index']}, {'ППС15', 'ППС19'})
        self.assertTrue(all(r['pps'] in r['search_fields'] for r in summary['index']))

    def test_cross_pps_identity_is_rejected_and_retry_is_idempotent(self):
        content = test_staffing.workbook_bytes([{'number': '701'}])
        first = self.apply('ППС15', [], content)
        self.assertTrue(self.apply('ППС15', [], content)['already_imported'])
        with self.assertRaises(ImportProblem):
            self.apply('ППС19', [{'number': '701'}, {'number': '802'}])
        with self.c.app.app_context():
            db = self.c.module.get_db()
            self.assertEqual(db.execute('SELECT COUNT(*) FROM staffing_imports').fetchone()[0], 1)
            self.assertEqual(db.execute('SELECT pps FROM workers WHERE personnel_no=?', ('701',)).fetchone()[0], 'ППС15')
            self.assertIsNone(db.execute('SELECT 1 FROM workers WHERE personnel_no=?', ('802',)).fetchone())

    def test_pps19_header_alias_and_notes_are_not_brigades(self):
        workbook = openpyxl.load_workbook(io.BytesIO(test_staffing.workbook_bytes([{}])))
        sheet = workbook['Явка']
        sheet.cell(1, 4, 'Табельный номер')
        sheet.cell(1, 25, 'Примечание')
        sheet.cell(2, 25, 'Текст примечания')
        content = io.BytesIO(); workbook.save(content)
        parsed = parse_attendance(content.getvalue(), 'ППС19.xlsx')
        self.assertEqual(parsed['issues'], [])
        self.assertEqual(parsed['rows'][0]['crew_name'], '')
        self.assertEqual(parsed['summary']['without_crew'], 1)

    def test_preview_binds_source_label(self):
        content = test_staffing.workbook_bytes([{}])
        def post(path, label, token=''):
            return self.c.admin.post(path, data={'file': (io.BytesIO(content), 'ППС.xlsx'),
                'source_label': label, 'preview_token': token}, headers={'X-CSRF-Token': 'staffing-csrf'})
        preview = post('/api/staffing/import/preview', 'ППС19')
        self.assertEqual(preview.status_code, 200)
        token = preview.json['preview_token']
        self.assertEqual(post('/api/staffing/import/apply', 'ППС15', token).status_code, 400)
        self.assertEqual(post('/api/staffing/import/apply', 'ППС19', token).status_code, 200)
        self.assertEqual(self.c.group_table()[0]['pps'], 'ППС19')
