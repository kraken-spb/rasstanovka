import importlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl


class StaffingExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bootstrap = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {
            'DATABASE_PATH': str(Path(cls.bootstrap.name) / 'bootstrap.db'),
            'ADMIN_PASSWORD': 'export-test-password-123',
            'SECRET_KEY': 'export-test-secret-at-least-thirty-two-characters',
        }):
            cls.module = importlib.import_module('app')

    @classmethod
    def tearDownClass(cls):
        cls.bootstrap.cleanup()

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.addCleanup(setattr, self.module, 'DATABASE_PATH', self.module.DATABASE_PATH)
        self.module.DATABASE_PATH = Path(temporary.name) / 'export.db'
        with self.module.app.app_context():
            self.module.init_db()
            db = self.module.get_db()
            self.admin_id = db.execute("SELECT id FROM users WHERE role='admin'").fetchone()[0]
            self.foreman_id = self.user(db, 'export-foreman', 'foreman')
            self.other_id = self.user(db, 'export-other', 'foreman')
            self.viewer_id = self.user(db, 'export-viewer', 'viewer')
            self.crew_id = db.execute("INSERT INTO crews(name,owner_user_id,created_at,linear_itr,brigadier) VALUES ('Экспорт',?,'now','ИТР бригады','Бригадир бригады')", (self.foreman_id,)).lastrowid
            self.other_crew_id = db.execute("INSERT INTO crews(name,owner_user_id,created_at,linear_itr,brigadier) VALUES ('Другая',?,'now','Чужой ИТР','Чужой бригадир')", (self.other_id,)).lastrowid
            object_id = db.execute("INSERT INTO objects(name) VALUES ('Группа А')").lastrowid
            self.site_id = db.execute("INSERT INTO subobjects(object_id,name) VALUES (?, 'Подобъект 1')", (object_id,)).lastrowid
            self.visible_worker = self.worker(db, '=Формула', '000123', contractor='Подрядчик', profession='Должность', gsp='ГСП')
            self.legacy_worker = self.worker(db, 'Ночной сотрудник', '000124')
            self.hidden_worker = self.worker(db, 'Чужой сотрудник', '000125')
            self.inactive_worker = self.worker(db, 'Неактивный сотрудник', '000126', active=0)
            category = db.execute("INSERT INTO gdlr_categories(name,name_key,edit_token,updated_by,updated_at) VALUES ('Ручная ГДЛР','ручная гдлр','token',?,'now')", (self.admin_id,)).lastrowid
            db.execute('INSERT INTO employee_gdlr(worker_id,category_id,edit_token,updated_by,updated_at) VALUES (?,?,\'token\',?,\'now\')', (self.visible_worker, category, self.admin_id))
            db.execute("INSERT INTO staffing_row_details(worker_id,linear_itr_override,brigadier_override,edit_token,updated_by,updated_at) VALUES (?, 'ИТР вручную', 'Бригадир вручную', 'token', ?, 'now')", (self.visible_worker, self.admin_id))
            self.assignment(db, self.visible_worker, '1 смена', self.other_id, self.crew_id, '=Работодатель')
            self.assignment(db, self.legacy_worker, 'Ночная смена', self.other_id, self.crew_id, 'Работодатель ночи')
            self.assignment(db, self.hidden_worker, '1 смена', self.other_id, self.other_crew_id, 'Чужой работодатель')
            self.assignment(db, self.inactive_worker, '1 смена', self.foreman_id, None, 'Исторический работодатель')
            db.commit()
        self.admin = self.client(self.admin_id)
        self.foreman = self.client(self.foreman_id)
        self.viewer = self.client(self.viewer_id)

    def user(self, db, username, role):
        return db.execute('INSERT INTO users(username,password_hash,full_name,role,created_at) VALUES (?,\'hash\',?,?,\'now\')', (username, username, role)).lastrowid

    def worker(self, db, name, number, contractor='', profession='', gsp='', active=1):
        return db.execute('''INSERT INTO workers(full_name,personnel_no,contractor,employer,profession,gsp_profession,category,active)
            VALUES (?,?,?,?,?,?,?,?)''', (name, number, contractor, 'Старый работодатель', profession, gsp, 'Исходная ГДЛР', active)).lastrowid

    def assignment(self, db, worker_id, shift, foreman_id, crew_id, employer):
        db.execute('''INSERT INTO assignments(work_date,shift,subobject_id,worker_id,employer,foreman_user_id,created_at,crew_id,edit_token)
            VALUES ('2026-09-13',?,?,?,?,?,'now',?,'token')''', (shift, self.site_id, worker_id, employer, foreman_id, crew_id))

    def client(self, user_id):
        client = self.module.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = user_id
            session['csrf_token'] = 'export-csrf'
        return client

    def export(self, client, shift='all', day='2026-09-13'):
        return client.get('/api/staffing/export', query_string={'date': day, 'shift': shift})

    def workbook(self, response):
        return openpyxl.load_workbook(io.BytesIO(response.data), data_only=False)

    def test_all_shifts_exports_actual_history_with_exact_fields(self):
        with self.module.app.app_context():
            before = [tuple(row) for row in self.module.get_db().execute('SELECT * FROM assignments ORDER BY id')]
        response = self.export(self.admin)
        self.assertEqual(response.status_code, 200)
        with self.module.app.app_context():
            after = [tuple(row) for row in self.module.get_db().execute('SELECT * FROM assignments ORDER BY id')]
        self.assertEqual(after, before)
        self.assertEqual(response.headers['X-Export-Count'], '4')
        book = self.workbook(response)
        self.assertEqual(book.sheetnames, ['Список сотрудников', 'Сводная таблица'])
        expected = ['№\nп/п','Группа подобъектов','Подобъект','Компания подрядчик','Организация-работодатель','ФИО работника','Таб. № с префиксом','Должность по штатному расписанию','Профессия ГСП','Категория ГДЛР','ФИО линейного ИТР','ФИО бригадира','Смена','СМУ','Выполняемые операции']
        self.assertEqual([cell.value for cell in book['Список сотрудников'][1]], expected)
        rows = list(book['Список сотрудников'].iter_rows(min_row=2, values_only=False))
        visible = next(row for row in rows if row[5].value == '=Формула')
        self.assertEqual([cell.value for cell in visible], [1, 'Группа А', 'Подобъект 1', 'Подрядчик', '=Работодатель', '=Формула', '000123', 'Должность', 'ГСП', 'Ручная ГДЛР', 'ИТР вручную', 'Бригадир вручную', 'День', None, None])
        self.assertEqual(visible[5].data_type, 's')
        self.assertEqual(visible[4].data_type, 's')
        self.assertEqual(visible[6].number_format, '@')
        self.assertEqual(book['Список сотрудников'].max_row, 5)
        self.assertEqual([row[0].value for row in rows], [1, 2, 3, 4])
        night = next(row for row in rows if row[5].value == 'Ночной сотрудник')
        self.assertEqual(night[12].value, 'Ночь')
        self.assertEqual(book['Список сотрудников'].freeze_panes, 'A2')
        self.assertEqual(book['Список сотрудников'].tables['StaffingSource'].autoFilter.ref, 'A1:O5')

    def test_department_and_operations_follow_assignment_date_and_shift(self):
        department = 'Строительно-монтажный участок № 15.2'
        operations = '=Монтаж трубопровода\nСварка стыков\nКонтроль соединений'
        with self.module.app.app_context():
            db = self.module.get_db()
            db.execute('UPDATE workers SET department=? WHERE id=?', (department, self.visible_worker))
            self.assignment(db, self.visible_worker, '2 смена', self.other_id, self.crew_id, 'Работодатель ночи')
            for worker, day, shift, description in [
                (self.visible_worker, '2026-09-13', '1 смена', operations),
                (self.visible_worker, '2026-09-13', '2 смена', 'Ночная сборка'),
                (self.visible_worker, '2026-09-12', '1 смена', 'Вчерашние операции'),
                (self.legacy_worker, '2026-09-13', '2 смена', 'Ночные операции'),
            ]:
                db.execute('INSERT INTO staffing_performed_work VALUES (?,?,?,?,?,?,?)',
                           (day, worker, shift, description, 'work-token', self.admin_id, 'now'))
            db.commit()
        response = self.export(self.admin)
        self.assertEqual(response.headers['X-Export-Count'], '5')
        sheet = self.workbook(response)['Список сотрудников']
        rows = {(r[6].value, r[12].value): r for r in sheet.iter_rows(min_row=2)}
        day = rows[('000123', 'День')]
        self.assertEqual([c.value for c in day[13:]], [department, operations])
        self.assertEqual(day[14].data_type, 's')
        self.assertTrue(day[14].alignment.wrap_text)
        self.assertGreaterEqual(sheet.row_dimensions[day[0].row].height, 48)
        self.assertEqual(rows[('000123', 'Ночь')][14].value, 'Ночная сборка')
        self.assertEqual(rows[('000124', 'Ночь')][14].value, 'Ночные операции')
        self.assertIsNone(rows[('000125', 'День')][14].value)
        night = self.workbook(self.export(self.admin, '2 смена'))['Список сотрудников']
        self.assertEqual({r[14].value for r in night.iter_rows(min_row=2)}, {'Ночная сборка', 'Ночные операции'})

    def test_foreman_sees_only_own_or_own_crew_assignments(self):
        response = self.export(self.foreman, '1 смена')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['X-Export-Count'], '2')
        book = self.workbook(response)
        self.assertEqual(book.sheetnames, ['Список сотрудников', 'Сводная таблица'])
        self.assertEqual({row[12].value for row in book['Список сотрудников'].iter_rows(min_row=2)}, {'День'})
        values = [row[5].value for row in book['Список сотрудников'].iter_rows(min_row=2)]
        self.assertEqual(values, ['=Формула', 'Неактивный сотрудник'])

    def test_night_filter_exports_two_tabs_with_shift_column(self):
        response = self.export(self.admin, '2 смена')
        self.assertEqual(response.status_code, 200)
        book = self.workbook(response)
        self.assertEqual(book.sheetnames, ['Список сотрудников', 'Сводная таблица'])
        self.assertEqual(book['Список сотрудников'].max_row, 2)
        self.assertEqual(book['Список сотрудников']['F2'].value, 'Ночной сотрудник')
        self.assertEqual(book['Список сотрудников']['M2'].value, 'Ночь')

    def test_permissions_invalid_period_and_empty_export(self):
        self.assertEqual(self.module.app.test_client().get('/api/staffing/export').status_code, 401)
        viewer_export = self.export(self.viewer)
        self.assertEqual(viewer_export.status_code, 200)
        self.assertEqual(viewer_export.headers['X-Export-Count'], '4')
        self.assertEqual(list(self.workbook(viewer_export)['Список сотрудников'].values),
                         list(self.workbook(self.export(self.admin))['Список сотрудников'].values))
        self.assertEqual(self.viewer.put('/api/crews/1/assignments', json={},
                         headers={'X-CSRF-Token': 'export-csrf'}).status_code, 403)
        for query in ({'date': '20260913', 'shift': 'all'}, {'date': '2026-02-30', 'shift': 'all'}, {'date': '2026-09-13', 'shift': 'Ночная смена'}):
            self.assertEqual(self.admin.get('/api/staffing/export', query_string=query).status_code, 400)
        empty = self.export(self.admin, day='2026-09-14')
        self.assertEqual(empty.status_code, 404)
        self.assertIn('нет расставленных', empty.get_json()['error'])


if __name__ == '__main__':
    unittest.main()
