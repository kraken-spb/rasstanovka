import os
import unittest
from io import BytesIO
from uuid import uuid4

from flask import Flask
from werkzeug.exceptions import BadRequest
from openpyxl import load_workbook
from workforce_core import email_value
from workforce_registry_export import registry_workbook, HEADERS


class EmployeeEmailValueTest(unittest.TestCase):
    def test_optional_trimmed_and_invalid_values(self):
        self.assertEqual(email_value(''), '')
        self.assertEqual(email_value(' Worker+site@example.ru '), 'Worker+site@example.ru')
        for value in ('abc', 'a@@b.ru', 'a b@example.ru', 'a@example', 'a@b.ru\nother@b.ru', 'a'*255+'@b.ru', ['a@b.ru']):
            with self.subTest(value=value), Flask(__name__).test_request_context(), self.assertRaises(BadRequest):
                email_value(value)

    def test_export_email_is_text_and_follows_phone(self):
        self.assertEqual(HEADERS.index('E-mail'), HEADERS.index('Телефон')+1)
        book=load_workbook(BytesIO(registry_workbook([{'full_name':'Работник','email':'=test@example.ru'}],'2026-09-19','rotation')))
        cell=book.active.cell(5,HEADERS.index('E-mail')+1)
        self.assertEqual((cell.value,cell.data_type),('=test@example.ru','s'))


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select isolated PostgreSQL database.')
class EmployeeEmailApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import test_workforce_api as fixture
        fixture.WorkforceApiTest.setUpClass()

    def setUp(self):
        import test_workforce_api as fixture
        self.f=fixture.WorkforceApiTest();self.f.setUp();self.addCleanup(self.f.tearDown)
        self.f.add_source_record(self.f.worker,'urp:П15')

    def test_save_clear_stale_audit_and_role_scoped_reads(self):
        f=self.f
        before=f.request('get',f'people/{f.worker}').json['profile']
        data=dict(email='Worker@example.ru',token=before['token'],reason='Обновление контакта',request_key=str(uuid4()))
        saved=f.request('patch',f'people/{f.worker}/profile',data,role='rotation')
        self.assertEqual(saved.status_code,200,saved.json)
        self.assertEqual(saved.json['email'],'Worker@example.ru')
        self.assertEqual(f.request('patch',f'people/{f.worker}/profile',dict(data,email='other@example.ru',request_key=str(uuid4()))).status_code,409)
        audit=f.db.native("SELECT after_json FROM workforce_audit WHERE worker_id=%s AND entity_type='profile' AND action='update' ORDER BY changed_at DESC LIMIT 1",(f.worker,)).fetchone()
        self.assertEqual(audit['after_json']['email'],'Worker@example.ru')
        f.app.config['TESTING']=False
        path='people?section=rotation&date=2026-09-19&department=TEST-SMU'
        self.assertEqual(f.request('get',path).json['rows'][0]['email'],'Worker@example.ru')
        for role in ('foreman','viewer'):
            with self.subTest(role=role):
                self.assertNotIn('email',f.request('get',path,role=role).json['rows'][0])
                self.assertNotIn('email',f.request('get',f'people/{f.worker}',role=role).json['profile'])
                book=load_workbook(BytesIO(f.request('get',path.replace('people?','people/export?'),role=role).data))
                self.assertIsNone(book.active.cell(5,HEADERS.index('E-mail')+1).value)
        cleared=f.request('patch',f'people/{f.worker}/profile',dict(email='',token=saved.json['token'],reason='Удаление контакта',request_key=str(uuid4())))
        self.assertEqual(cleared.status_code,200,cleared.json)
        self.assertEqual(cleared.json['email'],'')

    def test_invalid_edit_and_foreign_worker_are_rejected(self):
        f=self.f
        before=f.request('get',f'people/{f.worker}').json['profile']
        data=dict(email='not-an-email',token=before['token'],reason='Проверка',request_key=str(uuid4()))
        self.assertEqual(f.request('patch',f'people/{f.worker}/profile',data).status_code,400)
        self.assertEqual(f.request('get',f'people/{f.worker}').json['profile']['token'],before['token'])
        data.update(email='worker@example.ru')
        self.assertEqual(f.request('patch',f'people/{f.ids[1]}/profile',data,role='rotation').status_code,404)
