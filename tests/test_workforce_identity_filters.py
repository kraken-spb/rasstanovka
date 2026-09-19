import os
import unittest
from urllib.parse import urlencode

import test_workforce_api as fixtures


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class WorkforceIdentityFiltersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        self.f.app.config['TESTING'] = False
        self.name = 'Иванов Иван ' + self.f.suffix
        self.numbers = ['000' + self.f.suffix, '001' + self.f.suffix]
        for worker, number in zip(self.f.ids, self.numbers):
            self.f.db.native('UPDATE workers SET full_name=%s,personnel_no=%s WHERE id=%s', (self.name, number, worker))
            self.f.add_source_record(worker, 'urp:К15', role='recruitment', sheet='ПВП')

    def get(self, values=(), role='admin', endpoint='people'):
        query = [('date', '2026-09-17'), ('section', 'recruitment'),
                 ('department', 'TEST-SMU'), ('department', 'OTHER-SMU'), *values]
        response = self.f.request('get', endpoint + '?' + urlencode(query), role=role)
        self.assertEqual(response.status_code, 200, response.data[:300])
        return response

    def test_exact_multiselection_homonyms_intersection_and_export(self):
        names = [('full_name', self.name), ('full_name', 'Не существующий сотрудник')]
        self.assertEqual(self.get(names).json['totals']['total'], 2)
        numbers = [('personnel_no', n) for n in self.numbers]
        self.assertEqual(self.get(numbers).json['totals']['total'], 2)
        selected = names + [('personnel_no', self.numbers[0])]
        result = self.get(selected).json
        self.assertEqual([r['id'] for r in result['rows']], [self.f.worker])
        self.assertEqual(result['rows'][0]['personnel_no'], self.numbers[0])
        self.assertEqual(self.get([('personnel_no', self.numbers[0][1:])]).json['totals']['total'], 0)
        self.assertEqual(self.get(names + [('stage', 'stage.onsite')]).json['rows'], [])
        self.assertEqual(self.get(names + [('limit', '1'), ('offset', '1')]).json['totals']['total'], 2)
        self.assertEqual(self.get(names, role='recruitment').json['totals']['total'], 1)
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(self.get(selected, endpoint='people/export').data))
        self.assertEqual(book.active.max_row, 5)
        self.assertIn(self.numbers[0], [c.value for c in book.active[5]])

    def test_options_cover_all_pages_respect_other_filters_and_hide_internal_numbers(self):
        names = self.get([('identity_options', 'full_name'), ('limit', '1'), ('offset', '1')]).json['options']
        self.assertEqual(names, [self.name])
        numbers = self.get([('identity_options', 'personnel_no'), ('limit', '1')]).json['options']
        self.assertEqual(numbers, self.numbers)
        self.assertEqual(self.get([('identity_options', 'personnel_no')], role='recruitment').json['options'], self.numbers[:1])
        self.assertEqual(self.get([('identity_options', 'full_name'), ('personnel_no', self.numbers[1])], role='recruitment').json['options'], [])
        self.f.db.native('UPDATE workers SET personnel_is_internal=TRUE WHERE id=%s', (self.f.ids[1],))
        self.assertEqual(self.get([('identity_options', 'personnel_no')]).json['options'], ['__none__', self.numbers[0]])
        self.assertEqual(self.get([('personnel_no', self.numbers[1])]).json['rows'], [])
        missing = self.get([('personnel_no', '__none__')]).json
        self.assertEqual([r['id'] for r in missing['rows']], [self.f.ids[1]])
        self.assertEqual(missing['rows'][0]['personnel_no'], '')

    def test_invalid_fields_and_excessive_selection_are_rejected(self):
        for values in [[('identity_options', 'password')], [('full_name', 'x' * 501)],
                       [('personnel_no', '\x01')], [('full_name', str(i)) for i in range(101)]]:
            with self.subTest(values=values[:2]):
                self.assertEqual(self.f.request('get', 'people?' + urlencode(values)).status_code, 400)
