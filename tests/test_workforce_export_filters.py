"""Download filters select the displayed reference values without widening access."""
from io import BytesIO
import os
import unittest
from urllib.parse import urlencode
from uuid import uuid4

from openpyxl import load_workbook
import test_workforce_api as fixtures
from workforce_registry_export import HEADERS


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class WorkforceExportFiltersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        f = self.f
        for worker in f.ids:
            f.add_source_record(worker, 'urp:П15')
        self.codes = {kind: f.create_catalog(kind, 'Фильтр ' + kind)['code']
                      for kind in ('citizenship', 'travelpoint', 'profession', 'project')}
        f.db.native('UPDATE workers SET profession_code=%s,profession=%s WHERE id=%s',
                    (self.codes['profession'], 'Профессия для выгрузки', f.worker))
        f.db.native('UPDATE workforce_profiles SET citizenship_code=%s,origin_code=%s WHERE worker_id=%s',
                    (self.codes['citizenship'], self.codes['travelpoint'], f.worker))
        smu = f.db.native('INSERT INTO smu_catalog(name) VALUES (%s) RETURNING id', ('Фильтр ' + f.suffix,)).fetchone()[0]
        f.db.native("INSERT INTO employee_smu(worker_id,smu_id,source_department) VALUES (%s,%s,'TEST-SMU')", (f.worker, smu))
        f.db.native('INSERT INTO workforce_smu_projects(smu_id,project_code) VALUES (%s,%s)', (smu, self.codes['project']))

    def export(self, filters=(), role='admin'):
        return self.f.request('get', 'people/export?' + urlencode([
            ('date', '2026-09-18'), ('section', 'rotation'), ('department', 'TEST-SMU'),
            ('department', 'OTHER-SMU'), *filters]), role=role)

    def count(self, filters=(), role='admin'):
        response = self.export(filters, role)
        self.assertEqual(response.status_code, 200, response.data[:300])
        return int(response.headers['X-Export-Row-Count'])

    def test_reference_filters_combine_and_empty_multi_values_do_not_widen_scope(self):
        filters = [(key, self.codes[kind]) for key, kind in [('project','project'), ('profession','profession'),
                    ('citizenship','citizenship'), ('origin','travelpoint')]]
        for selected in [[pair] for pair in filters] + [filters]:
            with self.subTest(selected=selected):
                self.assertEqual(self.count(selected), 1)
        self.assertEqual(self.count([('project', '__none__')]), 1)
        self.assertEqual(self.count([('project', '__none__'), ('project', self.codes['project'])]), 2)
        self.assertEqual(self.count([('project', '__none__')], 'rotation'), 0)
        self.assertEqual(self.count(filters, 'rotation'), 1)
        self.assertEqual(self.count([('project', self.codes['project']), ('origin', '__none__')]), 0)
        self.assertEqual(self.count([('origin', '__none__')]), 1)
        self.assertEqual(self.export([('origin', self.codes['travelpoint'])], 'foreman').status_code, 403)
        self.assertEqual(self.export([('origin', '__none__')], 'viewer').status_code, 403)

    def movement(self, worker, direction, planned, basis='basis.ticket', result=None):
        f = self.f
        return f.db.native('''INSERT INTO workforce_movements(worker_id,direction,planned_date,basis_code,
            result_code,created_by,updated_by,request_key) VALUES (%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id''',
            (worker,direction,planned,basis,result,f.users['admin']['id'],f.users['admin']['id'],uuid4())).fetchone()[0]

    def test_trip_filters_match_first_displayed_plan_and_exclude_cancelled_and_replaced(self):
        f = self.f
        old = self.movement(f.worker, 'departure', '2026-09-14')
        cancelled = self.movement(f.worker, 'departure', '2026-09-15', result='result.cancelled')
        selected = self.movement(f.worker, 'arrival', '2026-09-20')
        f.db.native('UPDATE workforce_movements SET rescheduled_from=%s WHERE id=%s', (old, selected))
        self.movement(f.worker, 'departure', '2026-10-01', 'basis.schedule')
        self.assertEqual(self.count([('movement_direction','arrival'), ('movement_basis','basis.ticket')]), 1)
        self.assertEqual(self.count([('movement_direction','departure')]), 0)
        self.assertEqual(self.count([('movement_basis','basis.schedule')]), 0)
        self.assertEqual(self.count([('movement_direction','__none__')]), 1)
        response = self.export([('movement_direction','arrival')])
        with BytesIO(response.data) as stream:
            book = load_workbook(stream)
            self.assertEqual(book.active.cell(5, HEADERS.index('Заезд / выезд') + 1).value, 'Заезд')
            book.close()

    def test_schedule_and_existing_filters_match_export_values_and_missing_links(self):
        f = self.f
        actor = f.users['admin']['id']
        schedule = f.db.native('''INSERT INTO workforce_rotation_schedules(name,onsite_days,leave_days,travel_days,updated_by)
            VALUES (%s,30,30,2,%s) RETURNING id''', ('График '+f.suffix, actor)).fetchone()[0]
        f.db.native('''INSERT INTO workforce_rotations(worker_id,schedule_id,schedule_snapshot,start_date,planned_end_date,
            leave_end_date,next_arrival_date,created_by,updated_by,request_key)
            VALUES (%s,%s,'{}','2026-09-01','2026-09-30','2026-10-30','2026-11-01',%s,%s,%s)''',
            (f.worker,schedule,actor,actor,uuid4()))
        self.assertEqual(self.count([('rotation_schedule',str(schedule))]), 1)
        self.assertEqual(self.count([('rotation_schedule','__none__')]), 1)
        self.assertEqual(self.count([('employment','employment.staff')]), 1)
        self.assertEqual(self.count([('employment','__none__')]), 1)
        self.assertEqual(self.count([('category','__none__')]), 2)
        employer = f.db.native('SELECT employer_id FROM workforce_profiles WHERE worker_id=%s',(f.worker,)).fetchone()[0]
        self.assertEqual(self.count([('employer',str(employer))]), 2)
        self.assertEqual(self.count([('employer','__none__')]), 0)

    def test_bad_filters_fail_and_inactive_reference_is_still_selectable(self):
        f = self.f
        f.db.native('UPDATE workforce_catalog SET active=FALSE WHERE code=%s',(self.codes['profession'],))
        self.assertEqual(self.count([('profession',self.codes['profession'])]), 1)
        for values in [[('rotation_schedule','bad')], [('employer','bad')], [('category','-1')],
                       [('movement_direction','other')], [('citizenship','project.bad')],
                       [('profession','profession.'+str(n)) for n in range(101)]]:
            self.assertEqual(self.export(values).status_code, 400)
