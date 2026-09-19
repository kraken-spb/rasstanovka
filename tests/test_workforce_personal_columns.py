import os
import unittest
from io import BytesIO

from openpyxl import load_workbook

import test_workforce_api as fixtures
from workforce_registry_export import HEADERS


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class WorkforcePersonalColumnsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        self.f.add_source_record(self.f.worker, 'urp:П15')
        self.city = self.f.create_catalog('travelpoint', 'Город из справочника')
        self.citizenship = self.f.create_catalog('citizenship', 'Гражданство из справочника')
        self.f.db.native('''UPDATE workforce_profiles SET phone=%s,origin_code=%s,origin_city=%s,citizenship_code=%s
            WHERE worker_id=%s''', ('=+7(999)123-45-67', self.city['code'], 'Старое текстовое значение',
                                    self.citizenship['code'], self.f.worker))

    def row(self, role='admin'):
        response = self.f.request('get', 'people?date=2026-09-16&department=TEST-SMU', role=role)
        self.assertEqual(response.status_code, 200, response.data)
        return next(item for item in response.json['rows'] if item['id'] == self.f.worker)

    def test_private_columns_are_role_scoped_after_a_warm_admin_cache(self):
        f = self.f
        f.app.config['TESTING'] = False
        admin = self.row()
        self.assertEqual((admin['phone'], admin['origin_city'], admin['citizenship']),
                         ('=+7(999)123-45-67', self.city['label'], self.citizenship['label']))
        for role in ('foreman', 'viewer'):
            with self.subTest(role=role):
                private = self.row(role)
                self.assertNotIn('phone', private)
                self.assertNotIn('origin_city', private)
                self.assertEqual(private['citizenship'], self.citizenship['label'])
        # The cached common projection must not be mutated by the restricted response.
        again = self.row()
        self.assertEqual((again['phone'], again['origin_city'], again['citizenship']),
                         ('=+7(999)123-45-67', self.city['label'], self.citizenship['label']))

    def test_catalog_city_wins_over_legacy_and_scope_stays_enforced(self):
        f = self.f
        self.assertEqual(self.row('rotation')['origin_city'], self.city['label'])
        f.db.native("UPDATE workforce_profiles SET origin_code=NULL,origin_city='Старый базовый город' WHERE worker_id=%s",
                    (f.worker,))
        self.assertEqual(self.row('recruitment')['origin_city'], 'Старый базовый город')
        self.assertEqual(f.request('get', f'people/{f.ids[1]}', role='rotation').status_code, 404)
        f.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s", (f.users['rotation']['id'],))
        response = f.request('get', 'people?date=2026-09-16', role='rotation')
        self.assertEqual((response.json['rows'], response.json['totals']['total']), ([], 0))

    def test_export_uses_header_lookup_text_cells_and_removes_private_values_for_foreman(self):
        f = self.f
        path = 'people/export?date=2026-09-16&section=rotation&department=TEST-SMU'
        admin = f.request('get', path)
        self.assertEqual(admin.status_code, 200, admin.data[:200])
        book = load_workbook(BytesIO(admin.data))
        sheet = book.active
        phone = sheet.cell(5, HEADERS.index('Телефон') + 1)
        origin = sheet.cell(5, HEADERS.index('Город отправления') + 1)
        citizenship = sheet.cell(5, HEADERS.index('Гражданство') + 1)
        self.assertEqual((phone.value, phone.data_type, origin.value, citizenship.value),
                         ('=+7(999)123-45-67', 's', self.city['label'], self.citizenship['label']))
        book.close()

        restricted = f.request('get', path, role='foreman')
        self.assertEqual(restricted.status_code, 200, restricted.data[:200])
        book = load_workbook(BytesIO(restricted.data))
        sheet = book.active
        self.assertIsNone(sheet.cell(5, HEADERS.index('Телефон') + 1).value)
        self.assertIsNone(sheet.cell(5, HEADERS.index('Город отправления') + 1).value)
        self.assertEqual(sheet.cell(5, HEADERS.index('Гражданство') + 1).value, self.citizenship['label'])
        book.close()
