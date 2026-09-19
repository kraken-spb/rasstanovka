import json
import os
import unittest
from urllib.parse import urlencode
from io import BytesIO

from flask import Flask
from werkzeug.exceptions import BadRequest

from table_sorting import validate_sort, WORKFORCE_SORT_FIELDS


class SortValidationTest(unittest.TestCase):
    def test_invalid_duplicate_unknown_direction_and_large_sort(self):
        with Flask(__name__).test_request_context():
            for value in [None,{},[{'field':'w.id; DROP TABLE users','direction':'asc'}],
                          [{'field':'name','direction':'sideways'}],
                          [{'field':'name','direction':'asc'}]*2,
                          [{'field':key,'direction':'asc'} for key in list(WORKFORCE_SORT_FIELDS)[:9]],
                          [{'field':'name','direction':'asc','sql':'x'}]]:
                with self.subTest(value=value),self.assertRaises(BadRequest):validate_sort(value,WORKFORCE_SORT_FIELDS)
            self.assertEqual(validate_sort([],WORKFORCE_SORT_FIELDS),[])


@unittest.skipIf(os.getenv('CREW_POSTGRES_TEST_ENV'),'SQLite account checks run in a separate process.')
class SortPreferencesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from tests import test_preferences as fixtures
        fixtures.PreferencesTest.setUpClass()

    @classmethod
    def tearDownClass(cls):
        from tests import test_preferences as fixtures
        fixtures.PreferencesTest.tearDownClass()

    def test_accounts_sections_reload_reset_and_csrf(self):
        from tests import test_preferences as fixtures
        f=fixtures.PreferencesTest();f.setUp();self.addCleanup(f.doCleanups)
        saved={'staffingSort':[{'field':'crew_number','direction':'asc'},{'field':'name','direction':'desc'}],
               'rotationSort':[{'field':'department','direction':'desc'}],
               'recruitmentSort':[{'field':'stage_date','direction':'asc'}]}
        self.assertEqual(f.save(saved).status_code,200)
        self.assertEqual(f.client_for(f.admin).get('/api/preferences/staffing').json['settings'],saved)
        self.assertEqual(f.clients['foreman'].get('/api/preferences/staffing').json['settings'],{})
        self.assertEqual(f.save({'rotationSort':[]}).status_code,200)
        saved['rotationSort']=[]
        self.assertEqual(f.client_for(f.admin).get('/api/preferences/staffing').json['settings'],saved)
        self.assertEqual(f.save({'staffingSort':[]},csrf=False).status_code,403)
        self.assertEqual(f.save({'staffingSort':[{'field':'stage_date','direction':'asc'}]}).status_code,400)
        self.assertEqual(f.save({'rotationSort':[{'field':'name','direction':'asc'}]},role='viewer').status_code,200)
        self.assertEqual(f.clients['viewer'].get('/api/preferences/staffing').json['settings'],{'rotationSort':[{'field':'name','direction':'asc'}]})
        self.assertEqual(f.save({'groupMode':'crew'},role='viewer').status_code,403)
        self.assertEqual(f.save({'rotationSort':[{'field':'phone','direction':'asc'}]},role='viewer').status_code,403)
        self.assertEqual(f.save({'rotationSort':[{'field':'email','direction':'asc'}]},role='viewer').status_code,403)


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'),'Select disposable PostgreSQL database.')
class SortPostgresTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import test_workforce_api as fixtures
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        import test_workforce_api as fixtures
        self.f=fixtures.WorkforceApiTest();self.f.setUp();self.addCleanup(self.f.tearDown)
        self.prefix='Сортировка '+self.f.suffix
        self.ids=[]
        for i,(department,number) in enumerate([('СМУ 2','10'),('СМУ 2','2'),('СМУ 10','11'),('СМУ 10','3')]):
            row=self.f.db.native('INSERT INTO workers(full_name,personnel_no,department) VALUES (%s,%s,%s) RETURNING id',
                (self.prefix+' '+str(i), self.f.suffix+number, department)).fetchone()
            self.ids.append(row['id'])
            self.f.db.native('INSERT INTO workforce_registry_memberships(worker_id,service,created_by) VALUES (%s,%s,%s)',(row['id'],'rotation',self.f.users['admin']['id']))
        self.sort=[dict(field='department',direction='asc'),dict(field='personnel',direction='asc')]

    def get(self,levels=None,**values):
        query={'date':'2026-09-18','q':self.prefix,'sort':json.dumps(self.sort if levels is None else levels),**values}
        response=self.f.request('get','people?'+urlencode(query))
        self.assertEqual(response.status_code,200,response.data)
        return response.json

    def test_whole_selection_before_pages_natural_order_and_cache(self):
        self.f.app.config['TESTING']=False
        first=self.get(limit=2,offset=0)['rows'];second=self.get(limit=2,offset=2)['rows']
        self.assertEqual([r['id'] for r in first+second],[self.ids[1],self.ids[0],self.ids[3],self.ids[2]])
        reversed_sort=[dict(field='department',direction='desc'),dict(field='personnel',direction='desc')]
        self.assertEqual([r['id'] for r in self.get(reversed_sort)['rows']],list(reversed([self.ids[1],self.ids[0],self.ids[3],self.ids[2]])))

    def test_every_displayed_column_supports_sorting_and_nulls_last(self):
        for key in sorted(WORKFORCE_SORT_FIELDS):
            with self.subTest(key=key): self.get([dict(field=key,direction='desc')],limit=1)
        self.f.db.native('UPDATE workforce_profiles SET arrival_date=%s WHERE worker_id=%s',('2026-09-02',self.ids[0]))
        self.assertEqual(self.get([dict(field='arrival_date',direction='desc')])['rows'][0]['id'],self.ids[0])
        self.assertEqual(self.get([dict(field='arrival_date',direction='asc')])['rows'][0]['id'],self.ids[0])

    def test_cached_order_invalidates_on_edits_and_scope_changes(self):
        f=self.f;f.app.config['TESTING']=False
        self.assertEqual(self.get()['rows'][0]['id'],self.ids[1])
        f.db.native("UPDATE workers SET department='А' WHERE id=%s",(self.ids[2],))
        self.assertEqual(self.get()['rows'][0]['id'],self.ids[2])
        f.db.native("INSERT INTO user_smu_access(user_id,mode,departments_json,edit_token,updated_by,updated_at) VALUES (%s,'selected','[]','test',%s,'now')",(f.users['admin']['id'],f.users['admin']['id']))
        self.assertEqual(self.get()['rows'],[])

    def test_persisted_default_isolated_sections_export_and_private_fields(self):
        f=self.f
        f.db.native("INSERT INTO user_staffing_preferences(user_id,settings_json,updated_at) VALUES (%s,%s,'now')",(f.users['admin']['id'],json.dumps({'rotationSort':self.sort,'recruitmentSort':[dict(field='name',direction='desc')]})))
        q=urlencode(dict(date='2026-09-18',q=self.prefix,section='rotation'))
        response=f.request('get','people?'+q)
        self.assertEqual(response.status_code,200,response.data)
        ids=[r['id'] for r in response.json['rows']]
        self.assertEqual(ids,[self.ids[1],self.ids[0],self.ids[3],self.ids[2]])
        response=f.request('get','people/export?'+q)
        self.assertEqual(response.status_code,200,response.data)
        from openpyxl import load_workbook
        sheet=load_workbook(BytesIO(response.data)).worksheets[0]
        from workforce_registry_export import HEADERS
        name_column=HEADERS.index('ФИО')
        names=[row[name_column] for row in sheet.iter_rows(values_only=True) if isinstance(row[name_column],str) and row[name_column].startswith(self.prefix)]
        self.assertEqual(names,[self.prefix+' '+str(i) for i in [1,0,3,2]])
        bad=urlencode(dict(date='2026-09-18',sort=json.dumps([dict(field='phone',direction='asc')])))
        self.assertEqual(f.request('get','people?'+bad,role='viewer').status_code,403)
        bad=urlencode(dict(date='2026-09-18',sort=json.dumps([dict(field='email',direction='asc')])))
        self.assertEqual(f.request('get','people?'+bad,role='viewer').status_code,403)
