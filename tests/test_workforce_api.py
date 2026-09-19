"""Transactional integration tests use only an explicitly selected staging database."""
import os
from pathlib import Path
import unittest
from uuid import uuid4

from flask import Flask, g, jsonify, request
from werkzeug.exceptions import HTTPException


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select isolated PostgreSQL staging explicitly.')
class WorkforceApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from tools.migrate_sqlite_to_postgres import load_environment
        load_environment(Path(os.environ['CREW_POSTGRES_TEST_ENV']))
        if os.environ.get('APP_ENVIRONMENT') != 'staging':
            raise RuntimeError('Refusing workforce integration tests outside staging.')

    def setUp(self):
        from postgres_db import PostgresConnection
        from workforce_api import register_workforce_routes
        class RollbackConnection(PostgresConnection):
            def execute(self, query, parameters=()):
                if query.upper() in ('BEGIN', 'BEGIN IMMEDIATE') and self.in_transaction:
                    return self.native('SELECT 1')
                return super().execute(query, parameters)

            def commit(self):
                pass

            def __exit__(self, *_):
                return False

        self.db = RollbackConnection()
        self.db.execute('BEGIN IMMEDIATE')
        self.suffix = uuid4().hex
        self.users = {}
        for role in ('admin', 'foreman', 'rotation', 'recruitment', 'hr_viewer', 'viewer'):
            row = self.db.native('''INSERT INTO users(username,password_hash,full_name,role,created_at)
                VALUES (%s,'unusable','Проверка прав',%s,'2026-09-16') RETURNING id,role,active''',
                (self.suffix + role, role)).fetchone()
            self.users[role] = dict(row)
            if role in ('foreman', 'rotation', 'recruitment'):
                self.db.native('''INSERT INTO user_smu_access(user_id,mode,departments_json,edit_token,updated_by,updated_at)
                    VALUES (%s,'selected',%s,%s,%s,'2026-09-16')''',
                    (row['id'], '["TEST-SMU"]', self.suffix, row['id']))
        self.ids = []
        for department in ('TEST-SMU', 'OTHER-SMU'):
            row = self.db.native('''INSERT INTO workers(full_name,personnel_no,employer,department)
                VALUES ('Тестовый Сотрудник',%s,'Тестовый работодатель',%s) RETURNING id''',
                (self.suffix + department, department)).fetchone()
            self.ids.append(row['id'])
        self.worker = self.ids[0]
        self.db.native("UPDATE workforce_profiles SET employment_code='employment.staff' WHERE worker_id=%s", (self.worker,))
        self.app = Flask(__name__)
        self.app.config['TESTING'] = True

        @self.app.before_request
        def actor():
            g.user = self.users[request.headers.get('Test-Role', 'admin')]

        @self.app.errorhandler(HTTPException)
        def failure(error):
            return jsonify(error=error.description), error.code

        # Authentication/CSRF are covered through the real app separately. The domain
        # authorizer is exercised even with a permissive outer decorator.
        register_workforce_routes(self.app, lambda: self.db, lambda *roles: lambda fn: fn)
        self.client = self.app.test_client()

    def tearDown(self):
        from workforce_read_cache import _lock, _values
        with _lock:
            _values.clear()
        self.db.rollback()
        self.db.close()

    def test_registry_cache_invalidation_and_fresh_scope(self):
        from unittest.mock import patch
        import workforce_read_cache
        self.app.config['TESTING'] = False
        path = 'people?department=TEST-SMU&date=2026-09-16'
        first = self.request('get', path, role='foreman').get_json()
        self.assertEqual([row['id'] for row in first['rows']], [self.worker])
        with patch.object(workforce_read_cache, 'cached', wraps=workforce_read_cache.cached) as cache:
            second = self.request('get', path, role='foreman').get_json()
            self.assertEqual(first, second)
            self.assertEqual(cache.call_count, 2)
        self.db.native('UPDATE workers SET full_name=%s WHERE id=%s', ('Изменённое ФИО', self.worker))
        changed = self.request('get', path, role='foreman').get_json()
        self.assertEqual(changed['rows'][0]['full_name'], 'Изменённое ФИО')
        self.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s", (self.users['foreman']['id'],))
        revoked = self.request('get', path, role='foreman').get_json()
        self.assertEqual(revoked['rows'], [])
        self.assertEqual(revoked['totals']['total'], 0)

    def add_source_record(self, worker_id, source_key, role='rotation', sheet='Явка', active=True):
        # A reviewed initial batch can contain both services; membership belongs to each row.
        if not hasattr(self, 'source_batch'):
            self.source_batch = self.db.native('''INSERT INTO workforce_import_batches
                (request_key,file_sha256,filename,service,source_key,report_date,state,preview_json,created_by)
                VALUES (%s,%s,'Согласованные источники.xlsx','reviewed',%s,'2026-09-16','applied','{}',%s)
                RETURNING id''', (uuid4(), self.suffix * 2, 'test:' + self.suffix, self.users['admin']['id'])).fetchone()['id']
            self.source_row = 0
        self.source_row += 1
        return self.db.native('''INSERT INTO workforce_source_records
            (batch_id,worker_id,filename,sheet,source_row,source_role,raw_json,mapped_json,source_key,active)
            VALUES (%s,%s,'Исходник.xlsx',%s,%s,%s,'{}','{}',%s,%s) RETURNING id''',
            (self.source_batch, worker_id, sheet, self.source_row, role, source_key, active)).fetchone()['id']

    def test_registry_sections_use_row_sources_without_duplicate_people(self):
        self.add_source_record(self.worker, 'urp:П15')
        self.add_source_record(self.worker, 'urp:П19', sheet='Неявка')
        self.add_source_record(self.worker, 'urp:К19', role='recruitment', sheet='Патенты ПВП')
        # Current employment/stage and a familiar sheet title do not establish source ownership.
        self.db.native("UPDATE workforce_profiles SET employment_code='employment.recruitment' WHERE worker_id=%s", (self.worker,))
        self.add_source_record(self.ids[1], 'unrecognized:' + self.suffix, sheet='Явка')
        self.add_source_record(self.ids[1], 'urp:К15', role='recruitment', sheet='ПВП', active=False)
        for section in ('rotation', 'recruitment'):
            with self.subTest(section=section):
                response = self.request('get', 'people?date=2026-09-16&section=' + section + '&q=Тестовый Сотрудник')
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual([row['id'] for row in response.json['rows']], [self.worker])
                self.assertEqual(response.json['totals']['total'], 1)

    def test_registry_sections_keep_source_less_manual_people_in_compatible_list(self):
        self.db.native('''INSERT INTO manual_employees(worker_id,created_by,created_at,request_key,payload_hash)
            VALUES (%s,%s,'2026-09-16',%s,'test')''', (self.worker, self.users['admin']['id'], uuid4().hex))
        path = 'people?date=2026-09-16&department=TEST-SMU'
        compatible = self.request('get', path).json
        self.assertEqual([row['id'] for row in compatible['rows']], [self.worker])
        self.assertEqual(self.request('get', path + '&section=').json, compatible)
        for section in ('rotation', 'recruitment'):
            self.assertEqual(self.request('get', path + '&section=' + section).json['totals']['total'], 0)
        self.add_source_record(self.worker, 'urp:П15')
        self.assertEqual(self.request('get', path + '&section=rotation').json['totals']['total'], 1)
        self.assertEqual(self.request('get', path + '&section=recruitment').json['totals']['total'], 0)

    def test_registry_section_pagination_totals_filters_scope_and_outstaff(self):
        outstaff = self.db.native('''INSERT INTO workers(full_name,personnel_no,department)
            VALUES ('Аутстафф Проверка',%s,'TEST-SMU') RETURNING id''', (self.suffix + 'outstaff',)).fetchone()['id']
        inactive = self.db.native('''INSERT INTO workers(full_name,personnel_no,department,active)
            VALUES ('Неактивный Проверка',%s,'TEST-SMU',0) RETURNING id''', (self.suffix + 'inactive',)).fetchone()['id']
        self.add_source_record(self.worker, 'urp:П15')
        self.add_source_record(self.ids[1], 'urp:П19')
        self.add_source_record(outstaff, 'urp:П19', role='outstaff', sheet='Аутстаффинг')
        self.add_source_record(outstaff, 'urp:П19', role='outstaff', sheet='Аустаффинг')
        self.add_source_record(inactive, 'urp:П15')
        for worker, stage in ((self.worker, 'stage.onsite'), (outstaff, 'stage.pvp')):
            self.db.native('''INSERT INTO workforce_stage_events
                (worker_id,stage_code,effective_date,confirmed,reason,created_by,request_key)
                VALUES (%s,%s,'2026-09-16',TRUE,'Проверка фильтра',%s,%s)''',
                (worker, stage, self.users['admin']['id'], uuid4()))
        path = 'people?date=2026-09-16&section=rotation&limit=1'
        pages = [self.request('get', path + '&offset=' + str(offset), role='rotation').json for offset in (0, 1, 2)]
        self.assertEqual({row['id'] for page in pages for row in page['rows']}, {self.worker, outstaff})
        self.assertEqual([len(page['rows']) for page in pages], [1, 1, 0])
        expected = {'total': 2, 'onsite': 1, 'pvp': 1, 'inbound': 0, 'outbound': 0, 'on_leave': 0, 'unconfirmed': 0}
        for page in pages:
            self.assertEqual(page['totals'], expected)
        self.assertEqual(self.request('get', path + '&department=OTHER-SMU', role='rotation').json['totals']['total'], 0)
        self.assertEqual(self.request('get', path + '&department=OTHER-SMU').json['totals']['total'], 1)
        selected = self.request('get', path + '&stage=stage.pvp&q=^Аутстафф&regex=1', role='rotation').json
        self.assertEqual([row['id'] for row in selected['rows']], [outstaff])
        self.assertEqual(selected['totals']['pvp'], 1)
        self.assertEqual(self.request('get', path + '&active=0', role='rotation').json['totals']['total'], 1)

    def test_registry_section_cache_tracks_membership_only_changes(self):
        self.app.config['TESTING'] = False
        rotation = self.add_source_record(self.worker, 'urp:П15')
        self.add_source_record(self.worker, 'urp:К15', role='recruitment', sheet='ПВП')
        path = 'people?date=2026-09-16&department=TEST-SMU&section='
        first = self.request('get', path + 'rotation').json
        self.assertEqual(first['totals']['total'], 1)
        self.assertEqual(self.request('get', path + 'recruitment').json['totals']['total'], 1)
        self.db.native('UPDATE workforce_source_records SET active=FALSE WHERE id=%s', (rotation,))
        removed = self.request('get', path + 'rotation').json
        self.assertEqual(removed['rows'], [])
        self.assertEqual(removed['totals']['total'], 0)
        self.assertNotEqual(removed['revision'], first['revision'])
        self.assertEqual(self.request('get', path + 'recruitment').json['totals']['total'], 1)
        replacement = self.add_source_record(self.worker, 'urp:П19')
        self.assertEqual(self.request('get', path + 'rotation').json['totals']['total'], 1)
        self.db.native('DELETE FROM workforce_source_records WHERE id=%s', (replacement,))
        self.assertEqual(self.request('get', path + 'rotation').json['totals']['total'], 0)

    def test_registry_rejects_unknown_section(self):
        response = self.request('get', 'people?section=staff')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Неизвестный раздел', response.json['error'])

    def add_stage_record(self, worker_id, stage, day='2026-09-16', confirmed=True, retracted=False):
        return self.db.native('''INSERT INTO workforce_stage_events
            (worker_id,stage_code,effective_date,confirmed,retracted,reason,created_by,request_key)
            VALUES (%s,%s,%s,%s,%s,'Проверка фильтра состояний',%s,%s) RETURNING id''',
            (worker_id, stage, day, confirmed, retracted, self.users['admin']['id'], uuid4())).fetchone()['id']

    def test_registry_stage_multiselect_combines_unknown_and_resets_with_cache(self):
        self.app.config['TESTING'] = False
        unknown = self.db.native('''INSERT INTO workers(full_name,personnel_no,department)
            VALUES ('Состояние не подтверждено',%s,'TEST-SMU') RETURNING id''', (self.suffix + 'unknown',)).fetchone()['id']
        leave = self.db.native('''INSERT INTO workers(full_name,personnel_no,department)
            VALUES ('Межвахтовый отпуск',%s,'TEST-SMU') RETURNING id''', (self.suffix + 'leave',)).fetchone()['id']
        self.add_stage_record(self.worker, 'stage.onsite')
        self.add_stage_record(self.ids[1], 'stage.pvp')
        self.add_stage_record(leave, 'stage.leave')
        path = 'people?date=2026-09-16&q=' + self.suffix
        cases = [
            ('&stage=stage.onsite', {self.worker}),
            ('&stage=stage.onsite&stage=stage.pvp', set(self.ids)),
            ('&stage=unconfirmed&stage=stage.onsite', {self.worker, unknown}),
            ('&stage=unconfirmed', {unknown}),
            ('&stage=stage.pvp&stage=stage.onsite&stage=stage.pvp', set(self.ids)),
            ('&stage=&stage=stage.onsite', {self.worker}),
            ('&stage=&stage=', {*self.ids, unknown, leave}),
            ('', {*self.ids, unknown, leave}),
        ]
        for query, expected in cases:
            with self.subTest(query=query):
                response = self.request('get', path + query)
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual({row['id'] for row in response.json['rows']}, expected)
                self.assertEqual(len(response.json['rows']), len(expected))
                self.assertEqual(response.json['totals']['total'], len(expected))
        mixed = self.request('get', path + '&stage=stage.onsite&stage=unconfirmed').json
        self.assertEqual(mixed['totals']['onsite'], 1)
        self.assertEqual(mixed['totals']['unconfirmed'], 1)

    def test_registry_stage_multiselect_uses_latest_confirmed_state_on_report_date(self):
        self.add_stage_record(self.worker, 'stage.leave', '2026-09-14')
        self.add_stage_record(self.worker, 'stage.pvp', '2026-09-15')
        self.add_stage_record(self.worker, 'stage.inbound', '2026-09-15')
        self.add_stage_record(self.worker, 'stage.onsite', '2026-09-16', confirmed=False)
        self.add_stage_record(self.worker, 'stage.pvp', '2026-09-16', retracted=True)
        self.add_stage_record(self.worker, 'stage.onsite', '2026-09-17')
        for day, query, expected in [
            ('2026-09-13', '&stage=unconfirmed&stage=stage.onsite', {self.worker}),
            ('2026-09-14', '&stage=stage.leave&stage=stage.inbound', {self.worker}),
            ('2026-09-16', '&stage=stage.leave&stage=stage.inbound', {self.worker}),
            ('2026-09-16', '&stage=stage.pvp&stage=stage.onsite', set()),
            ('2026-09-16', '&stage=unconfirmed', set()),
            ('2026-09-17', '&stage=stage.pvp&stage=stage.onsite', {self.worker}),
            ('2026-09-17', '&stage=stage.inbound', set()),
        ]:
            with self.subTest(day=day, query=query):
                response = self.request('get', f'people?date={day}&department=TEST-SMU' + query, role='rotation')
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual({row['id'] for row in response.json['rows']}, expected)
                self.assertEqual(response.json['totals']['total'], len(expected))

    def test_registry_stage_multiselect_keeps_section_scope_totals_and_pagination(self):
        onsite = self.db.native('''INSERT INTO workers(full_name,personnel_no,department)
            VALUES ('Явка обоих источников',%s,'TEST-SMU') RETURNING id''', (self.suffix + 'onsite',)).fetchone()['id']
        pvp = self.db.native('''INSERT INTO workers(full_name,personnel_no,department)
            VALUES ('ПВП комплектации',%s,'TEST-SMU') RETURNING id''', (self.suffix + 'pvp',)).fetchone()['id']
        for worker in (self.worker, self.ids[1], onsite):
            self.add_source_record(worker, 'urp:П15')
        for worker in (onsite, pvp):
            self.add_source_record(worker, 'urp:К19', role='recruitment', sheet='ПВП')
        self.add_stage_record(onsite, 'stage.onsite')
        self.add_stage_record(self.ids[1], 'stage.onsite')
        self.add_stage_record(pvp, 'stage.pvp')
        path = 'people?date=2026-09-16&q=' + self.suffix + '&stage=stage.onsite&stage=unconfirmed'
        pages = [self.request('get', path + f'&section=rotation&limit=1&offset={offset}', role='rotation').json
                 for offset in (0, 1, 2)]
        self.assertEqual([len(page['rows']) for page in pages], [1, 1, 0])
        self.assertEqual({row['id'] for page in pages for row in page['rows']}, {self.worker, onsite})
        for page in pages:
            self.assertEqual(page['totals'], {'total': 2, 'onsite': 1, 'pvp': 0, 'inbound': 0, 'outbound': 0, 'on_leave': 0, 'unconfirmed': 1})
        self.assertEqual(self.request('get', path + '&section=rotation').json['totals']['total'], 3)
        self.assertEqual(self.request('get', path + '&section=rotation&department=OTHER-SMU', role='rotation').json['totals']['total'], 0)
        recruitment = self.request('get', path + '&section=recruitment', role='recruitment').json
        self.assertEqual(recruitment['rows'], [])
        self.assertEqual(recruitment['totals']['total'], 0)

    def test_registry_stage_multiselect_rejects_unknown_codes(self):
        for query in ('stage=unknown', 'stage=stage.onsite&stage=unknown', 'stage=unconfirmed&stage=employment.staff',
                      'stage=stage.onsite,stage.pvp'):
            with self.subTest(query=query):
                response = self.request('get', 'people?' + query)
                self.assertEqual(response.status_code, 400)
                self.assertIn('состояния', response.json['error'])

    def test_postgres_dashboard_colors_are_catalog_backed_and_scoped(self):
        from personnel_dashboard import register_personnel_dashboard
        from gdlr_api import register_gdlr_routes
        register_personnel_dashboard(self.app, lambda:self.db, lambda *roles:lambda fn:fn)
        register_gdlr_routes(self.app, lambda:self.db, lambda *roles:lambda fn:fn, lambda:'2026-09-17T00:00:00Z')
        self.db.native("UPDATE users SET role='super_admin' WHERE id=%s", (self.users['admin']['id'],))
        made = self.client.post('/api/gdlr-categories', json={'name':'Тест цвета '+self.suffix, 'color':'#17aa88'})
        self.assertEqual(made.status_code, 201, made.data)
        category_id = made.json['id']
        self.db.native('''INSERT INTO employee_gdlr(worker_id,category_id,edit_token,updated_by,updated_at)
            VALUES (%s,%s,'test',%s,'now') ON CONFLICT(worker_id) DO UPDATE SET category_id=excluded.category_id''',
            (self.worker, category_id, self.users['admin']['id']))
        site = self.db.native('SELECT id FROM subobjects LIMIT 1').fetchone()[0]
        for worker_id in self.ids:
            for shift in ('1 смена','2 смена'):
                self.db.native('''INSERT INTO assignments(work_date,shift,subobject_id,worker_id,foreman_user_id,created_at)
                    VALUES ('2026-09-16',%s,%s,%s,%s,'now')''', (shift,site,worker_id,self.users['foreman']['id']))
        path = '/api/personnel-dashboard?start=2026-09-16&end=2026-09-16'
        result = self.client.get(path, headers={'Test-Role':'foreman'})
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.json['counts'], [1])
        self.assertEqual(result.json['category_series'], [{'id':str(category_id),'name':'Тест цвета '+self.suffix,
            'color':'#17AA88','counts':[1]}])
        category = next(row for row in self.client.get('/api/gdlr-categories').json['rows'] if row['id']==category_id)
        patch = {'name':category['name'],'active':True,'color':'#dd2299','expected_token':category['edit_token']}
        self.assertEqual(self.client.patch('/api/gdlr-categories/'+str(category_id),json=patch).status_code, 200)
        self.assertEqual(self.client.get(path,headers={'Test-Role':'foreman'}).json['category_series'][0]['color'], '#DD2299')
        self.assertEqual(self.client.patch('/api/gdlr-categories/'+str(category_id),json=patch).status_code, 409)

    def test_staffing_source_has_one_row_and_requires_ready_base(self):
        from staffing_import import postgres_member_source_sql
        source = postgres_member_source_sql()
        self.db.native('UPDATE workforce_profiles SET workforce_managed=true,staffing_ready=true WHERE worker_id=%s', (self.worker,))
        rows = self.db.native('SELECT w.id,sm.source_row FROM ' + source + ' WHERE w.id=%s', (self.worker,)).fetchall()
        self.assertEqual([row['id'] for row in rows], [self.worker])
        self.db.native('UPDATE workforce_profiles SET staffing_ready=false WHERE worker_id=%s', (self.worker,))
        self.assertEqual(self.db.native('SELECT w.id FROM ' + source + ' WHERE w.id=%s', (self.worker,)).fetchall(), [])

    def test_profile_master_correction_preserves_unknown_schedule_and_identity(self):
        self.db.native("INSERT INTO workforce_catalog(code,kind,label) VALUES (%s,'profession','Новая должность')", ('profession.'+self.suffix,))
        self.db.native("UPDATE workforce_profiles SET rotation_schedule='Вахта 1 — длительность неизвестна' WHERE worker_id=%s", (self.worker,))
        before = self.request('get', f'people/{self.worker}', role='rotation').get_json()['profile']
        body = {'full_name':'Уточнённое ФИО', 'profession':'Новая должность', 'rotation_schedule_id':'',
                'token':before['token'], 'reason':'Сверено с первичным документом', 'request_key':str(uuid4())}
        self.assertEqual(self.request('patch', f'people/{self.worker}/profile', body, role='foreman').status_code, 403)
        response = self.request('patch', f'people/{self.worker}/profile', body, role='rotation')
        self.assertEqual(response.status_code, 200, response.data)
        after = response.get_json()
        self.assertEqual(after['id'], before['id'])
        self.assertEqual(after['uuid'], before['uuid'])
        self.assertEqual(after['full_name'], 'Уточнённое ФИО')
        self.assertEqual(after['rotation_schedule'], before['rotation_schedule'])

    def create_catalog(self, kind, label):
        response = self.request('post', 'catalog/'+kind,
            {'label':label+' '+self.suffix,'reason':'Проверка справочника','request_key':str(uuid4())})
        self.assertEqual(response.status_code, 201, response.data)
        return response.json

    def test_profile_catalog_links_preserve_text_roles_and_disabled_history(self):
        profession = self.create_catalog('profession', 'Должность')
        point = self.create_catalog('travelpoint', 'Пункт отправления')
        path = f'people/{self.worker}/profile'
        before = self.request('get', f'people/{self.worker}').json['profile']
        body = {'profession_code':profession['code'],'origin_code':point['code'],
            'token':before['token'],'reason':'Сверка','request_key':str(uuid4())}
        self.assertEqual(self.request('patch', path, body, 'foreman').status_code, 403)
        saved = self.request('patch', path, body, 'rotation')
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(saved.json['profession'], profession['label'])
        self.assertEqual(saved.json['origin_city'], point['label'])
        self.assertEqual(saved.json['profession_code'], profession['code'])
        self.assertEqual(self.request('patch', path, body, 'rotation').status_code, 200)  # same request replay
        body['request_key'] = str(uuid4())
        self.assertEqual(self.request('patch', path, body, 'rotation').status_code, 409)
        renamed = self.request('patch', 'catalog/profession/'+profession['code'],
            {'label':'Переименована '+self.suffix,'active':False,'token':profession['edit_token'],
             'reason':'Архивирование должности','request_key':str(uuid4())})
        self.assertEqual(renamed.status_code, 200, renamed.data)
        current = self.request('get', f'people/{self.worker}').json['profile']
        self.assertEqual(current['profession_code'], profession['code'])
        self.assertEqual(current['profession'], renamed.json['label'])
        self.assertNotEqual(current['token'], saved.json['token'])
        change = {'profession_code':profession['code'],'phone':'123','token':current['token'],
            'reason':'Контакт','request_key':str(uuid4())}
        result = self.request('patch', path, change, 'rotation')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.json['profession_code'], profession['code'])
        other = self.request('get', f'people/{self.ids[1]}').json['profile']
        change.update(token=other['token'],request_key=str(uuid4()))
        self.assertEqual(self.request('patch', f'people/{self.ids[1]}/profile', change).status_code, 400)
        self.db.native('UPDATE workers SET profession=%s WHERE id=%s', ('Неизвестная из Excel '+self.suffix,self.worker))
        unknown = self.request('get', f'people/{self.worker}').json['profile']
        self.assertIsNone(unknown['profession_code'])
        update = {'notes':'Комментарий','token':unknown['token'],'reason':'Проверка','request_key':str(uuid4())}
        preserved = self.request('patch', path, update, 'rotation')
        self.assertEqual(preserved.status_code, 200, preserved.data)
        self.assertEqual(preserved.json['profession'], unknown['profession'])
        update.update(profession='Произвольный новый текст',token=preserved.json['token'],request_key=str(uuid4()))
        self.assertEqual(self.request('patch', path, update, 'rotation').status_code, 400)
        update.pop('profession')
        update.update(profession_code=None,request_key=str(uuid4()))
        cleared = self.request('patch', path, update, 'rotation')
        self.assertEqual(cleared.status_code, 200, cleared.data)
        self.assertEqual(cleared.json['profession'], '')

    def test_movement_catalog_links_survive_rename_archive_and_reschedule(self):
        point = self.create_catalog('travelpoint','Пункт поездки')
        created = self.request('post', f'people/{self.worker}/movement', self.movement(direction='arrival',
            actual_date=None, planned_date='2026-10-01', result_code=None, destination_kind='pvp',
            origin_code=point['code'], destination_code=point['code']), 'rotation')
        self.assertEqual(created.status_code, 201, created.data)
        row = created.json
        self.assertEqual(row['direction_code'], 'direction.arrival')
        self.assertEqual(row['destination_kind_code'], 'destination.pvp')
        self.assertEqual(row['origin'], point['label'])
        self.assertEqual(row['destination'], point['label'])
        renamed = self.request('patch','catalog/travelpoint/'+point['code'],
            {'label':'Новое название '+self.suffix,'active':False,'token':point['edit_token'],
             'reason':'Архивирование','request_key':str(uuid4())})
        self.assertEqual(renamed.status_code, 200, renamed.data)
        current = self.db.native('SELECT * FROM workforce_movements WHERE id=%s',(row['id'],)).fetchone()
        self.assertEqual(current['origin_code'], point['code'])
        self.assertEqual(current['destination_code'], point['code'])
        patch = {'origin_code':point['code'],'destination_code':point['code'],'notes':'Дополнение',
            'token':str(current['edit_token']),'reason':'Дополнение','request_key':str(uuid4())}
        updated = self.request('patch', f'people/{self.worker}/movement/{row["id"]}', patch, 'rotation')
        self.assertEqual(updated.status_code, 200, updated.data)
        moved = self.request('post', f'people/{self.worker}/movements/{row["id"]}/reschedule',
            {'planned_date':'2026-10-03','token':updated.json['edit_token'],'reason':'Перенос','request_key':str(uuid4())}, 'rotation')
        self.assertEqual(moved.status_code, 201, moved.data)
        self.assertEqual(moved.json['current']['origin_code'], point['code'])
        self.assertEqual(moved.json['current']['destination_code'], point['code'])
        rejected = self.request('post', f'people/{self.worker}/movement', self.movement(origin_code=point['code']), 'rotation')
        self.assertEqual(rejected.status_code, 400)
        ref = self.request('get','reference').json['catalog']
        self.assertEqual({r['code'] for r in ref if r['kind']=='direction'}, {'direction.arrival','direction.departure'})
        self.assertEqual(self.request('post','catalog/direction',{'label':'Произвольное','request_key':str(uuid4()),'reason':'Проверка'}).status_code,400)

    def test_disabled_pvp_reference_can_be_retained_but_not_newly_assigned(self):
        place = self.create_catalog('place','ПВП')
        body = {'place_id':place['id'],'planned_arrival':'2026-10-01','reason':'План','request_key':str(uuid4())}
        created = self.request('post',f'people/{self.worker}/pvp',body,'recruitment')
        self.assertEqual(created.status_code,201,created.data)
        self.db.native('UPDATE workforce_pvp_places SET active=false WHERE id=%s',(place['id'],))
        row = created.json
        saved = self.request('patch',f'people/{self.worker}/pvp/{row["id"]}',
            {'place_id':place['id'],'notes':'Примечание к архивной записи','token':row['edit_token'],
             'reason':'Дополнение','request_key':str(uuid4())},'recruitment')
        self.assertEqual(saved.status_code,200,saved.data)
        self.assertEqual(saved.json['place_id'],place['id'])
        body['request_key']=str(uuid4())
        self.assertEqual(self.request('post',f'people/{self.worker}/pvp',body,'recruitment').status_code,400)

    def test_catalog_exact_binding_does_not_create_or_guess_import_values(self):
        label = 'Пункт импорта '+self.suffix
        before = self.db.native("SELECT count(*) FROM workforce_catalog WHERE kind='travelpoint'").fetchone()[0]
        self.db.native('UPDATE workforce_profiles SET origin_city=%s WHERE worker_id=%s',(label,self.worker))
        current = self.request('get',f'people/{self.worker}').json['profile']
        self.assertIsNone(current['origin_code'])
        self.assertEqual(self.db.native("SELECT count(*) FROM workforce_catalog WHERE kind='travelpoint'").fetchone()[0],before)
        made = self.request('post','catalog/travelpoint',{'label':label,'reason':'Подтверждённое значение','request_key':str(uuid4())})
        self.assertEqual(made.status_code,201,made.data)
        self.assertEqual(self.request('get',f'people/{self.worker}').json['profile']['origin_code'],made.json['code'])
        self.db.native('UPDATE workforce_profiles SET origin_city=%s WHERE worker_id=%s',(label+' другой',self.worker))
        self.assertIsNone(self.request('get',f'people/{self.worker}').json['profile']['origin_code'])
        for role in ('foreman','rotation','recruitment','hr_viewer'):
            self.assertEqual(self.request('post','catalog/travelpoint',{'label':'Запрещено','reason':'Проверка','request_key':str(uuid4())},role).status_code,403)

    def test_background_report_owned_scope_and_download(self):
        from tempfile import TemporaryDirectory
        from unittest.mock import patch
        from workforce_jobs import process_one
        key = str(uuid4())
        body = {'date':'2026-09-16','request_key':key}
        created = self.request('post', 'export-jobs', body, role='foreman')
        self.assertEqual(created.status_code, 202, created.data)
        self.assertEqual(self.request('post', 'export-jobs', body, role='foreman').status_code, 200)
        self.assertEqual(self.request('get', f'export-jobs/{key}', role='rotation').status_code, 404)
        with TemporaryDirectory() as directory, patch.dict(os.environ, {'WORKFORCE_EXPORT_DIR':directory}):
            self.assertTrue(process_one(self.app, lambda:self.db, key))
            ready = self.request('get', f'export-jobs/{key}', role='foreman')
            self.assertEqual(ready.get_json()['state'], 'ready', ready.data)
            download = self.request('get', f'export-jobs/{key}/download', role='foreman')
            self.assertEqual(download.status_code, 200)
            self.assertTrue(download.data.startswith(b'PK'))
            download.close()
            self.db.native("UPDATE workers SET department='OTHER-SMU' WHERE id=%s", (self.worker,))
            self.assertEqual(self.request('get', f'export-jobs/{key}/download', role='foreman').status_code, 403)

    def request(self, method, path, body=None, role='admin'):
        return getattr(self.client, method)('/api/workforce/' + path, json=body, headers={'Test-Role': role})

    def movement(self, **changes):
        return {'direction': 'departure', 'actual_date': '2026-09-16', 'result_code': 'result.happened',
                'reason': 'Подтверждено перевахтой', 'request_key': str(uuid4()), **changes}

    def board_row(self, worker_id=None, role='admin', day='2026-09-16'):
        worker_id = worker_id or self.worker
        result = self.request('get', f'people?date={day}&q={self.suffix}', role=role)
        self.assertEqual(result.status_code, 200, result.data)
        return next(row for row in result.json['rows'] if row['id'] == worker_id)

    def board_move(self, stage, people=None, **changes):
        row = self.board_row()
        return {'date': '2026-09-16', 'effective_date': '2026-09-16', 'stage_code': stage,
                'reason': 'Подтверждено на доске', 'people': people or [{'id': row['id'], 'token': row['stage_token']}],
                'request_key': str(uuid4()), **changes}

    def test_board_transition_is_dated_audited_replayable_and_preserves_plans(self):
        from workforce_core import departure_warnings
        planned = self.request('post', f'people/{self.worker}/movement', self.movement(
            actual_date=None, planned_date='2026-10-10', result_code=None), 'rotation').json
        self.db.native('UPDATE workforce_profiles SET staffing_ready=FALSE WHERE worker_id=%s', (self.worker,))
        body = self.board_move('stage.onsite')
        created = self.request('post', 'transitions', body, 'rotation')
        self.assertEqual(created.status_code, 201, created.json)
        self.assertEqual(created.json['changed'], 1)
        self.assertEqual(self.request('post', 'transitions', body, 'rotation').status_code, 200)
        self.assertEqual(self.db.native('SELECT count(*) FROM workforce_stage_events WHERE worker_id=%s', (self.worker,)).fetchone()[0], 1)
        row = self.board_row()
        self.assertEqual(row['stage_code'], 'stage.onsite')
        self.assertTrue(row['staffing_ready'])
        self.assertIsNone(self.board_row(day='2026-09-15')['stage_code'])
        self.assertEqual(self.request('post', 'transitions', self.board_move('stage.leave'), 'rotation').status_code, 201)
        self.assertIn(self.worker, departure_warnings(self.db, '2026-09-16', [self.worker]))
        history = self.request('get', f'people/{self.worker}', role='rotation').json['history']
        self.assertEqual(sum(item['action'] == 'transition' for item in history), 2)
        current_plan = self.db.native('SELECT * FROM workforce_movements WHERE id=%s', (planned['id'],)).fetchone()
        self.assertEqual(current_plan['planned_date'].isoformat(), '2026-10-10')
        self.assertIsNone(current_plan['actual_date'])
        self.assertEqual(self.request('post', 'transitions', self.board_move('stage.onsite'), 'rotation').status_code, 201)
        self.assertEqual(departure_warnings(self.db, '2026-09-16', [self.worker]), {})

    def test_board_bulk_preflight_rejects_stale_scope_and_date_without_partial_changes(self):
        people = [{'id': worker_id, 'token': self.board_row(worker_id)['stage_token']} for worker_id in self.ids]
        body = self.board_move('stage.onsite', people)
        self.assertEqual(self.request('post', 'transitions', body, 'rotation').status_code, 404)
        self.assertEqual(self.request('post', 'transitions', {**body, 'effective_date': 'not-a-date'}).status_code, 400)
        self.request('post', f'people/{self.ids[1]}/stage', {'stage_code': 'stage.leave', 'effective_date': '2026-09-16',
            'confirmed': True, 'reason': 'Другая служба уже уточнила этап', 'request_key': str(uuid4())})
        stale = self.request('post', 'transitions', body)
        self.assertEqual(stale.status_code, 409, stale.json)
        self.assertIsNone(self.board_row()['stage_code'])
        self.assertEqual(self.db.native('SELECT count(*) FROM workforce_audit WHERE worker_id=%s', (self.worker,)).fetchone()[0], 0)
        wrong_date = self.board_move('stage.onsite', date='2026-09-15', effective_date='2026-09-15')
        self.assertEqual(self.request('post', 'transitions', wrong_date).status_code, 409)
        people = [{'id': worker_id, 'token': self.board_row(worker_id)['stage_token']} for worker_id in self.ids]
        succeeded = self.request('post', 'transitions', self.board_move('stage.onsite', people))
        self.assertEqual(succeeded.status_code, 201, succeeded.json)
        self.assertEqual(succeeded.json['changed'], 2)

    def test_board_accepts_future_single_and_bulk_transitions_as_of_effective_date(self):
        single = self.board_move('stage.onsite', effective_date='2099-01-01')
        created = self.request('post', 'transitions', single, 'rotation')
        self.assertEqual(created.status_code, 201, created.json)
        self.assertIsNone(self.board_row(day='2026-09-16')['stage_code'])
        self.assertEqual(self.board_row(day='2099-01-01')['stage_code'], 'stage.onsite')
        event = self.db.native('''SELECT effective_date FROM workforce_stage_events
            WHERE worker_id=%s ORDER BY sequence DESC LIMIT 1''', (self.worker,)).fetchone()
        self.assertEqual(event['effective_date'].isoformat(), '2099-01-01')
        self.assertEqual(self.db.native("SELECT count(*) FROM workforce_audit WHERE worker_id=%s AND action='transition'",
                                        (self.worker,)).fetchone()[0], 1)

        people = [{'id': worker_id, 'token': self.board_row(worker_id)['stage_token']} for worker_id in self.ids]
        bulk = self.board_move('stage.onsite', people, effective_date='2099-01-02')
        response = self.request('post', 'transitions', bulk, 'admin')
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['changed'], 2)
        self.assertEqual({self.board_row(worker_id, day='2099-01-02')['stage_code'] for worker_id in self.ids},
                         {'stage.onsite'})
        stale = {**bulk, 'request_key': str(uuid4()), 'people': [people[0], {**people[1], 'token': '0' * 64}]}
        self.assertEqual(self.request('post', 'transitions', stale, 'admin').status_code, 409)
    def test_board_service_roles_and_fresh_permissions(self):
        body = self.board_move('stage.onsite')
        for role in ('foreman', 'viewer', 'hr_viewer', 'recruitment'):
            self.assertEqual(self.request('post', 'transitions', body, role).status_code, 403, role)
        self.assertEqual(self.board_row(role='foreman')['transition_targets'], [])
        self.assertEqual(self.board_row(role='recruitment')['transition_targets'], ['stage.pvp'])
        self.assertEqual(self.request('post', 'transitions', self.board_move('stage.pvp'), 'recruitment').status_code, 201)
        self.assertEqual(self.board_row(role='recruitment')['transition_targets'], ['stage.inbound'])
        inbound = self.board_move('stage.inbound')
        self.assertEqual(self.request('post', 'transitions', inbound, 'recruitment').status_code, 201)
        self.assertEqual(self.request('post', 'transitions', inbound, 'recruitment').status_code, 200)
        self.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s", (self.users['recruitment']['id'],))
        self.assertEqual(self.request('post', 'transitions', inbound, 'recruitment').status_code, 404)
        self.db.native("UPDATE users SET role='foreman' WHERE id=%s", (self.users['rotation']['id'],))
        self.assertEqual(self.request('post', 'transitions', self.board_move('stage.onsite'), 'rotation').status_code, 403)

    def test_board_planned_trip_list_requires_a_pending_plan(self):
        self.app.config['TESTING'] = False
        self.request('post', 'transitions', self.board_move('stage.leave'), 'rotation')
        path = 'people?date=2026-09-16&department=TEST-SMU&queue=plans'
        self.assertEqual(self.request('get', path, role='rotation').json['rows'], [])
        created = self.request('post', f'people/{self.worker}/movement', self.movement(
            actual_date=None, planned_date='2026-10-10', result_code=None), 'rotation').json
        self.assertEqual([row['id'] for row in self.request('get', path, role='rotation').json['rows']], [self.worker])
        self.assertEqual(self.board_row(role='rotation')['movement']['planned_date'], '2026-10-10')
        self.request('patch', f'people/{self.worker}/movement/{created["id"]}', {
            'token': created['edit_token'], 'result_code': 'result.cancelled', 'reason': 'Поездка отменена',
            'request_key': str(uuid4())}, 'rotation')
        self.assertEqual(self.request('get', path, role='rotation').json['rows'], [])
        self.assertIsNone(self.board_row(role='rotation')['movement'])

    def test_board_rejects_aba_and_backdated_or_inactive_transitions(self):
        original = self.board_move('stage.onsite')
        self.request('post', 'transitions', self.board_move('stage.leave'), 'rotation')
        self.request('post', 'transitions', self.board_move('stage.pvp'), 'rotation')
        self.request('post', 'transitions', self.board_move('stage.leave'), 'rotation')
        self.assertEqual(self.request('post', 'transitions', original, 'rotation').status_code, 409)
        self.assertEqual(self.request('post', 'transitions', self.board_move('stage.onsite', effective_date='2026-09-15'), 'rotation').status_code, 400)
        current = self.board_move('stage.onsite')
        self.db.native('UPDATE workers SET active=0 WHERE id=%s', (self.worker,))
        self.assertEqual(self.request('post', 'transitions', current, 'rotation').status_code, 409)

    def test_board_page_counts_and_assignment_badge_do_not_leak_cached_permissions(self):
        self.request('post', 'transitions', self.board_move('stage.pvp'), 'recruitment')
        self.app.config['TESTING'] = False
        path = 'people?date=2026-09-16&department=TEST-SMU&stage=stage.pvp&limit=1'
        admin = self.request('get', path).json
        foreman = self.request('get', path, role='foreman').json
        self.assertEqual(admin['totals']['total'], 1)
        self.assertEqual(foreman['rows'][0]['transition_targets'], [])
        self.assertIn('stage.onsite', admin['rows'][0]['transition_targets'])
        self.assertFalse(admin['rows'][0]['assigned'])
        site = self.db.native('SELECT id FROM subobjects LIMIT 1').fetchone()[0]
        self.db.native('''INSERT INTO assignments(work_date,shift,subobject_id,worker_id,foreman_user_id,created_at)
            VALUES ('2026-09-16','1 смена',%s,%s,%s,'now')''', (site, self.worker, self.users['admin']['id']))
        self.assertTrue(self.request('get', path).json['rows'][0]['assigned'])
        self.assertEqual(self.request('get', path + '&offset=1').json['rows'], [])

    def upload_source(self, rows, role='rotation', source='urp:П15', day='2026-09-16', extra_headers=()):
        from io import BytesIO
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = 'Явка' if source.startswith('urp:П') else 'ПВП'
        sheet.append(['ФИО', 'Таб. номер', 'Подразделение', 'Категория ГДЛР', 'Должность', 'Организация', *extra_headers])
        for row in rows:
            sheet.append(row)
        stream = BytesIO()
        book.save(stream)
        self.last_upload_bytes = stream.getvalue()
        stream.seek(0)
        return self.client.post('/api/workforce/imports/preview', data={
            'source_key': source, 'date': day, 'file': (stream, 'Перевахтовка.xlsx')}, headers={'Test-Role': role})

    def test_import_preview_confirm_replay_scope_and_stale(self):
        from pathlib import Path
        from unittest.mock import patch
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        tab = str(uuid4().int)[:18]
        rows = [['Новый Проверочный Импорт', tab, 'TEST-SMU', category, 'Монтажник', 'ЛГСС'],
                ['За пределами доступа', tab + '1', 'OTHER-SMU', category, 'Монтажник', 'ЛГСС']]
        preview = self.upload_source(rows)
        self.assertEqual(preview.status_code, 200, preview.json)
        self.assertEqual(preview.json['counts']['added'], 1)
        self.assertEqual(preview.json['skipped_count'], 1)
        self.assertFalse(self.db.native('SELECT 1 FROM workers WHERE personnel_no=%s', (tab,)).fetchone())
        path = 'imports/' + preview.json['id'] + '/apply'
        data = {'token': preview.json['token'], 'confirmed': True, 'decisions': []}
        self.assertEqual(self.request('post', path, data, 'recruitment').status_code, 403)
        with patch('postgres_backup.create', return_value=Path('verified.dump')) as backup:
            first = self.request('post', path, data, 'rotation')
            self.assertEqual(first.status_code, 200, first.json)
            repeated = self.request('post', path, data, 'rotation')
            self.assertEqual(repeated.json, first.json)
            self.assertEqual(backup.call_count, 1)
        worker_id = self.db.native('SELECT id FROM workers WHERE personnel_no=%s', (tab,)).fetchone()['id']
        self.assertEqual(self.db.native('SELECT count(*) n FROM workforce_source_records WHERE worker_id=%s', (worker_id,)).fetchone()['n'], 1)
        rows[0][0] = 'Имя Исправлено Вручную'
        next_preview = self.upload_source(rows, day='2026-09-17')
        self.assertEqual(next_preview.json['counts']['matched'], 1)
        self.assertEqual(next_preview.json['items'][0]['changes'][0]['field'], 'name')
        self.db.native("UPDATE workforce_profiles SET phone='Изменено после предпросмотра',edit_token=gen_random_uuid() WHERE worker_id=%s", (worker_id,))
        stale = self.request('post', 'imports/' + next_preview.json['id'] + '/apply',
            {'token': next_preview.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
        self.assertEqual(stale.status_code, 409, stale.json)

    def test_import_preserves_absent_identity_and_gdlr_binding(self):
        from pathlib import Path
        from unittest.mock import patch
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        tab = str(uuid4().int)[:18]
        rows = [['Подтверждённый Импорт', tab, 'TEST-SMU', category, 'Монтажник', 'ЛГСС']]
        first = self.upload_source(rows)
        with patch('postgres_backup.create', return_value=Path('verified.dump')):
            result = self.request('post', 'imports/' + first.json['id'] + '/apply',
                {'token': first.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
            self.assertEqual(result.status_code, 200, result.json)
            worker_id = self.db.native('SELECT id FROM workers WHERE personnel_no=%s', (tab,)).fetchone()['id']
            rows[0][3] = 'Категория из Excel не должна менять базу'
            second = self.upload_source(rows, day='2026-09-17')
            result = self.request('post', 'imports/' + second.json['id'] + '/apply',
                {'token': second.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
            self.assertEqual(result.status_code, 200, result.json)
            current = self.db.native('SELECT gc.name FROM employee_gdlr eg JOIN gdlr_categories gc ON gc.id=eg.category_id WHERE worker_id=%s', (worker_id,)).fetchone()['name']
            self.assertEqual(current, category)
            missing = self.upload_source([['Другой Новый Импорт', tab + '2', 'TEST-SMU', category, 'Рабочий', 'ЛГСС']], day='2026-09-18')
            self.assertEqual([row['id'] for row in missing.json['missing']], [worker_id])
            result = self.request('post', 'imports/' + missing.json['id'] + '/apply',
                {'token': missing.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
            self.assertEqual(result.json['removed_from_source'], 1)
            self.assertEqual(self.db.native('SELECT active FROM workers WHERE id=%s', (worker_id,)).fetchone()['active'], 1)
            self.assertFalse(self.db.native('SELECT staffing_ready FROM workforce_profiles WHERE worker_id=%s', (worker_id,)).fetchone()['staffing_ready'])

    def test_applied_preview_cannot_expose_workers_after_scope_reduction(self):
        from io import BytesIO
        from unittest.mock import patch
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        preview = self.upload_source([['Закрытая карточка', str(uuid4().int)[:18], 'TEST-SMU', category, 'Рабочий', 'ЛГСС']])
        data = {'token':preview.json['token'], 'confirmed':True, 'decisions':[]}
        path = 'imports/' + preview.json['id'] + '/apply'
        with patch('postgres_backup.create', return_value=Path('verified.dump')):
            self.assertEqual(self.request('post', path, data, 'rotation').status_code, 200)
        self.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s", (self.users['rotation']['id'],))
        replay = self.client.post('/api/workforce/imports/preview', data={
            'source_key':'urp:П15', 'date':'2026-09-16',
            'file':(BytesIO(self.last_upload_bytes),'Перевахтовка ППС-15.xlsx')}, headers={'Test-Role':'rotation'})
        self.assertEqual(replay.status_code, 403, replay.json)
        self.assertNotIn('Закрытая карточка', replay.get_data(as_text=True))
        self.assertEqual(self.request('post', path, data, 'rotation').status_code, 403)

    def test_hidden_smu_identity_requires_privileged_review_without_disclosure(self):
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        name = 'Совпадение ' + self.suffix
        self.db.native("UPDATE workers SET full_name=%s,employer='ЛГСС' WHERE id=%s", (name, self.ids[1]))
        self.db.native("UPDATE workforce_profiles SET birth_date='1990-01-01' WHERE worker_id=%s", (self.ids[1],))
        preview = self.upload_source([[name, '', 'TEST-SMU', category, 'Рабочий', 'ЛГСС', '01.01.1990']], extra_headers=['Дата рождения'])
        self.assertEqual(preview.status_code, 200, preview.json)
        item = preview.json['items'][0]
        self.assertTrue(item['restricted_identity'])
        self.assertIsNone(item['worker_id'])
        self.assertEqual(item['candidates'], [])
        result = self.request('post', 'imports/' + preview.json['id'] + '/apply', {
            'token':preview.json['token'], 'confirmed':True,
            'decisions':[{'index':0,'worker_id':'new','reason':'Создать повторно'}]}, 'rotation')
        self.assertEqual(result.status_code, 403, result.json)
        self.assertEqual(self.db.native('SELECT count(*) FROM workers WHERE full_name=%s', (name,)).fetchone()[0], 1)

    def test_hidden_identity_created_after_preview_blocks_apply(self):
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        name = 'Конкурирующая личность ' + self.suffix
        preview = self.upload_source([[name, '', 'TEST-SMU', category, 'Рабочий', 'ЛГСС', '01.01.1990']], extra_headers=['Дата рождения'])
        self.assertEqual(preview.json['counts']['added'], 1)
        self.db.native("UPDATE workers SET full_name=%s,employer='ЛГСС' WHERE id=%s", (name, self.ids[1]))
        self.db.native("UPDATE workforce_profiles SET birth_date='1990-01-01' WHERE worker_id=%s", (self.ids[1],))
        result = self.request('post', 'imports/' + preview.json['id'] + '/apply', {
            'token':preview.json['token'], 'confirmed':True, 'decisions':[]}, 'rotation')
        self.assertEqual(result.status_code, 403, result.json)
        self.assertEqual(self.db.native('SELECT count(*) FROM workers WHERE full_name=%s', (name,)).fetchone()[0], 1)

    def test_import_missing_manual_record_stays_in_composition(self):
        from pathlib import Path
        from unittest.mock import patch
        category = self.db.native('SELECT name FROM gdlr_categories WHERE active=1 LIMIT 1').fetchone()['name']
        tab = str(uuid4().int)[:18]
        with patch('postgres_backup.create', return_value=Path('verified.dump')):
            preview = self.upload_source([['Ручная запись для сверки', tab, 'TEST-SMU', category, 'Рабочий', 'ЛГСС']])
            result = self.request('post', 'imports/' + preview.json['id'] + '/apply',
                {'token': preview.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
            self.assertEqual(result.status_code, 200, result.json)
            worker = self.db.native('SELECT id FROM workers WHERE personnel_no=%s', (tab,)).fetchone()['id']
            self.db.native('''INSERT INTO manual_employees(worker_id,created_by,created_at,request_key,payload_hash)
                VALUES (%s,%s,'2026-09-16',%s,'test')''', (worker, self.users['admin']['id'], uuid4().hex))
            preview = self.upload_source([['Другой человек для сверки', tab + '1', 'TEST-SMU', category, 'Рабочий', 'ЛГСС']], day='2026-09-17')
            self.assertTrue(preview.json['missing'][0]['manually_created'])
            result = self.request('post', 'imports/' + preview.json['id'] + '/apply',
                {'token': preview.json['token'], 'confirmed': True, 'decisions': []}, 'rotation')
            self.assertEqual(result.json['manual_preserved'], 1)
            self.assertEqual(result.json['removed_from_source'], 0)
            self.assertTrue(self.db.native('SELECT staffing_ready FROM workforce_profiles WHERE worker_id=%s', (worker,)).fetchone()['staffing_ready'])

    def test_staffing_reader_does_not_receive_private_hr_data(self):
        self.db.native("UPDATE workforce_profiles SET birth_date='1990-01-01',phone='private',notes='private' WHERE worker_id=%s", (self.worker,))
        point = self.create_catalog('travelpoint', 'Закрытый пункт поездки')
        self.db.native('UPDATE workforce_profiles SET origin_code=%s WHERE worker_id=%s', (point['code'],self.worker))
        movement = self.request('post', f'people/{self.worker}/movement', self.movement(origin_code=point['code'],destination_code=point['code']))
        self.assertEqual(movement.status_code, 201, movement.data)
        response = self.request('get', f'people/{self.worker}', role='foreman')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertFalse(response.json['private_details'])
        for field in ('birth_date', 'phone', 'notes', 'origin_city', 'origin_code'):
            self.assertNotIn(field, response.json['profile'])
        for row in response.json['movements']:
            for field in ('origin', 'destination', 'origin_code', 'destination_code'):
                self.assertNotIn(field, row)
        self.assertEqual(response.json['history'], [])
        self.assertEqual(response.json['sources'], [])
        admin = self.request('get', f'people/{self.worker}')
        self.assertEqual(admin.json['profile']['phone'], 'private')

    def test_smu_visibility_and_ready_base_role(self):
        response = self.request('get', 'people?date=2026-09-16', role='rotation')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual([row['id'] for row in response.json['rows']], [self.worker])
        self.assertEqual(self.request('get', f'people/{self.ids[1]}', role='rotation').status_code, 404)
        for kind, body in [('movement', self.movement()), ('stage', {'stage_code': 'stage.leave',
                'effective_date': '2026-09-16', 'confirmed': True, 'reason': 'Причина', 'request_key': str(uuid4())})]:
            self.assertEqual(self.request('post', f'people/{self.worker}/{kind}', body, 'foreman').status_code, 403)

    def test_service_boundaries_and_missing_scope(self):
        self.assertEqual(self.request('post', f'people/{self.worker}/movement', self.movement(), 'recruitment').status_code, 403)
        document = {'document_code': 'document.patent', 'reason': 'Патент получен', 'request_key': str(uuid4())}
        self.assertEqual(self.request('post', f'people/{self.worker}/document', document, 'rotation').status_code, 403)
        self.assertEqual(self.request('post', f'people/{self.worker}/document', document, 'recruitment').status_code, 201)
        self.db.native('DELETE FROM user_smu_access WHERE user_id=%s', (self.users['recruitment']['id'],))
        self.assertEqual(self.request('get', 'people?date=2026-09-16', role='recruitment').json['totals']['total'], 0)

    def test_replay_stale_edit_and_audit(self):
        body = self.movement()
        first = self.request('post', f'people/{self.worker}/movement', body, 'rotation')
        self.assertEqual(first.status_code, 201, first.json)
        again = self.request('post', f'people/{self.worker}/movement', body, 'rotation')
        self.assertEqual(first.json['id'], again.json['id'])
        altered = {**body, 'actual_date': '2026-09-17'}
        self.assertEqual(self.request('post', f'people/{self.worker}/movement', altered, 'rotation').status_code, 409)
        path = f'people/{self.worker}/movement/{first.json["id"]}'
        patch = {'token': first.json['edit_token'], 'notes': 'Уточнение', 'reason': 'Обновление', 'request_key': str(uuid4())}
        self.assertEqual(self.request('patch', path, patch, 'rotation').status_code, 200)
        self.assertEqual(self.request('patch', path, {**patch, 'request_key': str(uuid4())}, 'rotation').status_code, 409)
        history = self.request('get', f'people/{self.worker}').json['history']
        self.assertEqual(len([row for row in history if row['entity_type'] == 'movement']), 2)
        self.assertEqual(len([row for row in history if row['entity_type'] == 'stage']), 1)
        self.assertTrue(all(row['actor_id'] == self.users['rotation']['id'] for row in history))

    def test_confirmed_trip_updates_presence_and_correction_retracts_old_fact(self):
        created = self.request('post', f'people/{self.worker}/movement', self.movement(direction='arrival'), 'rotation')
        self.assertEqual(created.status_code, 201, created.json)
        before = self.request('get', 'people?date=2026-09-15', role='rotation').json
        self.assertIsNone(before['rows'][0]['stage_code'])
        current = self.request('get', 'people?date=2026-09-16', role='rotation').json
        self.assertEqual(current['rows'][0]['stage_code'], 'stage.onsite')
        corrected = self.request('patch', f'people/{self.worker}/movement/{created.json["id"]}',
            {'token': created.json['edit_token'], 'actual_date': '2026-09-17', 'reason': 'Дата была ошибочной', 'request_key': str(uuid4())}, 'rotation')
        self.assertEqual(corrected.status_code, 200, corrected.json)
        self.assertIsNone(self.request('get', 'people?date=2026-09-16', role='rotation').json['rows'][0]['stage_code'])
        self.assertEqual(self.request('get', 'people?date=2026-09-17', role='rotation').json['rows'][0]['stage_code'], 'stage.onsite')
        cancelled = self.request('patch', f'people/{self.worker}/movement/{created.json["id"]}',
            {'token': corrected.json['edit_token'], 'result_code': 'result.cancelled', 'reason': 'Подтверждение было ошибочным', 'request_key': str(uuid4())}, 'rotation')
        self.assertEqual(cancelled.status_code, 200, cancelled.json)
        self.assertIsNone(self.request('get', 'people?date=2026-09-17', role='rotation').json['rows'][0]['stage_code'])
        self.assertEqual(self.db.native('SELECT count(*) n FROM workforce_stage_events WHERE worker_id=%s AND retracted', (self.worker,)).fetchone()['n'], 2)

    def test_departure_warning_does_not_infer_from_elapsed_plan(self):
        from workforce_core import departure_warnings
        self.request('post', f'people/{self.worker}/movement', self.movement(), 'rotation')
        self.assertEqual(departure_warnings(self.db, '2026-09-15', [self.worker]), {})
        self.assertIn(self.worker, departure_warnings(self.db, '2026-09-16', [self.worker]))
        self.request('post', f'people/{self.worker}/movement', self.movement(direction='arrival', actual_date='2026-09-17', result_code=None), 'rotation')
        self.assertIn(self.worker, departure_warnings(self.db, '2026-09-18', [self.worker]))
        self.request('post', f'people/{self.worker}/movement', self.movement(direction='arrival', actual_date='2026-09-18'), 'rotation')
        self.assertEqual(departure_warnings(self.db, '2026-09-18', [self.worker]), {})

    def test_stage_asof_ignores_unconfirmed_and_future_events(self):
        for day, stage, confirmed in [('2026-09-14', 'stage.leave', True), ('2026-09-15', 'stage.onsite', False),
                                      ('2026-09-17', 'stage.pvp', True)]:
            response = self.request('post', f'people/{self.worker}/stage', {'effective_date': day, 'stage_code': stage,
                'confirmed': confirmed, 'reason': 'Проверка состояния', 'request_key': str(uuid4())}, 'rotation')
            self.assertEqual(response.status_code, 201, response.json)
        rows = self.request('get', 'people?date=2026-09-16', role='rotation').json['rows']
        self.assertEqual(rows[0]['stage_code'], 'stage.leave')
        rows = self.request('get', 'people?date=2026-09-17', role='rotation').json['rows']
        self.assertEqual(rows[0]['stage_code'], 'stage.pvp')

    def test_profile_stale_token_and_employer_sync(self):
        before = self.request('get', f'people/{self.worker}').json['profile']
        self.db.native('UPDATE workers SET employer=%s WHERE id=%s', ('Работодатель изменён', self.worker))
        current = self.request('get', f'people/{self.worker}').json['profile']
        self.assertEqual(current['employer'], 'Работодатель изменён')
        data = {'token': before['token'], 'phone': '+7 900 000 00 00', 'reason': 'Контакт', 'request_key': str(uuid4())}
        self.assertEqual(self.request('patch', f'people/{self.worker}/profile', data, 'rotation').status_code, 409)
        data.update(token=current['token'], request_key=str(uuid4()))
        response = self.request('patch', f'people/{self.worker}/profile', data, 'rotation')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['phone'], data['phone'])
        self.assertEqual(self.request('patch', f'people/{self.worker}/profile', data, 'recruitment').status_code, 403)

    def test_validation_and_regex_errors_are_explicit(self):
        self.assertEqual(self.request('post', f'people/{self.worker}/movement', self.movement(actual_date=None), 'rotation').status_code, 400)
        self.assertEqual(self.request('post', f'people/{self.worker}/movement', self.movement(basis_code='document.patent'), 'rotation').status_code, 400)
        response = self.request('get', 'people?date=2026-09-16&regex=1&q=%5B', role='rotation')
        self.assertEqual(response.status_code, 400, response.json)

    def test_catalog_admin_only_and_scope_required(self):
        data = {'label': 'Тестовое гражданство ' + self.suffix, 'reason': 'Справочник', 'request_key': str(uuid4())}
        self.assertEqual(self.request('post', 'catalog/citizenship', data, 'foreman').status_code, 403)
        self.assertEqual(self.request('post', 'catalog/citizenship', data, 'recruitment').status_code, 403)
        result = self.request('post', 'catalog/citizenship', data)
        self.assertEqual(result.status_code, 201, result.json)
        self.assertTrue(result.json['code'].startswith('citizenship.'))

    def test_arrival_rescheduling_preserves_chain_and_rejects_completed(self):
        plan = self.movement(direction='arrival', actual_date=None, planned_date='2026-09-20', result_code=None)
        first = self.request('post', f'people/{self.worker}/movement', plan, 'rotation').json
        path = f'people/{self.worker}/movements/{first["id"]}/reschedule'
        data = {'planned_date': '2026-09-22', 'token': first['edit_token'],
                'reason': 'Изменение билета', 'request_key': str(uuid4())}
        response = self.request('post', path, data, 'rotation')
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['previous']['planned_date'], '2026-09-20')
        self.assertEqual(response.json['previous']['result_code'], 'result.postponed')
        second = response.json['current']
        self.assertEqual(second['planned_date'], '2026-09-22')
        self.assertEqual(second['rescheduled_from'], first['id'])
        self.assertIsNone(second['result_code'])
        repeated = self.request('post', path, data, 'rotation')
        self.assertEqual(repeated.json['current']['id'], second['id'])
        data2 = {**data, 'planned_date': '2026-09-24', 'token': second['edit_token'], 'request_key': str(uuid4())}
        latest = self.request('post', f'people/{self.worker}/movements/{second["id"]}/reschedule', data2, 'rotation')
        self.assertEqual(latest.status_code, 201, latest.json)
        third = latest.json['current']
        complete = {'actual_date': '2026-09-24', 'result_code': 'result.happened', 'token': third['edit_token'],
                    'reason': 'Прибыл', 'request_key': str(uuid4())}
        completed = self.request('patch', f'people/{self.worker}/movement/{third["id"]}', complete, 'rotation')
        self.assertEqual(completed.status_code, 200, completed.json)
        prohibited = {**data2, 'token': completed.json['edit_token'], 'request_key': str(uuid4())}
        self.assertEqual(self.request('post', f'people/{self.worker}/movements/{third["id"]}/reschedule', prohibited, 'rotation').status_code, 400)

    def test_rotation_extension_uses_schedule_snapshot_and_preserves_ticket(self):
        schedule_data = {'name': 'График ' + self.suffix, 'onsite_days': 30, 'leave_days': 30, 'travel_days': 2,
                         'reason': 'Утверждён график', 'request_key': str(uuid4())}
        schedule = self.request('post', 'rotation-schedules', schedule_data).json
        data = {'schedule_id': schedule['id'], 'start_date': '2026-09-01', 'reason': 'Новая вахта', 'request_key': str(uuid4())}
        result = self.request('post', f'people/{self.worker}/rotations', data, 'rotation')
        self.assertEqual(result.status_code, 201, result.json)
        first = result.json
        self.assertEqual(first['planned_end_date'], '2026-09-30')
        self.assertEqual(first['leave_end_date'], '2026-10-30')
        self.assertEqual(first['next_arrival_date'], '2026-11-01')
        # Editing the directory must not silently alter a previously saved cycle.
        self.db.native('UPDATE workforce_rotation_schedules SET leave_days=45 WHERE id=%s', (schedule['id'],))
        ticket = self.request('post', f'people/{self.worker}/movement', self.movement(direction='arrival',
            planned_date='2026-11-01', actual_date=None, result_code=None, basis_code='basis.ticket'), 'rotation').json
        extension = {'new_end_date': '2026-10-05', 'token': first['edit_token'], 'reason': 'Продление по согласованию', 'request_key': str(uuid4())}
        response = self.request('post', f'people/{self.worker}/rotations/{first["id"]}/extend', extension, 'rotation')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['rotation']['next_arrival_date'], '2026-11-06')
        self.assertEqual(response.json['trip_plans_to_review'][0]['id'], ticket['id'])
        current = self.request('get', f'people/{self.worker}').json
        self.assertEqual(current['movements'][0]['planned_date'], '2026-11-01')
        self.assertEqual(current['history'][0]['before_json']['planned_end_date'], '2026-09-30')
        self.assertEqual(current['history'][0]['after_json']['planned_end_date'], '2026-10-05')
        extension['request_key'] = str(uuid4())
        self.assertEqual(self.request('post', f'people/{self.worker}/rotations/{first["id"]}/extend', extension, 'rotation').status_code, 409)

    def test_report_columns_scope_dates_and_formula_injection(self):
        import io
        from openpyxl import load_workbook
        from workforce_export import FIRST_HEADERS, SECOND_HEADERS
        self.db.native("UPDATE workers SET category='Арматурщик',full_name='=HYPERLINK(\"https://example.invalid\")' WHERE id=%s", (self.worker,))
        self.db.native("UPDATE workforce_profiles SET leave_end_date='2026-09-20' WHERE worker_id=%s", (self.worker,))
        event = {'stage_code': 'stage.leave', 'effective_date': '2026-09-01', 'confirmed': True,
                 'reason': 'Межвахтовый отпуск', 'request_key': str(uuid4())}
        self.request('post', f'people/{self.worker}/stage', event, 'rotation')
        response = self.request('get', 'export?date=2026-09-16', role='rotation')
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['Явка и аутстаффинг', 'Неявка, заезд и ПВП'])
        self.assertEqual([cell.value for cell in book.worksheets[0][1]], FIRST_HEADERS)
        self.assertEqual([cell.value for cell in book.worksheets[1][1]], SECOND_HEADERS)
        sheet = book.worksheets[1]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['K2'].value.date().isoformat(), '2026-09-22')
        self.assertEqual(sheet['L2'].value, 'Заезд')
        self.assertEqual(sheet['F2'].data_type, 's')
        self.assertEqual(sheet['F2'].value, '=HYPERLINK("https://example.invalid")')


if __name__ == '__main__':
    unittest.main()


