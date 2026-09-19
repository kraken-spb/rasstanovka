import io
import unittest

import openpyxl

from tests import test_staffing_export as export_tests


EMPTY = '__staffing_export_empty__'


class StaffingExportReferenceFiltersTest(export_tests.StaffingExportTest):
    def setUp(self):
        super().setUp()
        self.second_site = None
        self.first_object = None
        self.second_object = None
        self.second_worker = None
        self.empty_worker = None
        self.free_worker = None

    def add_reference_rows(self):
        if self.second_site is not None:
            return
        with self.module.app.app_context():
            db = self.module.get_db()
            self.first_object = db.execute('SELECT object_id FROM subobjects WHERE id=?', (self.site_id,)).fetchone()[0]
            self.second_object = db.execute("INSERT INTO objects(name) VALUES ('Группа Б')").lastrowid
            self.second_site = db.execute("INSERT INTO subobjects(object_id,name) VALUES (?, 'Подобъект 1')", (self.second_object,)).lastrowid
            self.second_worker = self.worker(db, 'Второй фильтр', '000127', contractor='Другой подрядчик', profession='Монтажник')
            db.execute('''INSERT INTO assignments(work_date,shift,subobject_id,worker_id,employer,foreman_user_id,created_at,crew_id,edit_token)
                VALUES ('2026-09-13','1 смена',?,?,?,?,'now',?,'token-second')''',
                (self.second_site, self.second_worker, 'Работодатель второй', self.other_id, self.other_crew_id))
            self.empty_worker = self.worker(db, 'Без ИТР', '000128', profession='Слесарь')
            db.execute('''INSERT INTO assignments(work_date,shift,subobject_id,worker_id,employer,foreman_user_id,created_at,crew_id,edit_token)
                VALUES ('2026-09-13','1 смена',?,?,?,?,'now',NULL,'token-empty')''',
                (self.site_id, self.empty_worker, 'Работодатель пустой', self.other_id))
            self.free_worker = self.worker(db, 'Нерасставленный', '000129', profession='Слесарь')
            db.execute("INSERT INTO manual_employees(worker_id,created_by,created_at,request_key,payload_hash) VALUES (?,?,'now','reference-free','hash')",
                       (self.free_worker, self.admin_id))
            category_id = db.execute("SELECT id FROM gdlr_categories WHERE name='Ручная ГДЛР'").fetchone()[0]
            db.execute('UPDATE gdlr_categories SET staffing_allowed=1 WHERE id=?', (category_id,))
            db.execute("INSERT INTO employee_gdlr(worker_id,category_id,edit_token,updated_by,updated_at) VALUES (?,?,'free-token',?,'now')",
                       (self.free_worker, category_id, self.admin_id))
            db.commit()

    def response_rows(self, response):
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        return list(openpyxl.load_workbook(io.BytesIO(response.data))['Список сотрудников'].iter_rows(min_row=2, values_only=True))

    def export_with(self, client=None, **params):
        return (client or self.admin).get('/api/staffing/export', query_string={
            'date': '2026-09-13', 'shift': 'all', **params})

    def test_options_are_scoped_and_keep_duplicate_subobjects_distinct(self):
        self.add_reference_rows()
        response = self.admin.get('/api/staffing/export/options?date=2026-09-13&shift=all')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        subobjects = {item['value']: item['label'] for item in data['subobjects']}
        self.assertEqual(subobjects[str(self.site_id)], 'Группа А · Подобъект 1')
        self.assertEqual(subobjects[str(self.second_site)], 'Группа Б · Подобъект 1')
        self.assertIn({'value': '=Работодатель', 'label': '=Работодатель'}, data['employers'])
        self.assertIn({'value': 'ИТР вручную', 'label': 'ИТР вручную'}, data['linear_itrs'])
        self.assertIn({'value': 'Бригадир вручную', 'label': 'Бригадир вручную'}, data['brigadiers'])
        foreman = self.foreman.get('/api/staffing/export/options?date=2026-09-13&shift=all').get_json()
        self.assertNotIn(str(self.second_site), {item['value'] for item in foreman['subobjects']})
        self.assertEqual(self.export_with(self.foreman, object_id=str(self.second_object)).status_code, 404)

    def test_reference_filters_intersect_and_multiselect_apply_to_list_and_summary(self):
        self.add_reference_rows()
        response = self.export_with(object_id=[str(self.first_object), str(self.second_object)],
                                    employer=['=Работодатель', 'Работодатель второй'], profession=['Должность', 'Монтажник'],
                                    linear_itr=['ИТР вручную', 'Чужой ИТР'], brigadier=['Бригадир вручную', 'Чужой бригадир'])
        rows = self.response_rows(response)
        self.assertEqual(response.headers['X-Export-Count'], '2')
        self.assertEqual({row[5] for row in rows}, {'=Формула', 'Второй фильтр'})
        book = openpyxl.load_workbook(io.BytesIO(response.data), data_only=True)
        summary = book['Сводная таблица']
        self.assertEqual(summary.cell(summary.max_row, summary.max_column).value, len(rows))

    def test_empty_effective_value_and_unassigned_population_are_filterable(self):
        self.add_reference_rows()
        response = self.export_with(linear_itr=EMPTY)
        self.assertIn('Без ИТР', [row[5] for row in self.response_rows(response)])
        response = self.export_with(include_unassigned='1', object_id=EMPTY, subobject_id=EMPTY, profession='Слесарь')
        rows = self.response_rows(response)
        self.assertEqual([row[5] for row in rows], ['Нерасставленный'])
        self.assertEqual(response.headers['X-Export-Unassigned-Count'], '1')

    def test_snapshot_employer_is_not_worker_employer(self):
        response = self.export_with(employer='=Работодатель')
        self.assertEqual([row[5] for row in self.response_rows(response)], ['=Формула'])
        self.assertEqual(self.export_with(employer='Старый работодатель').status_code, 404)


if __name__ == '__main__':
    unittest.main()
