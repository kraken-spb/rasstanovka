import io
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.utils.escape import unescape
from native_pivot import summary_sheet, save_workbook_with_slicer
from staffing_export import HEADERS


class NativePivotTest(unittest.TestCase):
    def report(self, rows):
        book=Workbook()
        source=book.active
        source.title='Список сотрудников'
        source.append(HEADERS)
        for row in rows:
            source.append(row)
        source.auto_filter.ref=f'A1:P{source.max_row}'
        summary_sheet(book, [], '2026-09-14')
        output=io.BytesIO()
        save_workbook_with_slicer(book, output)
        return output.getvalue()

    def row(self, name='Иванов Иван', company='Подрядчик', employer='Работодатель', shift='День'):
        return [1,'Группа','Подобъект',company,employer,name,'000123','Должность','Профессия',
                'Категория','ИТР','Бригадир',shift,'СМУ','Монтаж\nСварка','Электромонтаж']

    def test_real_pivot_cache_source_fields_and_count_aggregation(self):
        payload=self.report([self.row(),self.row(name=None),self.row(name='Другой',shift='Ночь')])
        book=load_workbook(io.BytesIO(payload))
        self.assertEqual(book.sheetnames,['Список сотрудников','Сводная таблица'])
        source, summary=book.worksheets
        self.assertIsNone(source.auto_filter.ref)
        self.assertEqual(source.tables['StaffingSource'].ref,'A1:P4')
        self.assertEqual(len(summary._pivots),1)
        pivot=summary._pivots[0]
        self.assertEqual([unescape(pivot.cache.cacheFields[f.x].name) for f in pivot.rowFields],HEADERS[1:3])
        self.assertEqual([unescape(pivot.cache.cacheFields[f.x].name) for f in pivot.colFields],[HEADERS[i] for i in (3,4,12)])
        self.assertEqual((pivot.dataFields[0].fld,pivot.dataFields[0].subtotal),(5,'count'))
        self.assertEqual(pivot.cache.cacheSource.worksheetSource.name,'StaffingSource')
        self.assertEqual(pivot.cache.recordCount,3)
        self.assertEqual(len(pivot.cache.records.r),3)
        self.assertFalse(pivot.disableFieldList)
        self.assertTrue(pivot.cache.enableRefresh)
        self.assertTrue(pivot.cache.refreshOnLoad)
        self.assertEqual(summary.cell(summary.max_row,summary.max_column).value,2)
        with ZipFile(io.BytesIO(payload)) as archive:
            self.assertIn('xl/pivotTables/pivotTable1.xml',archive.namelist())
            self.assertIn('xl/pivotCache/pivotCacheRecords1.xml',archive.namelist())
            self.assertFalse(any('externalLink' in name for name in archive.namelist()))
            records=ET.fromstring(archive.read('xl/pivotCache/pivotCacheRecords1.xml'))
            self.assertEqual(unescape(records[0][14].get('v')),'Монтаж\nСварка')
            self.assertEqual(records[0][6].get('v'),'000123')

    def test_counts_come_from_source_sheet_with_distinct_company_employer_shift_axes(self):
        data=[self.row(company='А',employer='Один'),self.row(company='А',employer='Два',shift='Ночь'),
              self.row(company='Б',employer='Один'),self.row(company='',employer='')]
        book=load_workbook(io.BytesIO(self.report(data)))
        summary=book.worksheets[1]
        self.assertEqual(summary.cell(summary.max_row,summary.max_column).value,4)
        self.assertEqual({unescape(item.v) for item in summary._pivots[0].cache.cacheFields[3].sharedItems._fields if item.tagname == 's'},{'А','Б'})
        self.assertEqual({unescape(item.v) for item in summary._pivots[0].cache.cacheFields[4].sharedItems._fields if item.tagname == 's'},{'Один','Два'})

    def test_visible_slicer_uses_authorized_source_categories_and_valid_cache_indexes(self):
        rows = [self.row(name=str(i)) for i in range(4)]
        for row, category in zip(rows, ['Монтаж & сварка', None, 'Монтаж & сварка', '=Категория']):
            row[9] = category
        payload = self.report(rows)
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
              'x': 'http://schemas.microsoft.com/office/spreadsheetml/2009/9/main',
              'd': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
              'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        with ZipFile(io.BytesIO(payload)) as archive:
            cache = ET.fromstring(archive.read('xl/pivotCache/pivotCacheDefinition1.xml'))
            shared = cache.find('s:cacheFields', ns)[9].find('s:sharedItems', ns)
            self.assertEqual([item.get('v') for item in shared], ['Монтаж & сварка', None, '=Категория'])
            records = ET.fromstring(archive.read('xl/pivotCache/pivotCacheRecords1.xml'))
            self.assertEqual([int(row[9].get('v')) for row in records], [0, 1, 0, 2])
            slicer_cache = ET.fromstring(archive.read('xl/slicerCaches/slicerCache1.xml'))
            self.assertEqual(slicer_cache.get('sourceName'), HEADERS[9])
            self.assertEqual(slicer_cache.find('x:pivotTables/x:pivotTable', ns).attrib,
                             {'tabId': '2', 'name': 'StaffingPivot'})
            tabular = slicer_cache.find('x:data/x:tabular', ns)
            self.assertEqual(tabular.get('pivotCacheId'), cache.find('s:extLst/s:ext/x:pivotCacheDefinition', ns).get('pivotCacheId'))
            items = tabular.find('x:items', ns)
            self.assertEqual({int(i.get('x')) for i in items}, {0, 1, 2})
            self.assertTrue(all(i.get('s') == '1' for i in items))
            self.assertEqual(int(items[-1].get('x')), 1)  # Excel sorts blanks last.
            slicer = ET.fromstring(archive.read('xl/slicers/slicer1.xml'))[0]
            self.assertEqual(slicer.get('cache'), slicer_cache.get('name'))
            drawing = ET.fromstring(archive.read('xl/drawings/drawing1.xml'))
            self.assertEqual(drawing.find('.//d:from/d:row', ns).text, '0')
            self.assertEqual(drawing.find('.//d:to/d:row', ns).text, '7')
            book = ET.fromstring(archive.read('xl/workbook.xml'))
            self.assertEqual(book.find('s:definedNames/s:definedName[@name="Slicer_GDLR"]', ns).text, '#N/A')
            self.assertEqual(book.find('s:bookViews/s:workbookView', ns).get('activeTab'), '1')
            sheet = ET.fromstring(archive.read('xl/worksheets/sheet2.xml'))
            self.assertEqual(sheet.find('s:sheetFormatPr', ns).get('defaultRowHeight'), '15')
            self.assertTrue(all('ht' not in r.attrib for r in sheet.find('s:sheetData', ns)))
            self.assertTrue(all(float(c.get('width')) == 10 for c in sheet.find('s:cols', ns) if int(c.get('min')) > 1))
            self.assertIsNotNone(sheet.find('s:extLst/s:ext/x:slicerList/x:slicer', ns))
