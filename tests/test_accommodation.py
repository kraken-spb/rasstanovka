import os
import unittest
from datetime import datetime, time
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from accommodation import COMPANY, HEADERS, TITLE, clean_row, request_workbook


class AccommodationWorkbookTest(unittest.TestCase):
    def test_partial_arrival_preserves_only_known_date_or_time(self):
        base=dict(personnel_no='0001',full_name='Сотрудник',profession='',category='',
                  birth_date='',citizenship='',phone='',object='',worker_category='',
                  departure_date='',responsible='',basis='')
        for day,clock,expected,number_format in (
            ('','',None,'General'),('2026-09-20','',datetime(2026,9,20),'dd.mm.yyyy'),
            ('','08:30',time(8,30),'hh:mm'),
            ('2026-09-20','08:30',datetime(2026,9,20,8,30),'dd.mm.yyyy hh:mm')):
            with self.subTest(day=day,clock=clock):
                book=load_workbook(BytesIO(request_workbook([{**base,'arrival_date':day,'arrival_time':clock}],'2026-09-18')))
                self.assertEqual(book.active['K6'].value,expected)
                self.assertEqual(book.active['K6'].number_format,number_format)
                for address in ('D6','E6','F6','G6','H6','I6','J6','L6','M6','N6'):
                    self.assertIsNone(book.active[address].value)
                book.close()

    def test_exact_columns_dates_identifiers_and_safe_text(self):
        row = dict(personnel_no='000825',full_name='=Сотрудник',profession='+Монтажник',category='Монтажник ТТ',
                   birth_date='1980-05-04',citizenship='Россия',phone='+79990000000',object='УКПГ-45',
                   worker_category='Рабочий',arrival_date='2026-09-20',arrival_time='13:40',
                   departure_date='2026-09-22',responsible='Иванов Иван Иванович',basis='=Основание')
        book=load_workbook(BytesIO(request_workbook([row],'2026-09-18')));sheet=book.active
        self.assertEqual(sheet['E1'].value,COMPANY)
        self.assertEqual(sheet['E2'].value,TITLE)
        self.assertEqual(len(HEADERS),14)
        self.assertEqual([c.value for c in sheet[5]],list(HEADERS))
        self.assertEqual(sheet['B6'].value,'000825')
        self.assertEqual(sheet['C6'].data_type,'s')
        self.assertEqual(sheet['N6'].data_type,'s')
        self.assertEqual(sheet['K6'].value,datetime(2026,9,20,13,40))
        self.assertEqual(sheet['F6'].value,datetime(1980,5,4))
        self.assertEqual(sheet.freeze_panes,'D6')
        self.assertEqual(sheet.auto_filter.ref,'A5:N6')
        self.assertEqual(str(sheet.page_setup.paperSize),sheet.PAPERSIZE_A4)
        self.assertEqual(sheet.page_setup.orientation,'landscape')
        self.assertEqual(sheet.page_setup.fitToWidth,1)
        self.assertEqual(sheet.page_setup.fitToHeight,0)
        self.assertEqual(sheet.print_title_rows,'$1:$5')
        self.assertEqual(len(sheet._images),1)
        book.close()

    def test_long_reason_continues_without_losing_text_or_inflating_people_count(self):
        reason=('Продлённое размещение по согласованной заявке подразделения. '*40)[:2000]
        row=dict(personnel_no='0001',full_name='Иванов Иван Иванович',profession='Монтажник',
                 category='Монтажник ТТ',birth_date='',citizenship='Россия',phone='',object='УКПГ-45',
                 worker_category='Рабочий',arrival_date='2026-09-20',arrival_time='08:30',
                 departure_date='',responsible='Петров Пётр Петрович',basis=reason)
        book=load_workbook(BytesIO(request_workbook([row],'2026-09-18')));sheet=book.active
        self.assertGreater(sheet.max_row,6)
        text=''.join(str(sheet.cell(r,14).value or '') for r in range(6,sheet.max_row+1))
        self.assertEqual(''.join(text.split()),''.join(reason.split()))
        self.assertEqual(sheet['I3'].value,'Сотрудников: 1')
        self.assertTrue(all(sheet.row_dimensions[r].height<=310 for r in range(6,sheet.max_row+1)))
        book.close()


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'),'Select isolated PostgreSQL explicitly.')
class AccommodationApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import test_workforce_api as fixtures
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        import test_workforce_api as fixtures
        self.f=fixtures.WorkforceApiTest();self.f.setUp();self.addCleanup(self.f.tearDown)

    def prepare(self,role='admin',ids=None):
        return self.f.request('post','accommodation/prepare',{'ids':ids or [self.f.worker],'date':'2026-09-18'},role=role)

    def payload(self,role='admin'):
        result=self.prepare(role);self.assertEqual(result.status_code,200,result.json)
        rows=[{key:r[key] for key in ('id','token','object','worker_category','arrival_date','arrival_time','departure_date','responsible','basis')}
              for r in result.json['rows']]
        for r in rows:r.update(object='УКПГ-45',worker_category='Рабочий',arrival_date='2026-09-20',arrival_time='08:30',departure_date='2026-09-22',basis='По заявке')
        return {'date':'2026-09-18','request_key':str(uuid4()),'rows':rows}

    def export(self,payload,role='admin'):
        response=self.f.request('post','accommodation/export',payload,role=role)
        self.addCleanup(response.close)
        return response

    def test_prepare_scope_roles_and_unknown_defaults(self):
        for role in ('foreman','viewer'):
            self.assertEqual(self.prepare(role).status_code,403)
        for role in ('rotation','recruitment','hr_viewer','admin'):
            response=self.prepare(role);self.assertEqual(response.status_code,200,response.json)
            self.assertEqual(response.json['rows'][0]['arrival_time'],'')
            self.assertEqual(response.json['rows'][0]['worker_category'],'')
        self.assertEqual(self.prepare('rotation',self.f.ids).status_code,404)
        self.assertEqual(self.prepare(ids=[self.f.worker,self.f.worker]).status_code,400)

    def test_catalog_profile_filter_sort_bulk_and_registry_export(self):
        f=self.f
        profile=f.request('get',f'people/{f.worker}').json['profile']
        change={'accommodation_code':'accommodation.hostel','token':profile['token'],'reason':'Проживание', 'request_key':str(uuid4())}
        response=f.request('patch',f'people/{f.worker}/profile',change,role='rotation')
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['accommodation'],'Хостел')
        for worker in f.ids:f.add_source_record(worker,'urp:П15')
        filtered=f.request('get','people?date=2026-09-18&section=rotation&accommodation=accommodation.hostel&sort=[{"field":"accommodation","direction":"asc"}]')
        self.assertEqual([r['id'] for r in filtered.json['rows']],[f.worker])
        empty=f.request('get','people?date=2026-09-18&section=rotation&accommodation=__none__&q='+f.suffix)
        self.assertEqual([r['id'] for r in empty.json['rows']],[f.ids[1]])
        self.assertEqual(f.request('get','people?accommodation=employment.staff').status_code,400)
        bulk=f.request('post','bulk/prepare',{'ids':[f.worker]}).json
        self.assertIn('accommodation_code',bulk['fields'])
        download=f.request('get','people/export?date=2026-09-18&section=rotation&accommodation=accommodation.hostel')
        self.addCleanup(download.close)
        book=load_workbook(BytesIO(download.data));sheet=book.active
        accommodation_column=[cell.value for cell in sheet[4]].index('Проживание')+1
        self.assertEqual(sheet.cell(5,accommodation_column).value,'Хостел');book.close()

    def test_export_idempotency_immutability_audit_and_stale_source(self):
        f=self.f;payload=self.payload();result=self.export(payload)
        self.assertEqual(result.status_code,200,result.json if result.is_json else '')
        self.assertEqual(result.headers['X-Export-Row-Count'],'1')
        self.assertEqual(self.export(payload).status_code,200)
        self.assertEqual(f.db.native("SELECT count(*) FROM workforce_audit WHERE entity_id=%s AND entity_type='accommodation_request'",(payload['request_key'],)).fetchone()[0],1)
        self.assertIsNone(f.db.native('SELECT accommodation_code FROM workforce_profiles WHERE worker_id=%s',(f.worker,)).fetchone()[0])
        f.db.native('UPDATE workers SET full_name=%s WHERE id=%s',('Новое ФИО',f.worker))
        # Retry returns the issued snapshot, while a new request needs a fresh preview.
        replay=self.export(payload);self.assertEqual(replay.status_code,200)
        book=load_workbook(BytesIO(replay.data));self.assertEqual(book.active['C6'].value,'Тестовый Сотрудник');book.close()
        payload['request_key']=str(uuid4());self.assertEqual(self.export(payload).status_code,409)

    def test_empty_basis_exports_blank_cell_and_preserves_audit(self):
        payload=self.payload();payload['rows'][0]['basis']=''
        result=self.export(payload)
        self.assertEqual(result.status_code,200,result.json if result.is_json else '')
        book=load_workbook(BytesIO(result.data))
        self.assertEqual(book.active['N5'].value,'Основание')
        self.assertIsNone(book.active['N6'].value)
        book.close()
        self.assertEqual(self.f.db.native("SELECT count(*) FROM workforce_audit WHERE entity_id=%s AND entity_type='accommodation_request'",(payload['request_key'],)).fetchone()[0],1)

    def test_invalid_fields_times_dates_and_payload_do_not_create_audit(self):
        payload=self.payload()
        for key,value in [('arrival_time','25:00'),('arrival_time','08:30:00'),('arrival_time',None),
                          ('arrival_date','2026-02-30'),('departure_date','invalid'),
                          ('departure_date','2026-09-19'),('worker_category','неизвестно'),('object',123)]:
            with self.subTest(key=key,value=value):
                original=payload['rows'][0][key];payload['rows'][0][key]=value
                self.assertEqual(self.export(payload).status_code,400)
                payload['rows'][0][key]=original
        payload['rows'][0]['full_name']='Подмена'
        self.assertEqual(self.export(payload).status_code,400)
        self.assertEqual(self.f.db.native("SELECT count(*) FROM workforce_audit WHERE entity_id=%s",(payload['request_key'],)).fetchone()[0],0)

    def test_incomplete_request_exports_and_retains_audit_and_source_profile(self):
        payload=self.payload()
        for key in ('object','worker_category','arrival_date','arrival_time','departure_date','responsible','basis'):
            payload['rows'][0][key]=''
        result=self.export(payload)
        self.assertEqual(result.status_code,200,result.json if result.is_json else '')
        book=load_workbook(BytesIO(result.data))
        self.assertEqual(book.active['C6'].value,'Тестовый Сотрудник')
        for address in ('I6','J6','K6','L6','M6','N6'):self.assertIsNone(book.active[address].value)
        book.close()
        self.assertEqual(self.export(payload).status_code,200)
        audit=self.f.db.native("SELECT count(*) FROM workforce_audit WHERE entity_id=%s",(payload['request_key'],)).fetchone()[0]
        self.assertEqual(audit,1)
        self.assertIsNone(self.f.db.native('SELECT accommodation_code FROM workforce_profiles WHERE worker_id=%s',(self.f.worker,)).fetchone()[0])
        for day,clock,departure in (('2026-09-20','',''),('','08:30',''),('','','2026-09-22')):
            payload['request_key']=str(uuid4())
            payload['rows'][0].update(arrival_date=day,arrival_time=clock,departure_date=departure)
            self.assertEqual(self.export(payload).status_code,200)

    def test_revoked_scope_and_personal_export_permissions(self):
        payload=self.payload('rotation');self.assertEqual(self.export(payload,'rotation').status_code,200)
        self.f.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s",(self.f.users['rotation']['id'],))
        self.assertEqual(self.export(payload,'rotation').status_code,404)
        self.assertEqual(self.export(payload,'foreman').status_code,403)
        from user_roles import hr_request_allowed
        self.assertTrue(hr_request_allowed('workforce_accommodation_prepare','POST'))
        self.assertTrue(hr_request_allowed('workforce_accommodation_export','POST'))
        self.assertFalse(hr_request_allowed('workforce_profile','PATCH'))

    def test_category_classification_uses_current_database_hierarchy(self):
        f=self.f
        category=f.db.native("SELECT category_id FROM gdlr_hierarchy_nodes WHERE category_id IS NOT NULL AND name_key='монтажник тт'").fetchone()[0]
        f.db.native('''INSERT INTO employee_gdlr(worker_id,category_id,edit_token,updated_by,updated_at)
            VALUES (%s,%s,%s,%s,'2026-09-18') ON CONFLICT(worker_id) DO UPDATE SET category_id=excluded.category_id''',
            (f.worker,category,f.suffix,f.users['admin']['id']))
        self.assertEqual(self.prepare().json['rows'][0]['worker_category'],'Рабочий')
        root=f.db.native("SELECT category_id FROM gdlr_hierarchy_nodes WHERE name_key='итр'").fetchone()[0]
        if root is not None:
            f.db.native('UPDATE employee_gdlr SET category_id=%s WHERE worker_id=%s',(root,f.worker))
            self.assertEqual(self.prepare().json['rows'][0]['worker_category'],'ИТР')
