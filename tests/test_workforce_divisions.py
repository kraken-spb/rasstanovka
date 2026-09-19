import os
import unittest
from uuid import uuid4
from io import BytesIO
from openpyxl import load_workbook
import test_workforce_api as fixtures


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class DivisionsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f=fixtures.WorkforceApiTest();self.f.setUp();self.addCleanup(self.f.tearDown)
        self.db=self.f.db
        self.pps=self.db.native("INSERT INTO pps_catalog(name,name_key) VALUES (%s,%s) RETURNING id",('ППС test '+self.f.suffix,self.f.suffix)).fetchone()[0]

    def create(self,name='Геодезическая группа',parent=True):
        name+=' '+self.f.suffix
        response=self.f.request('post','divisions',{'name':name,'pps_id':self.pps if parent else None})
        self.assertEqual(response.status_code,201,response.json);return response.json

    def test_parent_unique_and_optimistic_edit_with_audit(self):
        row=self.create()
        self.assertEqual(self.f.request('post','divisions',{'name':row['name'],'pps_id':self.pps}).status_code,409)
        unbound=self.create(parent=False)
        self.assertNotEqual(row['id'],unbound['id'])
        data={'name':'Новая группа','pps_id':self.pps,'active':False,'expected_token':row['edit_token']}
        result=self.f.request('patch','divisions/'+str(row['id']),data)
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(self.f.request('patch','divisions/'+str(row['id']),data).status_code,409)
        self.assertEqual(self.db.native("SELECT count(*) FROM workforce_audit WHERE entity_type='division' AND entity_id=%s",(str(row['id']),)).fetchone()[0],2)

    def test_read_only_and_scoped_roles_cannot_edit(self):
        for role in ['foreman','rotation','recruitment','viewer','hr_viewer']:
            result=self.f.request('post','divisions',{'name':'Запрещено','pps_id':self.pps},role)
            self.assertEqual(result.status_code,403,(role,result.json))
        self.db.native("INSERT INTO user_smu_access(user_id,mode,departments_json,edit_token,updated_at,updated_by) VALUES (%s,'selected','[]','x','2026-09-19',%s)",(self.f.users['admin']['id'],self.f.users['admin']['id']))
        self.assertEqual(self.f.request('post','divisions',{'name':'Запрещено','pps_id':self.pps}).status_code,403)

    def test_smu_has_one_authoritative_name_and_parent(self):
        smu=self.db.native('INSERT INTO smu_catalog(name,pps_id) VALUES (%s,%s) RETURNING id',('СМУ test '+self.f.suffix,self.pps)).fetchone()[0]
        row=dict(self.db.native('SELECT * FROM workforce_divisions WHERE smu_id=%s',(smu,)).fetchone())
        self.assertEqual(row['pps_id'],self.pps)
        blocked=self.f.request('patch','divisions/'+str(row['id']),{'name':'Другой','pps_id':None,'expected_token':str(row['edit_token'])})
        self.assertEqual(blocked.status_code,409)
        self.db.native('UPDATE smu_catalog SET name=%s,pps_id=NULL WHERE id=%s',('Новое СМУ '+self.f.suffix,smu))
        updated=dict(self.db.native('SELECT * FROM workforce_divisions WHERE smu_id=%s',(smu,)).fetchone())
        self.assertIsNone(updated['pps_id']);self.assertNotEqual(updated['edit_token'],row['edit_token'])
        self.assertEqual(updated['name'],'Новое СМУ '+self.f.suffix)

    def test_new_smu_reuses_existing_unit_without_duplicate_or_lost_members(self):
        row=self.create('Новый участок')
        self.db.native('UPDATE workforce_profiles SET division_id=%s WHERE worker_id=%s',(row['id'],self.f.worker))
        smu=self.db.native('INSERT INTO smu_catalog(name,pps_id) VALUES (%s,%s) RETURNING id',(row['name'],self.pps)).fetchone()[0]
        unit=self.db.native('SELECT id FROM workforce_divisions WHERE smu_id=%s',(smu,)).fetchone()[0]
        self.assertEqual(unit,row['id'])
        self.assertEqual(self.db.native('SELECT division_id FROM workforce_profiles WHERE worker_id=%s',(self.f.worker,)).fetchone()[0],row['id'])

    def test_bulk_binding_filters_export_and_scope_keep_smu_unchanged(self):
        row=self.create()
        before=self.db.native('SELECT department FROM workers WHERE id=%s',(self.f.worker,)).fetchone()[0]
        prepared=self.f.request('post','bulk/table/prepare',{'ids':[self.f.worker],'date':'2026-09-19','field':'division_id'})
        self.assertEqual(prepared.status_code,200,prepared.json)
        data=prepared.json
        body={'date':data['date'],'people':[{k:r[k] for k in ('id','token')} for r in data['people']],
            'field':'division_id','value':str(row['id']),'reference_token':data['fields']['division_id']['token'],'extra':{},'request_key':str(uuid4())}
        result=self.f.request('post','bulk/table/apply',body)
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(self.db.native('SELECT department FROM workers WHERE id=%s',(self.f.worker,)).fetchone()[0],before)
        self.assertEqual(self.f.request('delete','divisions/'+str(row['id']),{'expected_token':row['edit_token']}).status_code,409)
        self.f.add_source_record(self.f.worker,'urp:П15')
        suffix='date=2026-09-19&section=rotation&division='+str(row['id'])
        result=self.f.request('get','people?'+suffix)
        self.assertEqual(result.status_code,200,result.json);self.assertEqual(result.json['totals']['total'],1)
        self.assertEqual(result.json['rows'][0]['division'],row['name'])
        pps=self.f.request('get','people?'+suffix+'&pps='+str(self.pps))
        self.assertEqual(pps.json['totals']['total'],1)
        exported=self.f.request('get','people/export?'+suffix)
        self.assertEqual(exported.status_code,200,exported.data[:100])
        book=load_workbook(BytesIO(exported.data));sheet=book.active
        column=[c.value for c in sheet[4]].index('Подразделение')+1
        self.assertEqual(sheet.cell(5,column).value,row['name']);book.close()
        hidden=self.f.request('get','people?'+suffix,role='foreman')
        self.assertEqual(hidden.json['totals']['total'],1)
        self.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s",(self.f.users['foreman']['id'],))
        self.assertEqual(self.f.request('get','people?'+suffix,role='foreman').json['totals']['total'],0)

    def test_archived_parent_not_assignable_and_unknown_filter_rejected(self):
        row=self.create();self.db.native('UPDATE pps_catalog SET active=0 WHERE id=%s',(self.pps,))
        data=self.f.request('post','bulk/table/prepare',{'ids':[self.f.worker],'date':'2026-09-19','field':'division_id'}).json
        self.assertNotIn(str(row['id']),[r['value'] for r in data['fields']['division_id']['options']])
        self.assertEqual(self.f.request('get','people?date=2026-09-19&division=xyz').status_code,400)
