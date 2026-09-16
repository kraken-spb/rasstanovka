import io
import json
import unittest
from openpyxl import load_workbook
import test_export_smu as export_fixtures
import test_placement_report as report_fixtures
import test_personnel_dashboard as dashboard_fixtures
import test_preferences as preference_fixtures


class MultiFilterTest(unittest.TestCase):
    def fixture(self, cls):
        cls.setUpClass(); self.addCleanup(cls.tearDownClass)
        case=cls(); case.setUp(); self.addCleanup(case.doCleanups)
        return case

    def test_export_multiple_values_compose_without_duplicate_rows(self):
        case=self.fixture(export_fixtures.ExportSmuTest); f=case.f
        rows=f.rows(f.export())
        departments=list(dict.fromkeys(r[13] for r in rows if r[13]))[:2]
        contractors=list(dict.fromkeys(r[3] for r in rows if r[3]))[:2]
        categories=list(dict.fromkeys(r[9] or '' for r in rows))[:2]
        query=[('date','2026-09-13'),('shift','1 смена'),('shift','2 смена'),('include_unassigned','1')]
        query += [('department',v) for v in departments]+[('contractor',v) for v in contractors]+[('category',v) for v in categories]
        result=f.fixture.admin.get('/api/staffing/export',query_string=query)
        actual=f.rows(result)
        expected=[r for r in rows if (not departments or r[13] in departments) and (not contractors or r[3] in contractors) and (r[9] or '') in categories]
        self.assertEqual(sorted(r[6] for r in actual),sorted(r[6] for r in expected))
        self.assertEqual(len(actual),int(result.headers['X-Export-Count']))
        self.assertEqual(len(load_workbook(io.BytesIO(result.data)).sheetnames),2)
        self.assertEqual(f.fixture.admin.get('/api/staffing/export',query_string=query+[('shift','bad')]).status_code,400)

    def test_report_pdf_and_calendar_include_multiple_and_blank_categories(self):
        case=self.fixture(report_fixtures.PlacementReportTest);client=case.client(case.admin)
        query=[('date','2026-09-13'),('pps','ППС15'),('pps',''),('category','Электромонтажники'),('category','')]
        response=client.get('/api/placement-report',query_string=query)
        self.assertEqual(response.status_code,200)
        data=response.get_json();self.assertEqual(data['totals']['total'],3)
        self.assertEqual({g['category'] for g in data['groups']},{'Электромонтажники'})
        pdf=client.get('/api/placement-report/pdf',query_string=query)
        self.assertEqual(pdf.status_code,200);self.assertTrue(pdf.data.startswith(b'%PDF'))
        calendar=client.get('/api/calendar',query_string=[('start','2026-09-13'),('days','1'),('category','Электромонтажники'),('category','')])
        self.assertEqual(calendar.status_code,200)
        self.assertEqual(client.get('/api/placement-report',query_string=[('date','2026-09-13')]+[('category',str(i)) for i in range(101)]).status_code,400)

    def test_dashboard_multiple_authors_deduplicates_people(self):
        case=self.fixture(dashboard_fixtures.PersonnelDashboardTest);c=case.case
        ids=c.worker_ids[:2]
        self.assertEqual(c.add(c.crew_a,ids).status_code,200)
        self.assertEqual(c.place(ids,c.sites[0]).status_code,200)
        self.assertEqual(c.place(ids[:1],c.sites[0],client=c.admin,shift='2 смена').status_code,200)
        response=c.admin.get('/api/personnel-dashboard',query_string=[('start','2026-09-11'),('end','2026-09-13'),('user',str(c.owner_a)),('user',str(c.admin_id))])
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.get_json()['counts'],[2,0,0]);self.assertEqual(response.get_json()['unique_count'],2)

    def test_saved_multi_preferences_are_validated_and_isolated(self):
        case=self.fixture(preference_fixtures.PreferencesTest)
        prefix='__multi_filter_v1__:'
        selection=prefix+json.dumps(['name:ЛГСС','name:УПМА'],ensure_ascii=False)
        self.assertEqual(case.save({'contractor':selection}).status_code,200)
        self.assertEqual(case.client_for(case.admin).get('/api/preferences/staffing').get_json()['settings']['contractor'],selection)
        self.assertEqual(case.clients['foreman'].get('/api/preferences/staffing').get_json()['settings'],{})
        for invalid in (prefix+'{}',prefix+'[1]',prefix+'["bad"]'):
            self.assertEqual(case.save({'shift':invalid}).status_code,400)


if __name__=='__main__':unittest.main()
