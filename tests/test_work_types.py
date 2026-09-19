import io
import unittest
import openpyxl
import test_performed_work as work_fixture
import test_staffing_export as export_fixture


class WorkTypesTest(unittest.TestCase):
    setUpClass=classmethod(work_fixture.PerformedWorkTest.setUpClass.__func__)
    tearDownClass=classmethod(work_fixture.PerformedWorkTest.tearDownClass.__func__)
    setUp=work_fixture.PerformedWorkTest.setUp
    client=work_fixture.PerformedWorkTest.client
    write=work_fixture.PerformedWorkTest.write
    apply=work_fixture.PerformedWorkTest.apply
    group_table=work_fixture.PerformedWorkTest.group_table
    payload=work_fixture.PerformedWorkTest.payload

    def kinds(self):return self.admin.get('/api/work-types').json['rows']
    def bind(self,rows,kind=None):
        return self.write('/api/staffing/work-type',{**self.payload(rows),'work_type_id':kind['id'] if kind else None,'work_type_token':kind['edit_token'] if kind else None})

    def test_selected_binding_preserves_description_and_summary_and_clear(self):
        self.apply();rows=self.group_table();kind=self.kinds()[0]
        self.assertEqual(self.write('/api/staffing/performed-work',self.payload(rows[:1],'Текст операций')).status_code,200)
        rows=self.group_table();old=self.payload(rows[:1]);result=self.bind(rows[:1],kind)
        self.assertEqual(result.status_code,200,result.json)
        values=self.group_table();self.assertEqual(values[0]['work_type_id'],kind['id'])
        self.assertEqual(values[0]['performed_work'],'Текст операций');self.assertIsNone(values[1]['work_type_id'])
        summary=self.admin.get('/api/staffing?date=2026-09-12&shift=all&view=summary').json['index']
        self.assertEqual(next(r for r in summary if r['id']==rows[0]['id'])['work_type'],kind['name'])
        self.assertEqual(self.write('/api/staffing/performed-work',old).status_code,409)
        self.assertEqual(self.bind(values[:1]).status_code,200)
        self.assertEqual(self.group_table()[0]['performed_work'],'Текст операций')
        later=self.admin.get('/api/staffing?date=2026-09-13&shift=all').json['rows']
        self.assertTrue(all(r['work_type_id'] is None for r in later))

    def test_catalog_permissions_duplicate_stale_and_disabled_binding(self):
        self.assertEqual({r['name'] for r in self.kinds()},{'Монтаж и сварка МК','Электромонтаж','Благоустройство'})
        headers={'X-CSRF-Token':'staffing-csrf'};kind=self.kinds()[0]
        self.assertEqual(self.foreman.post('/api/work-types',json={'name':'Новый'},headers=headers).status_code,403)
        self.assertEqual(self.admin.post('/api/work-types',json={'name':kind['name'].upper()},headers=headers).status_code,409)
        body={'name':kind['name'],'active':False,'expected_token':kind['edit_token']}
        self.assertEqual(self.admin.patch('/api/work-types/'+str(kind['id']),json=body,headers=headers).status_code,200)
        self.assertEqual(self.admin.patch('/api/work-types/'+str(kind['id']),json=body,headers=headers).status_code,409)
        self.apply();self.assertEqual(self.bind(self.group_table()[:1],kind).status_code,400)

    def test_batch_scope_and_stale_tokens(self):
        self.apply();rows=self.group_table();kind=self.kinds()[0]
        data={**self.payload(rows),'work_type_id':kind['id'],'work_type_token':kind['edit_token']}
        self.assertEqual(self.write('/api/staffing/work-type',data,self.foreman).status_code,403)
        self.assertEqual(self.admin.put('/api/staffing/work-type',json=data).status_code,403)
        self.assertEqual(self.bind(rows[:1],kind).status_code,200)
        self.assertEqual(self.write('/api/staffing/work-type',data).status_code,409)
        self.assertTrue(all(r['work_type_id'] is None for r in self.group_table()[1:]))


class WorkTypeExportTest(unittest.TestCase):
    setUpClass=classmethod(export_fixture.StaffingExportTest.setUpClass.__func__)
    tearDownClass=classmethod(export_fixture.StaffingExportTest.tearDownClass.__func__)
    setUp=export_fixture.StaffingExportTest.setUp
    client=export_fixture.StaffingExportTest.client
    user=export_fixture.StaffingExportTest.user
    worker=export_fixture.StaffingExportTest.worker
    assignment=export_fixture.StaffingExportTest.assignment

    def test_filter_applies_to_both_tabs_and_empty_and_scope(self):
        with self.module.app.app_context():
            db=self.module.get_db();kind=db.execute('SELECT * FROM work_types ORDER BY id').fetchone()
            for worker in [self.visible_worker,self.hidden_worker]:
                db.execute('''INSERT INTO staffing_performed_work(work_date,worker_id,shift,description,edit_token,updated_by,updated_at,work_type_id)
                    VALUES ('2026-09-13',?,'1 смена','','work',?,'now',?)''',(worker,self.admin_id,kind['id']))
            db.commit();key=kind['id'];name=kind['name']
        response=self.foreman.get('/api/staffing/export?date=2026-09-13&shift=all&work_type_id='+str(key))
        self.assertEqual(response.status_code,200,response.json);self.assertEqual(response.headers['X-Export-Count'],'1')
        book=openpyxl.load_workbook(io.BytesIO(response.data));sheet=book['Список сотрудников']
        self.assertEqual(sheet.cell(1,16).value,'Вид работ');self.assertEqual(sheet.cell(2,16).value,name)
        self.assertEqual(sheet.max_row,2);self.assertEqual(len(book.sheetnames),2)
        options=self.foreman.get('/api/staffing/export/options?date=2026-09-13&shift=all').json
        self.assertIn({'value':str(key),'label':name},options['work_types'])
        none=self.foreman.get('/api/staffing/export?date=2026-09-13&shift=all&work_type_id=__staffing_export_empty__')
        self.assertEqual(none.headers['X-Export-Count'],'2')
