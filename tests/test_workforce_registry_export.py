import os
import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import test_workforce_api as fixtures
from workforce_registry_export import HEADERS, registry_workbook


class RegistryWorkbookTest(unittest.TestCase):
    def test_dates_text_identifiers_formula_literals_and_empty_sheet(self):
        for section, title in [('rotation', 'Перевахта'), ('recruitment', 'Комплектация')]:
            with self.subTest(section=section):
                row = {'full_name': '=HYPERLINK("https://example.invalid")', 'personnel_no': '000825',
                       'profession': '+Опасная формула', 'arrival_date': '2026-09-01',
                       'movement': {'direction': 'arrival', 'planned_date': '2026-09-21', 'basis': 'по билету'}}
                book = load_workbook(BytesIO(registry_workbook([row], '2026-09-17', section)))
                sheet = book[title]
                self.assertEqual([cell.value for cell in sheet[4]], list(HEADERS))
                self.assertEqual(sheet.cell(5, HEADERS.index('ФИО') + 1).data_type, 's')
                self.assertEqual(sheet.cell(5, HEADERS.index('ФИО') + 1).value, row['full_name'])
                self.assertEqual(sheet.cell(5, HEADERS.index('Табельный номер') + 1).value, '000825')
                self.assertEqual(sheet.cell(5, HEADERS.index('Дата заезда') + 1).value, datetime(2026, 9, 1))
                self.assertEqual(sheet.cell(5, HEADERS.index('Плановая дата поездки') + 1).value, datetime(2026, 9, 21))
                self.assertEqual(sheet.cell(5, HEADERS.index('Состояние') + 1).value, 'Без подтверждённого состояния')
                self.assertEqual(sheet.freeze_panes, 'D5')
                self.assertEqual(sheet.auto_filter.ref, f'A4:{get_column_letter(len(HEADERS))}5')
                self.assertEqual(sheet['B2'].value, 1)
                book.close()
                empty = load_workbook(BytesIO(registry_workbook([], '2026-09-17', section)))
                self.assertEqual(empty[title].max_row, 4)
                self.assertEqual(empty[title]['B2'].value, 0)
                empty.close()


    def test_reuses_the_same_row_styles_without_losing_safe_text_or_dates(self):
        rows = [
            {'full_name': '=Первая', 'personnel_no': '001', 'arrival_date': '2026-09-01'},
            {'full_name': '+Вторая', 'personnel_no': '002', 'arrival_date': '2026-09-02'},
            {'full_name': '@Третья', 'personnel_no': '003', 'arrival_date': '2026-09-03'},
        ]
        book = load_workbook(BytesIO(registry_workbook(rows, '2026-09-17', 'rotation')))
        sheet = book.active
        self.assertEqual(sheet.cell(5, HEADERS.index('ФИО') + 1).data_type, 's')
        self.assertEqual(sheet['B5'].style_id, sheet['C5'].style_id)
        self.assertEqual(sheet['B5'].style_id, sheet['B7'].style_id)
        self.assertNotEqual(sheet['B5'].style_id, sheet['B6'].style_id)
        self.assertEqual(sheet.cell(5, HEADERS.index('Дата заезда') + 1).number_format, 'dd.mm.yyyy')
        arrival = HEADERS.index('Дата заезда') + 1
        self.assertEqual(sheet.cell(5, arrival).style_id, sheet.cell(7, arrival).style_id)
        self.assertNotEqual(sheet.cell(5, arrival).style_id, sheet.cell(6, arrival).style_id)
        self.assertEqual(sheet.cell(6, arrival).fill.fgColor.rgb, sheet['B6'].fill.fgColor.rgb)
        book.close()


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class RegistryExportApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        for worker in self.f.ids:
            self.f.add_source_record(worker, 'urp:П15')
        self.f.add_source_record(self.f.worker, 'urp:К15', role='recruitment', sheet='ПВП')

    def export(self, extra='', role='admin', section='rotation'):
        return self.f.request('get', f'people/export?date=2026-09-17&section={section}&' + extra, role=role)

    def test_export_all_filtered_rows_ignores_page_and_enforces_section_and_scope(self):
        f = self.f
        for role, section, count in [('admin', 'rotation', 2), ('rotation', 'rotation', 1),
                                     ('recruitment', 'recruitment', 1), ('foreman', 'rotation', 1)]:
            with self.subTest(role=role, section=section):
                response = self.export('department=TEST-SMU&department=OTHER-SMU&offset=1&limit=1', role, section)
                self.assertEqual(response.status_code, 200, response.data[:200])
                self.assertEqual(response.headers['X-Export-Row-Count'], str(count))
                self.assertEqual(response.headers['Cache-Control'], 'private, no-store')
                book = load_workbook(BytesIO(response.data))
                self.assertEqual(book.active.max_row - 4, count)
                self.assertEqual(book.active['B2'].value, count)
                book.close()
        f.db.native("UPDATE user_smu_access SET departments_json='[]' WHERE user_id=%s", (f.users['rotation']['id'],))
        self.assertEqual(self.export(role='rotation').headers['X-Export-Row-Count'], '0')
        f.db.native('UPDATE users SET active=0 WHERE id=%s', (f.users['rotation']['id'],))
        self.assertEqual(self.export(role='rotation').status_code, 403)

    def test_export_filters_match_registry_date_stage_regex_and_queue(self):
        f = self.f
        f.add_stage_record(f.worker, 'stage.pvp', '2026-09-16')
        f.add_stage_record(f.worker, 'stage.onsite', '2026-09-18')
        for extra in ['stage=stage.pvp&stage=unconfirmed', 'stage=stage.onsite',
                      'department=TEST-SMU&queue=pvp', 'conflicts=1', 'q=Тест.*&regex=1']:
            listing = f.request('get', 'people?date=2026-09-17&section=rotation&department=TEST-SMU&department=OTHER-SMU&' + extra)
            response = self.export('department=TEST-SMU&department=OTHER-SMU&' + extra)
            self.assertEqual(response.status_code, 200, response.data[:200])
            self.assertEqual(int(response.headers['X-Export-Row-Count']), listing.json['totals']['total'])
        self.assertEqual(self.export(section='').status_code, 400)
        self.assertEqual(self.export(section='bad').status_code, 400)
        with patch('workforce_registry_export.MAX_ROWS', 1):
            self.assertEqual(self.export('department=TEST-SMU&department=OTHER-SMU').status_code, 422)
        self.assertEqual(self.export('q=(&regex=1').status_code, 400)
