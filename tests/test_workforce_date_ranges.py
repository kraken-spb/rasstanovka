"""Date filters match displayed cells before pagination, including Excel and scope."""
from io import BytesIO
import os
import unittest
from urllib.parse import urlencode
from uuid import uuid4

from openpyxl import load_workbook
import test_workforce_api as fixtures


@unittest.skipUnless(os.getenv('CREW_POSTGRES_TEST_ENV'), 'Select an isolated PostgreSQL database.')
class WorkforceDateRangesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.WorkforceApiTest.setUpClass()

    def setUp(self):
        self.f = fixtures.WorkforceApiTest()
        self.f.setUp()
        self.addCleanup(self.f.tearDown)
        for worker in self.f.ids:
            self.f.add_source_record(worker, 'urp:П15')
        self.f.db.native("UPDATE workforce_profiles SET arrival_date='2024-02-29' WHERE worker_id=%s", (self.f.worker,))

    def get(self, values=(), role='admin', export=False):
        query = [('date','2026-09-18'),('section','rotation'),('department','TEST-SMU'),('department','OTHER-SMU'),*values]
        result = self.f.request('get', ('people/export?' if export else 'people?') + urlencode(query), role=role)
        self.assertEqual(result.status_code, 200, result.data[:300])
        return result if export else result.json

    def test_inclusive_bounds_open_bounds_nulls_pagination_cache_and_scope(self):
        period = [('arrival_date_from','2024-02-01'),('arrival_date_to','2024-02-29')]
        rows = self.get(period + [('limit','1')])
        self.assertEqual([r['id'] for r in rows['rows']], [self.f.worker])
        self.assertEqual(self.get(period + [('offset','1'),('limit','1')])['rows'], [])
        self.assertEqual(rows, self.get(period + [('limit','1')]))
        self.assertEqual(self.get([('arrival_date_from','2024-03-01')])['totals']['total'], 0)
        self.assertEqual(self.get([('arrival_date_to','2024-02-29')])['totals']['total'], 1)
        self.assertEqual(self.get(period + [('arrival_date_empty','include')])['totals']['total'], 2)
        self.assertEqual(self.get(period + [('arrival_date_empty','include')], role='rotation')['totals']['total'], 1)
        self.assertEqual(self.get([('arrival_date_empty','only')])['rows'][0]['id'], self.f.ids[1])
        self.assertEqual(self.get([('arrival_date_empty','exclude')])['rows'][0]['id'], self.f.worker)
        # No bounds with "include" must not leave unused SQL parameters.
        self.assertEqual(self.get([('forecast_departure_date_empty','include')])['totals']['total'], 2)
        export = self.get(period, export=True)
        book = load_workbook(BytesIO(export.data))
        self.assertEqual(export.headers['X-Export-Row-Count'], '1')
        self.assertEqual(book.active.max_row, 5)
        book.close()

    def test_planned_date_matches_first_pending_not_cancelled_replaced_or_later(self):
        f = self.f
        def movement(day, result=None, previous=None):
            return f.db.native('''INSERT INTO workforce_movements(worker_id,direction,planned_date,
                result_code,rescheduled_from,created_by,updated_by,request_key)
                VALUES (%s,'arrival',%s,%s,%s,%s,%s,%s) RETURNING id''',
                (f.worker,day,result,previous,f.users['admin']['id'],f.users['admin']['id'],uuid4())).fetchone()[0]
        old=movement('2026-09-01')
        movement('2026-09-02','result.cancelled')
        movement('2026-09-21',previous=old)
        movement('2026-10-10')
        selected=self.get([('planned_date_from','2026-09-21'),('planned_date_to','2026-09-21')])
        self.assertEqual(selected['rows'][0]['movement']['planned_date'],'2026-09-21')
        for start,end in [('2026-09-01','2026-09-20'),('2026-10-01','2026-10-31')]:
            self.assertEqual(self.get([('planned_date_from',start),('planned_date_to',end)])['rows'],[])
        self.assertEqual(self.get([('planned_date_empty','only')])['rows'][0]['id'],f.ids[1])

    def test_stage_date_is_effective_as_of_report_day(self):
        f=self.f
        for day,stage in [('2026-09-01','stage.leave'),('2026-10-01','stage.inbound')]:
            f.add_stage_record(f.worker,stage,day)
        rows=self.get([('stage_date_from','2026-09-01'),('stage_date_to','2026-09-30')])['rows']
        self.assertEqual([r['id'] for r in rows],[f.worker])
        self.assertEqual(rows[0]['effective_date'],'2026-09-01')
        self.assertEqual(self.get([('stage_date_from','2026-10-01')])['rows'],[])

    def test_rotation_dates_and_forecast_match_displayed_rotation(self):
        f=self.f; actor=f.users['admin']['id']
        schedule=f.db.native('''INSERT INTO workforce_rotation_schedules(name,onsite_days,leave_days,travel_days,updated_by)
            VALUES (%s,30,30,2,%s) RETURNING id''',(f.suffix,actor)).fetchone()[0]
        f.db.native('''INSERT INTO workforce_rotations(worker_id,schedule_id,schedule_snapshot,start_date,planned_end_date,
            leave_end_date,next_arrival_date,created_by,updated_by,request_key)
            VALUES (%s,%s,'{}','2026-09-01','2026-09-30','2026-10-30','2026-11-01',%s,%s,%s)''',
            (f.worker,schedule,actor,actor,uuid4()))
        f.db.native("UPDATE workforce_profiles SET forecast_departure_date='2026-12-31' WHERE worker_id=%s",(f.worker,))
        conditions=[('forecast_departure_date_from','2026-09-01'),('forecast_departure_date_to','2026-09-30'),
                    ('leave_end_date_from','2026-10-01'),('leave_end_date_to','2026-10-31'),
                    ('next_arrival_date_from','2026-11-01'),('next_arrival_date_to','2026-11-01')]
        self.assertEqual(self.get(conditions)['rows'][0]['id'],f.worker)
        self.assertEqual(self.get([('forecast_departure_date_from','2026-12-01')])['rows'],[])
        self.assertEqual(self.get([('leave_end_date_empty','only')])['rows'][0]['id'],f.ids[1])

    def test_invalid_ambiguous_or_reversed_filters_fail_closed(self):
        for key in ('stage_date','arrival_date','forecast_departure_date','planned_date','leave_end_date','next_arrival_date'):
            for values in [[(key+'_from','2026-02-30')],[(key+'_from','2026-10-01'),(key+'_to','2026-09-01')],
                           [(key+'_from','2026-09-01'),(key+'_from','2026-10-01')],[(key+'_empty','bad')],
                           [(key,'2026-09-01'),(key+'_to','2026-10-01')],[(key+'_from','2026-09-01'),(key+'_empty','only')]]:
                with self.subTest(key=key,values=values):
                    self.assertEqual(self.f.request('get','people?'+urlencode(values)).status_code,400)
